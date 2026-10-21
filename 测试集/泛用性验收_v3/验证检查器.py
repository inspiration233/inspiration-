# -*- coding: utf-8 -*-
"""用临时副本中的真实错误验证检查器；正式夹具永不修改。

运行：python 验证检查器.py [--root 案例根目录]
先要求正式基线通过，再分别注入错误。任一错误未被对应检查组拒绝即非零退出。
"""
import argparse
from contextlib import contextmanager
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.styles import PatternFill


ROOT = Path(__file__).resolve().parent
SPEC = importlib.util.spec_from_file_location("fixture_oracle", ROOT / "验收断言.py")
ORACLE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(ORACLE)


def artifact(root, cid, role, contains=None):
    d = root / cid
    names = json.loads((d / "case.json").read_text(encoding="utf-8"))["artifacts"][role]
    if contains:
        names = [name for name in names if contains in name]
    if not names:
        raise AssertionError(f"{cid}/{role}: 找不到反例目标文件")
    return d / names[0]


def edit(path, operation):
    wb = load_workbook(path)
    operation(wb)
    wb.save(path)
    wb.close()


def task_row(ws, tid):
    for row in ws:
        if row[0].value == tid:
            return row[0].row
    raise AssertionError(f"反例目标任务不存在: {tid}")


def change_cell(root, cid, role, tid, column, value, contains=None):
    def operation(wb):
        ws = wb.active
        ws.cell(task_row(ws, tid), column).value = value
    edit(artifact(root, cid, role, contains), operation)


def no_monthly(root):
    artifact(root, "W1", "expected_monthly").unlink()


def no_task(root):
    edit(artifact(root, "W1", "expected_annual"),
         lambda wb: wb.active.delete_rows(task_row(wb.active, "T002")))


def wrong_formula(root):
    change_cell(root, "W1", "expected_annual", "T002", 6, "=E999*100")


def no_column(root):
    edit(artifact(root, "W1", "expected_monthly"), lambda wb: wb.active.delete_cols(6))


def no_validation(root):
    edit(artifact(root, "W1", "split"), lambda wb: wb.active.data_validations.dataValidation.clear())


def wrong_validation(root):
    def operation(wb):
        wb.active.data_validations.dataValidation[0].formula1 = '"错误选项"'
    edit(artifact(root, "W1", "split"), operation)


def wrong_cf(root):
    def operation(wb):
        cf = wb.active.conditional_formatting
        first = next(iter(cf))
        cf[first][0].formula = ['"错误状态"']
    edit(artifact(root, "W1", "split"), operation)


def wrong_cf_color(root):
    def operation(wb):
        cf = wb.active.conditional_formatting
        cf[next(iter(cf))][0].dxf.fill = PatternFill(fill_type="solid", fgColor="FF0000")
    edit(artifact(root, "W1", "split"), operation)


def no_freeze(root):
    edit(artifact(root, "W1", "split"), lambda wb: setattr(wb.active, "freeze_panes", None))


def wrong_json_old(root):
    path = root / "W2" / "case.json"
    document = json.loads(path.read_text(encoding="utf-8"))
    document["expected_changes"][0]["old"] = 999
    path.write_text(json.dumps(document, ensure_ascii=False, indent=2), encoding="utf-8")


def wrong_width(root):
    edit(artifact(root, "W1", "split"),
         lambda wb: setattr(wb.active.column_dimensions["A"], "width", 99))


def wrong_header_merge(root):
    def operation(wb):
        ws = wb.active
        ws.unmerge_cells("C3:D3")
        ws.merge_cells("C20:D20")  # 合并数量没有减少，但表头的结构已错。
    edit(artifact(root, "W3", "split"), operation)


def wrong_feedback_unit(root):
    change_cell(root, "W3", "feedback", "A006", 3, "工程技术部")


def wrong_february(root):
    change_cell(root, "W2", "feedback", "组织安全培训演练", 3, 99)


def wrong_w2_monthly(root):
    path = artifact(root, "W2", "expected_monthly")
    def operation(wb):
        ws = wb.active
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value == "财务部" and ws.cell(r, 2).value == "组织安全培训演练":
                ws.cell(r, 3).value = 18
                return
        raise AssertionError("W2月汇总反例目标缺失")
    edit(path, operation)


def wrong_w6_snapshot(root):
    path = artifact(root, "W6", "expected_monthly")
    def operation(wb):
        ws = wb.active
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 2).value == "C003":
                ws.cell(r, 1).value = "3月"
                return
        raise AssertionError("W6月汇总反例目标缺失")
    edit(path, operation)


def no_f23_validation(root):
    edit(root / "F23" / "F23_条件格式验证.xlsx",
         lambda wb: wb.active.data_validations.dataValidation.clear())


def wrong_f23_range(root):
    def operation(wb):
        for dv in wb.active.data_validations.dataValidation:
            if dv.type == "decimal":
                dv.sqref = "B2:B8"
    edit(root / "F23" / "F23_条件格式验证.xlsx", operation)


def wrong_f23_formula(root):
    def operation(wb):
        for dv in wb.active.data_validations.dataValidation:
            if dv.type == "custom":
                dv.formula1 = "TRUE"
    edit(root / "F23" / "F23_条件格式验证.xlsx", operation)


def wrong_d40_formula(root):
    def operation(wb):
        wb["换列后"]["E2"] = "=SUM(A2:C2)"  # 错误地把原D=1000纳入求和。
    edit(root / "D40" / "D40_列换序引用.xlsx", operation)


@contextmanager
def checked_temporary_directory():
    temp_parent = Path(tempfile.gettempdir()).resolve()
    target = Path(tempfile.mkdtemp(prefix="excel-oracle-", dir=temp_parent)).resolve()
    if target.parent != temp_parent or not target.name.startswith("excel-oracle-"):
        raise RuntimeError(f"拒绝对未验证的临时路径进行操作或清理: {target}")
    try:
        yield target
    finally:
        # 递归清理前重新验证最终绝对路径，防止路径漂移到其他目录。
        resolved = target.resolve()
        if resolved != target or resolved.parent != temp_parent or not resolved.name.startswith("excel-oracle-"):
            raise RuntimeError(f"拒绝清理路径已改变的临时目录: {resolved}")
        shutil.rmtree(resolved)


MUTATIONS = (
    ("W1缺月度文件", "W1", "业务数据", no_monthly),
    ("W1整条T002缺失", "W1", "业务数据", no_task),
    ("W1公式引用错行", "W1", "业务数据", wrong_formula),
    ("W1月汇总缺列", "W1", "业务数据", no_column),
    ("W1下拉全部丢失", "W1", "原生属性", no_validation),
    ("W1下拉内容错误", "W1", "原生属性", wrong_validation),
    ("W1条件格式内容错误", "W1", "原生属性", wrong_cf),
    ("W1仅条件颜色错误", "W1", "原生属性", wrong_cf_color),
    ("W1冻结窗格丢失", "W1", "原生属性", no_freeze),
    ("W1列宽错误但非空", "W1", "原生属性", wrong_width),
    ("W3合并数量正确但范围错误", "W3", "原生属性", wrong_header_merge),
    ("W3反馈串入其他单位", "W3", "业务数据", wrong_feedback_unit),
    ("W2误改二月", "W2", "业务数据", wrong_february),
    ("W2一月汇总漏应用修改", "W2", "业务数据", wrong_w2_monthly),
    ("W2仅JSON旧值错误", "W2", "业务数据", wrong_json_old),
    ("W6使用反馈前月份", "W6", "业务数据", wrong_w6_snapshot),
    ("F23验证全部丢失", "F23", "公式与规则", no_f23_validation),
    ("F23验证范围不完整", "F23", "公式与规则", wrong_f23_range),
    ("F23自定义公式错误", "F23", "公式与规则", wrong_f23_formula),
    ("D40换列后计算对象改变", "D40", "公式与规则", wrong_d40_formula),
)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    args = parser.parse_args()
    baseline = ORACLE.run(args.root)
    failures = [message for kind, message, _ in baseline if kind != "pass"]
    cli = [sys.executable, "-B", str(ROOT / "自检.py"), "--quiet"]
    child_env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONDONTWRITEBYTECODE="1")
    baseline_cli = subprocess.run(cli + ["--root", str(args.root.resolve())],
                                  capture_output=True, text=True, encoding="utf-8", errors="replace",
                                  env=child_env)
    if baseline_cli.returncode != 0:
        failures.append(f"基线CLI退出码={baseline_cli.returncode}\n{baseline_cli.stdout}\n{baseline_cli.stderr}")
    if failures:
        print("基线尚未通过，不能把原有失败当作反例检测成功：")
        for message in failures:
            print("  " + message)
        return 1
    missed = []
    for name, cid, category, mutate in MUTATIONS:
        with checked_temporary_directory() as target:
            # 完整复制42案例，避免CLI仅因缺案例而退出非零造成虚假反例成功。
            for source in args.root.glob("*/case.json"):
                shutil.copytree(source.parent, target / source.parent.name)
            mutate(target)
            rejected = any(kind == "fail" and message.startswith(cid) and group == category
                           for kind, message, group in ORACLE.run(target))
            result_file = target / "cli-result.json"
            cli_result = subprocess.run(cli + ["--root", str(target), "--json", str(result_file)],
                                        capture_output=True, text=True, encoding="utf-8", errors="replace",
                                        env=child_env)
            cli_rejected = False
            if cli_result.returncode == 1 and result_file.exists():
                report = json.loads(result_file.read_text(encoding="utf-8"))
                cli_rejected = any(item["status"] == "fail"
                                   and item["message"].startswith(cid)
                                   and item["category"] == category
                                   for item in report["results"])
            if rejected and cli_rejected:
                print(f"已拒绝：{name}")
            else:
                missed.append(name)
                print(f"漏检：{name}（内部拒绝={rejected}，CLI拒绝={cli_rejected}）")
    print(f"检查器反例验证：{len(MUTATIONS)-len(missed)}已拒绝 / {len(missed)}漏检")
    return 1 if missed else 0


if __name__ == "__main__":
    raise SystemExit(main())
