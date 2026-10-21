# -*- coding: utf-8 -*-
"""夹具验收预言机：只读工作簿，不导入产品引擎、生成器或历史通过记录。

run(root) 返回 (pass/fail/skip, message, category) 元组列表。
这里验收的是已约定夹具及预期答案的自洽性，不代表产品功能已实现。
"""
from __future__ import annotations

from copy import copy, deepcopy
import json
from pathlib import Path
from xml.etree.ElementTree import tostring

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import MultiCellRange


ROLES = ("original", "split", "feedback", "expected_monthly", "expected_annual")
W1_HEADER = ("任务编号", "任务名称", "责任单位", "月份", "计划数量", "预算金额(元)", "状态")
W6_HEADER = ("编号", "任务名称", "责任单位", "起始月", "结束月", "数量")


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def equal(actual, expected, label):
    require(actual == expected, f"{label}: 实际 {actual!r}；期望 {expected!r}")


def xml(obj):
    return None if obj is None else tostring(copy(obj).to_tree(), encoding="unicode")


class Case:
    def __init__(self, root, cid):
        self.directory = Path(root) / cid
        self.cid = cid
        self.document = json.loads((self.directory / "case.json").read_text(encoding="utf-8"))
        self.books = []

    def files(self, role, count=None):
        values = self.document.get("artifacts", {}).get(role)
        require(isinstance(values, list) and values, f"{self.cid}: 缺少 artifacts.{role} 文件清单")
        require(all(isinstance(name, str) for name in values), f"{self.cid}/{role}: 文件名必须是字符串")
        require(len(set(values)) == len(values), f"{self.cid}/{role}: 文件清单重复")
        if count is not None:
            equal(len(values), count, f"{self.cid}/{role} 文件数")
        paths = []
        for name in values:
            path = (self.directory / name).resolve()
            require(path.is_relative_to(self.directory.resolve()), f"{self.cid}: 文件越出案例目录")
            require(path.is_file(), f"{self.cid}/{role}: 缺文件 {name}")
            paths.append(path)
        return paths

    def open(self, path):
        wb = load_workbook(path, data_only=False)
        self.books.append(wb)
        require(len(wb.worksheets) == 1, f"{path.name}: 本案例应恰有一个工作表")
        return wb.active

    def one(self, role):
        return self.open(self.files(role, 1)[0])

    def check_manifest(self):
        names = [p for role in ROLES for p in self.files(role)]
        equal(len(names), len(set(names)), f"{self.cid}: 不同角色不能共用同一文件")
        equal(set(names), set(p.resolve() for p in self.directory.glob("*.xlsx")),
              f"{self.cid}: artifacts 必须覆盖全部 xlsx，不能残留未声明的旧答案")

    def selection(self, month, rule, source):
        spec = self.document.get("selection", {})
        equal(spec.get("target_year"), 2026, f"{self.cid}: 年份")
        equal(spec.get("target_month"), month, f"{self.cid}: 目标月份")
        equal(spec.get("month_rule"), rule, f"{self.cid}: 月份归属口径")
        equal(spec.get("month_source"), source, f"{self.cid}: 月份来源")
        equal(spec.get("snapshot"), "after_feedback", f"{self.cid}: 结果时间点")

    def close(self):
        for wb in self.books:
            wb.close()


def rows(ws, first, width, key_col=0):
    require(ws.max_column == width, f"{ws.title}: 列数 {ws.max_column}，期望 {width}")
    result, coordinates = {}, {}
    for r in range(first, ws.max_row + 1):
        values = tuple(ws.cell(r, c).value for c in range(1, width + 1))
        if all(v is None for v in values):
            continue
        key = values[key_col]
        require(key not in (None, ""), f"{ws.title}!{r}: 任务标识为空")
        require(key not in result, f"{ws.title}: 任务 {key!r} 重复")
        result[key] = values
        coordinates[key] = r
    return result, coordinates


def native_rows(ws, source, formula_col, formula_for):
    width = source.max_column
    for r in range(1, 5):
        equal(tuple(ws.cell(r, c).value for c in range(1, width + 1)),
              tuple(source.cell(r, c).value for c in range(1, width + 1)),
              f"{ws.title}: 第{r}行前置区/表头")
    records, positions = rows(ws, 5, width)
    for tid, values in records.items():
        equal(values[formula_col - 1], formula_for(positions[tid]),
              f"{ws.title}/{tid}: 公式")
    # 公式的物理行号不作为业务内容；公式已独立逐条精确验证。
    business = {tid: tuple(v for i, v in enumerate(vals, 1) if i != formula_col)
                for tid, vals in records.items()}
    return business, positions


def check_w1(case):
    case.check_manifest()
    case.selection(6, "exact_month", "月份")
    equal(case.document.get("feedback_changes"), [
        {"task": "T002", "field": "计划数量", "old": 20, "new": 25},
        {"task": "T999", "field": "新增", "old": None, "new": "新增任务"},
        {"task": "T005", "field": "状态", "old": "进行中", "new": "已完成"},
    ], "W1 JSON反馈变更")
    equal(case.document.get("ground_truth", {}).get("june_tasks"), 1, "W1 JSON六月基线数")
    equal(case.document.get("expected_monthly_ids"), {"6月": ["T006", "T999"]},
          "W1 JSON反馈后六月集合")
    source = case.one("original")
    equal(tuple(c.value for c in source[4]), W1_HEADER, "W1原始表头")
    base, _ = native_rows(source, source, 6, lambda r: f"=E{r}*100")
    equal(set(base), {f"T{i:03d}" for i in range(1, 13)}, "W1原始任务集合")
    equal(base["T002"][4], 20, "W1基线T002数量")
    equal(base["T005"][5], "进行中", "W1基线T005状态")
    equal({k for k, v in base.items() if v[3] == "6月"}, {"T006"}, "W1六月基线集合")
    units = {v[2] for v in base.values()}
    seen = set()
    for path in case.files("split", len(units)):
        ws = case.open(path)
        require(ws.title in units and ws.title not in seen, f"W1拆分单位重复/未知: {ws.title}")
        seen.add(ws.title)
        actual, _ = native_rows(ws, source, 6, lambda r: f"=E{r}*100")
        equal(actual, {k: v for k, v in base.items() if v[2] == ws.title}, f"W1拆分/{ws.title}")
    equal(seen, units, "W1拆分单位集合")

    updated = dict(base)
    change = list(updated["T002"]); change[4] = 25; updated["T002"] = tuple(change)
    change = list(updated["T005"]); change[5] = "已完成"; updated["T005"] = tuple(change)
    updated["T999"] = ("T999", "新增：设备升级改造", "工程技术部", "6月", 5, "进行中")
    feedback_units = {"财务部", "工程技术部", "市场营销部"}
    seen = set()
    for path in case.files("feedback", 3):
        ws = case.open(path)
        require(ws.title in feedback_units and ws.title not in seen, f"W1反馈单位重复/未知: {ws.title}")
        seen.add(ws.title)
        actual, _ = native_rows(ws, source, 6, lambda r: f"=E{r}*100")
        equal(actual, {k: v for k, v in updated.items() if v[2] == ws.title}, f"W1反馈/{ws.title}")
    equal(seen, feedback_units, "W1反馈单位集合")
    monthly, _ = native_rows(case.one("expected_monthly"), source, 6, lambda r: f"=E{r}*100")
    equal(monthly, {k: v for k, v in updated.items() if v[3] == "6月"}, "W1六月结果")
    annual, _ = native_rows(case.one("expected_annual"), source, 6, lambda r: f"=E{r}*100")
    equal(annual, updated, "W1全年结果及未反馈任务")


def matrix_records(ws, header_row):
    header = ("责任单位", "任务名称", *(f"{m}月" for m in range(1, 13)), "合计")
    equal(tuple(c.value for c in ws[header_row]), header, f"{ws.title}: 月份矩阵表头")
    result = {}
    unit = None
    for r in range(header_row + 1, ws.max_row + 1):
        first = ws.cell(r, 1).value
        if first == "合计":
            for c in range(3, 16):
                letter = get_column_letter(c)
                equal(ws.cell(r, c).value, f"=SUM({letter}{header_row+1}:{letter}{r-1})",
                      f"{ws.title}!{letter}{r}: 总合计公式")
            continue
        unit = first if first is not None else unit
        task = ws.cell(r, 2).value
        if task is None and first is None:
            continue
        require(unit and task, f"{ws.title}!{r}: 缺单位或任务名称")
        key = (unit, task)
        require(key not in result, f"{ws.title}: 重复任务 {key}")
        result[key] = tuple(ws.cell(r, c).value for c in range(3, 15))
        equal(ws.cell(r, 15).value, f"=SUM(C{r}:N{r})", f"{ws.title}/{key}: 全年合计公式")
    return result


def split_matrix(ws):
    header = ("任务名称", *(f"{m}月" for m in range(1, 13)), "合计")
    equal(tuple(c.value for c in ws[1]), header, f"{ws.title}: 拆分月份表头")
    result = {}
    for r in range(2, ws.max_row + 1):
        task = ws.cell(r, 1).value
        require(task and task not in result, f"{ws.title}: 空任务或重复任务")
        result[task] = tuple(ws.cell(r, c).value for c in range(2, 14))
        equal(ws.cell(r, 14).value, f"=SUM(B{r}:M{r})", f"{ws.title}/{task}: 合计公式")
    return result


def check_w2(case):
    case.check_manifest()
    case.selection(1, "exact_month", "column_header")
    equal(case.document.get("expected_changes"), [
        {"unit": "财务部", "task": "组织安全培训演练", "field": "1月数量", "old": 18, "new": 99}
    ], "W2 JSON变更身份/字段/旧值/新值")
    base = matrix_records(case.one("original"), 2)
    units = {"综合办公室", "财务部", "工程技术部"}
    equal({k[0] for k in base}, units, "W2原始单位")
    equal(len(base), 12, "W2原始任务数")
    key = ("财务部", "组织安全培训演练")
    require(key in base, "W2基线缺少指定反馈任务")
    equal(base[key][:2], (18, 34), "W2反馈前1/2月基线")
    seen = set()
    for path in case.files("split", 3):
        ws = case.open(path)
        require(ws.title in units and ws.title not in seen, f"W2拆分单位重复/未知: {ws.title}")
        seen.add(ws.title)
        equal(split_matrix(ws), {k[1]: v for k, v in base.items() if k[0] == ws.title},
              f"W2拆分/{ws.title}")
    updated = dict(base)
    updated[key] = (99,) + base[key][1:]
    feedback = case.one("feedback")
    equal(feedback.title, "财务部", "W2反馈单位")
    equal(split_matrix(feedback), {k[1]: v for k, v in updated.items() if k[0] == "财务部"},
          "W2反馈仅1月修改，2月及其他任务保持")
    monthly = case.one("expected_monthly")
    equal(tuple(c.value for c in monthly[1]), ("责任单位", "任务名称", "1月数量"), "W2月汇总表头")
    actual = {}
    for r in range(2, monthly.max_row + 1):
        k = (monthly.cell(r, 1).value, monthly.cell(r, 2).value)
        require(k not in actual, f"W2月汇总重复任务: {k}")
        actual[k] = monthly.cell(r, 3).value
    equal(actual, {k: v[0] for k, v in updated.items()}, "W2一月汇总")
    equal(matrix_records(case.one("expected_annual"), 2), updated, "W2全年更新")


def check_w3(case):
    case.check_manifest()
    case.selection(3, "deadline_month", "完成时限")
    equal(case.document.get("expected_changes"), [
        {"task_id": "A002", "unit": "财务部", "field": "完成数", "old": 8, "new": 15}
    ], "W3 JSON变更身份/字段/旧值/新值")
    equal(case.document.get("expected_monthly_ids"), {"3月": ["A003"]}, "W3 JSON月度集合")
    source = case.one("original")
    base, _ = native_rows(source, source, 7, lambda r: f"=F{r}/E{r}")
    equal(set(base), {f"A{i:03d}" for i in range(1, 9)}, "W3原始任务集合")
    equal(base["A002"][5], 8, "W3反馈前A002完成数")
    units = {v[2] for v in base.values()}
    seen = set()
    for path in case.files("split", len(units)):
        ws = case.open(path)
        require(ws.title in units and ws.title not in seen, f"W3拆分单位重复/未知: {ws.title}")
        seen.add(ws.title)
        actual, _ = native_rows(ws, source, 7, lambda r: f"=F{r}/E{r}")
        equal(actual, {k: v for k, v in base.items() if v[2] == ws.title}, f"W3拆分/{ws.title}")
    updated = dict(base)
    values = list(base["A002"]); values[5] = 15; updated["A002"] = tuple(values)
    feedback = case.one("feedback")
    equal(feedback.title, "财务部", "W3反馈单位")
    actual, _ = native_rows(feedback, source, 7, lambda r: f"=F{r}/E{r}")
    equal(actual, {k: updated[k] for k in ("A002", "A006")}, "W3财务反馈任务范围及变更")
    month, _ = native_rows(case.one("expected_monthly"), source, 7, lambda r: f"=F{r}/E{r}")
    equal(month, {"A003": updated["A003"]}, "W3三月任务集合")
    annual, _ = native_rows(case.one("expected_annual"), source, 7, lambda r: f"=F{r}/E{r}")
    equal(annual, updated, "W3全年更新")


def check_w6(case):
    case.check_manifest()
    equal(case.document.get("expected_changes"), [
        {"task_id": "C003", "field": field, "old": "3月", "new": "5月"}
        for field in ("起始月", "结束月")
    ], "W6 JSON变更身份/字段/旧值/新值")
    equal(case.document.get("expected_monthly_ids"),
          {"1月": ["C001"], "2月": ["C002"], "3月": [], "4月": ["C004"], "5月": ["C003"]},
          "W6 JSON反馈后月度集合")
    spec = case.document.get("selection", {})
    for key, expected in {"target_year": 2026, "target_months": [1, 2, 3, 4, 5],
                          "month_source": "起始月", "month_rule": "start_month",
                          "snapshot": "after_feedback"}.items():
        equal(spec.get(key), expected, f"W6 selection.{key}")
    source = case.one("original")
    equal(tuple(c.value for c in source[1]), W6_HEADER, "W6原始表头")
    base, _ = rows(source, 2, 6)
    equal(set(base), {"C001", "C002", "C003", "C004"}, "W6原始任务集合")
    equal(base["C003"][3:5], ("3月", "3月"), "W6调整前月份")
    seen = set()
    units = {v[2] for v in base.values()}
    for path in case.files("split", 4):
        ws = case.open(path)
        equal(tuple(c.value for c in ws[1]), W6_HEADER, f"W6/{ws.title}表头")
        require(ws.title in units and ws.title not in seen, f"W6拆分单位重复/未知: {ws.title}")
        seen.add(ws.title)
        actual, _ = rows(ws, 2, 6)
        equal(actual, {k: v for k, v in base.items() if v[2] == ws.title}, f"W6拆分/{ws.title}")
    updated = dict(base)
    vals = list(base["C003"]); vals[3:5] = ["5月", "5月"]; updated["C003"] = tuple(vals)
    for role in ("feedback", "expected_annual"):
        ws = case.one(role)
        equal(tuple(c.value for c in ws[1]), W6_HEADER, f"W6/{role}表头")
        actual, _ = rows(ws, 2, 6)
        equal(actual, updated, f"W6/{role}: 只调整C003起止月")
    monthly = case.one("expected_monthly")
    equal(tuple(c.value for c in monthly[1]), ("月份", "编号", "任务名称", "责任单位", "数量"),
          "W6月度表头")
    actual, _ = rows(monthly, 2, 5, key_col=1)
    expected = {tid: (v[3], tid, v[1], v[2], v[5]) for tid, v in updated.items()}
    equal(actual, expected, "W6反馈后月度结果：3月空，5月只有C003")


def dimension(dimension):
    return (dimension.width if hasattr(dimension, "width") else dimension.height,
            dimension.hidden, dimension.outlineLevel, dimension.collapsed)


def check_native_fidelity(case):
    """与原始文件逐属性比较，保真不依赖生成器写入的检查声明。"""
    source = case.one("original")
    source_ids = {source.cell(r, 1).value: r for r in range(5, source.max_row + 1)}
    width = source.max_column
    expected_merges = {str(m) for m in source.merged_cells.ranges}
    contract = case.document.get("fidelity_contract", {})
    for key, expected in {
        "version": 1, "source_file": case.files("original", 1)[0].name,
        "source_sheet": source.title, "preserved_header_rows": [1, 2, 3, 4],
        "data_start_row": 5, "header_merges": sorted(expected_merges),
        "freeze_panes": source.freeze_panes,
        "derived_formula_column": 6 if case.cid == "W1" else 7,
        "derived_formula": "=E{row}*100" if case.cid == "W1" else "=F{row}/E{row}",
        "new_task_format_source_id": "T001" if case.cid == "W1" else "A001",
    }.items():
        equal(contract.get(key), expected, f"{case.cid}: 保真JSON {key}")
    file_contracts = contract.get("files", [])
    require(isinstance(file_contracts, list), f"{case.cid}: 保真JSON缺少文件列表")
    by_file = {item["file"]: item for item in file_contracts}
    expected_files = {p.name for role in ROLES[1:] for p in case.files(role)}
    equal(len(by_file), len(file_contracts), f"{case.cid}: 保真JSON文件不能重复")
    equal(set(by_file), expected_files, f"{case.cid}: 保真JSON输出文件集合")
    for role in ROLES[1:]:
        for path in case.files(role):
            ws = case.open(path)
            label = f"{case.cid}/{path.name}"
            equal({str(m) for m in ws.merged_cells.ranges}, expected_merges, f"{label}: 合并区域")
            equal(ws.freeze_panes, source.freeze_panes, f"{label}: 冻结位置")
            equal({k: dimension(v) for k, v in ws.column_dimensions.items()},
                  {k: dimension(v) for k, v in source.column_dimensions.items()}, f"{label}: 全部列宽/隐藏")
            for attr in ("page_setup", "page_margins", "print_options", "HeaderFooter"):
                equal(xml(getattr(ws, attr)), xml(getattr(source, attr)), f"{label}: {attr}")
            equal(xml(ws.sheet_properties.pageSetUpPr), xml(source.sheet_properties.pageSetUpPr),
                  f"{label}: 打印缩放")
            for attr in ("print_title_rows", "print_title_cols", "print_area"):
                equal(str(getattr(ws, attr)), str(getattr(source, attr)), f"{label}: {attr}")
            for r in range(1, ws.max_row + 1):
                if r <= 4:
                    sr = r
                else:
                    tid = ws.cell(r, 1).value
                    require(tid in source_ids or (case.cid == "W1" and tid == "T999"),
                            f"{label}: 无样式来源任务 {tid!r}")
                    sr = source_ids.get(tid, 5)
                equal(dimension(ws.row_dimensions[r]), dimension(source.row_dimensions[sr]),
                      f"{label}: 行{r}高度/隐藏")
                for c in range(1, width + 1):
                    cell, original = ws.cell(r, c), source.cell(sr, c)
                    if r <= 4:
                        equal(cell.value, original.value, f"{label}: 前置区/表头 {cell.coordinate}")
                    for attr in ("font", "fill", "border", "alignment", "protection"):
                        equal(xml(getattr(cell, attr)), xml(getattr(original, attr)),
                              f"{label}!{cell.coordinate}: {attr}")
                    equal(cell.number_format, original.number_format, f"{label}!{cell.coordinate}: 数字格式")
            target_range = str(MultiCellRange(f"G5:G{ws.max_row}")) if ws.max_row >= 5 else None
            declared = by_file[path.name]
            task_ids = [ws.cell(r, 1).value for r in range(5, ws.max_row + 1)]
            for key, expected in {
                "sheet": ws.title, "header_range": f"A1:{get_column_letter(width)}4",
                "data_range": f"A5:{get_column_letter(width)}{ws.max_row}" if task_ids else None,
                "task_ids": task_ids,
                "style_source_ids": {tid: tid if tid in source_ids else "T001" for tid in task_ids},
            }.items():
                equal(declared.get(key), expected, f"{label}: 保真JSON {key}")
            actual_dvs = ws.data_validations.dataValidation
            expected_dvs = deepcopy(source.data_validations.dataValidation) if target_range else []
            for dv in expected_dvs:
                dv.sqref = target_range
            equal([xml(v) for v in actual_dvs], [xml(v) for v in expected_dvs],
                  f"{label}: 下拉类型/范围/公式及属性")
            equal([str(MultiCellRange(v)) for v in declared.get("data_validation_ranges", [])],
                  [str(dv.sqref) for dv in actual_dvs], f"{label}: 保真JSON下拉范围")
            actual_cf = [(str(cf.sqref), [(xml(rule), xml(rule.dxf))
                                        for rule in ws.conditional_formatting[cf]])
                         for cf in ws.conditional_formatting]
            expected_cf = [(target_range, [(xml(rule), xml(rule.dxf))
                                          for rule in source.conditional_formatting[cf]])
                           for cf in source.conditional_formatting] if target_range else []
            equal(actual_cf, expected_cf, f"{label}: 条件格式规则/范围/颜色")
            equal([str(MultiCellRange(v)) for v in declared.get("conditional_format_ranges", [])],
                  [str(cf.sqref) for cf in ws.conditional_formatting], f"{label}: 保真JSON条件格式范围")


def check_d40(root):
    path = Path(root) / "D40" / "D40_列换序引用.xlsx"
    require(path.is_file(), "D40缺工作簿")
    wb = load_workbook(path)
    try:
        equal(wb.sheetnames, ["原始", "换列后"], "D40工作表集合")
        old, new = wb["原始"], wb["换列后"]
        equal(tuple(c.value for c in old[1]), ("A值", "B值", "C值", "D值", "SUM(A:C)"), "D40原表头")
        equal(tuple(c.value for c in new[1]), ("A值", "D值", "C值", "B值", "原SUM内容"), "D40换列表头")
        equal(old.max_row, 6, "D40原始行数")
        equal(new.max_row, old.max_row, "D40换列行数")
        for r in range(2, 7):
            values = tuple(old.cell(r, c).value for c in range(1, 5))
            equal(values, (2, 10, 100, 1000), f"D40原始第{r}行")
            equal(tuple(new.cell(r, c).value for c in range(1, 5)),
                  (values[0], values[3], values[2], values[1]), f"D40换列第{r}行")
            equal(old.cell(r, 5).value, f"=SUM(A{r}:C{r})", f"D40原始公式{r}")
            equal(new.cell(r, 5).value, f"=A{r}+D{r}+C{r}", f"D40换列公式{r}")
            equal(values[0] + values[1] + values[2], 112, f"D40独立算术结果{r}")
    finally:
        wb.close()


def check_f23(root):
    path = Path(root) / "F23" / "F23_条件格式验证.xlsx"
    require(path.is_file(), "F23缺工作簿")
    wb = load_workbook(path)
    try:
        ws = wb.active
        equal(ws.max_row, 9, "F23任务行数")
        values = [ws.cell(r, 2).value for r in range(2, 10)]
        require(all(isinstance(v, (float, int)) and 0 <= v <= 1 for v in values),
                "F23完成率必须为0~1数值")
        dvs = ws.data_validations.dataValidation
        equal(len(dvs), 2, "F23验证规则数（不能缺失或重复）")
        by_type = {dv.type: dv for dv in dvs}
        equal(set(by_type), {"decimal", "custom"}, "F23验证类型")
        decimal, custom = by_type["decimal"], by_type["custom"]
        equal((str(decimal.sqref), decimal.operator, decimal.formula1, decimal.formula2),
              ("B2:B9", "between", "0", "1"), "F23小数验证范围/上下限")
        equal((str(custom.sqref), custom.formula1),
              ("B2:B9", "AND(B2>=0,B2<=1,ISNUMBER(B2))"), "F23自定义验证范围/公式")
        cf = [(str(item.sqref), rule) for item in ws.conditional_formatting
              for rule in ws.conditional_formatting[item]]
        equal(len(cf), 3, "F23条件格式规则数")
        signatures = {(rng, rule.type, rule.operator, tuple(rule.formula)) for rng, rule in cf}
        equal(signatures, {
            ("B2:B9", "cellIs", "greaterThanOrEqual", ("0.8",)),
            ("B2:B9", "cellIs", "lessThan", ("0.5",)),
            ("D2:D9", "expression", None, ("DATE(2026,10,1)-D2>0",)),
        }, "F23条件格式范围/条件")
        for _, rule in cf:
            require(rule.dxf is not None, "F23条件格式缺样式")
            if rule.type == "cellIs":
                expected = "C6EFCE" if rule.operator == "greaterThanOrEqual" else "FFC7CE"
                equal(str(rule.dxf.fill.fgColor.rgb)[-6:], expected, "F23填充色")
            else:
                equal(str(rule.dxf.font.color.rgb)[-6:], "FF0000", "F23过期字体颜色")
    finally:
        wb.close()


def run(root):
    """每个检查组仅输出一条结果；任何缺任务/错值/错公式均为 fail。"""
    results = []
    for cid, check in (("W1", check_w1), ("W2", check_w2), ("W3", check_w3), ("W6", check_w6)):
        case = None
        try:
            case = Case(root, cid)
            check(case)
            results.append(("pass", f"{cid}: 基线→拆分→反馈→月度/全年预期内容一致", "业务数据"))
        except Exception as exc:
            results.append(("fail", f"{cid}: {exc}", "业务数据"))
        finally:
            if case:
                case.close()
    for cid in ("W1", "W3"):
        case = None
        try:
            case = Case(root, cid)
            check_native_fidelity(case)
            results.append(("pass", f"{cid}: 已核对全部输出的标题/表头、合并、样式、尺寸、冻结、打印及DV/CF",
                            "原生属性"))
        except Exception as exc:
            results.append(("fail", f"{cid}保真: {exc}", "原生属性"))
        finally:
            if case:
                case.close()
    for cid, check in (("D40", check_d40), ("F23", check_f23)):
        try:
            check(root)
            results.append(("pass", f"{cid}: 公式/验证规则与独立预期一致", "公式与规则"))
        except Exception as exc:
            results.append(("fail", f"{cid}: {exc}", "公式与规则"))
    return results
