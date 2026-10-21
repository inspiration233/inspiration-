# -*- coding: utf-8 -*-
"""检查验收夹具自身的一致性；不代表产品引擎已通过验收。"""
import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from 验收断言 import run as check_contracts

ROOT = Path(__file__).resolve().parent
CASE_IDS = (
    [f"W{i}" for i in range(1, 7)]
    + [f"S{i:02d}" for i in range(7, 19)]
    + [f"F{i}" for i in range(19, 25)]
    + [f"A{i}" for i in range(25, 29)]
    + [f"B{i}" for i in range(29, 32)]
    + [f"C{i}" for i in range(32, 39)]
    + [f"D{i}" for i in range(39, 42)]
    + ["RT42"]
)
CONTRACT_CASES = {"W1", "W2", "W3", "W6", "D40", "F23"}


def inspect(root=ROOT):
    """返回分类结果，供CLI及反例验证使用。每次重新打开当前文件。"""
    results = []

    def record(kind, message, category):
        results.append((kind, message, category))

    actual_cases = {p.parent.name for p in root.glob("*/case.json")}
    expected_cases = set(CASE_IDS)
    record(
        "pass" if actual_cases == expected_cases else "fail",
        f"案例清单：{len(actual_cases)}/42；缺少={sorted(expected_cases-actual_cases)}；"
        f"多出={sorted(actual_cases-expected_cases)}",
        "文件结构",
    )
    for cid in sorted(actual_cases | expected_cases):
        d = root / cid
        try:
            cj = json.loads((d / "case.json").read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            record("fail", f"{cid}: case.json 无法读取：{exc}", "文件结构")
            continue
        if not isinstance(cj, dict) or cj.get("case_id") != cid:
            record("fail", f"{cid}: case.json 的 case_id 与目录不符", "文件结构")
            continue
        files = sorted(d.glob("*.xlsx"))
        if not files:
            record("fail", f"{cid}: 缺少 Excel 文件", "文件结构")
        for path in files:
            try:
                wb = load_workbook(path)
                wb.close()
                record("pass", f"{cid}/{path.name}: 可读取", "文件结构")
            except Exception as exc:
                record("fail", f"{cid}/{path.name}: 无法读取：{exc}", "文件结构")

        artifacts = cj.get("artifacts")
        if artifacts:
            schema_valid = (isinstance(artifacts, dict)
                            and all(isinstance(group, list) for group in artifacts.values())
                            and all(isinstance(n, str) for group in artifacts.values() for n in group))
            names = [name for group in artifacts.values() for name in group] if schema_valid else []
            valid = schema_valid and all(Path(n).name == n for n in names)
            actual_names = {p.name for p in files}
            expected_names = set(names)
            ok = valid and len(names) == len(expected_names) and actual_names == expected_names
            record(
                "pass" if ok else "fail",
                f"{cid}: artifacts文件清单；缺少={sorted(expected_names-actual_names)}；"
                f"多出={sorted(actual_names-expected_names)}",
                "文件结构",
            )
        elif cj.get("type") == "完整流程":
            record("fail", f"{cid}: 完整流程必须指定 artifacts 文件清单", "文件结构")

        if cid not in CONTRACT_CASES and cid not in {"W4", "W5", "D41"}:
            record("skip", f"{cid}: 尚未实现独立内容断言；文件可读不等于业务通过", "检查覆盖")
        for question in cj.get("ambiguity", []):
            record("warn", f"{cid}: {question}", "业务待确认")
        if cj.get("checks") and cid not in CONTRACT_CASES:
            record("skip", f"{cid}: checks 中的属性描述尚未转为自动断言", "检查覆盖")

        try:
            if cid == "W5":
                wb = load_workbook(d / "W5_原始_Sheet月份.xlsx")
                counts = [sum(row[0] is not None for row in ws.iter_rows(min_row=2, values_only=True))
                          for ws in wb.worksheets]
                truth = cj.get("ground_truth", {})
                ok = (wb.sheetnames == ["1月", "2月", "3月"]
                      and counts == [4, 5, 6]
                      and truth.get("tasks_per_sheet") == counts
                      and truth.get("total_tasks") == sum(counts))
                wb.close()
                record("pass" if ok else "fail", f"W5: 三张表的任务数与JSON一致：{counts}", "业务数据")
            elif cid == "W4":
                wb = load_workbook(d / "W4_原始_纵向合并.xlsx")
                ws = wb.active
                truth = cj.get("ground_truth", {})
                ok = ws.max_row == 14 and truth.get("total_subtasks") == 13
                wb.close()
                record("pass" if ok else "fail", "W4: 原始表13行子任务与JSON一致", "业务数据")
                record("skip", "W4: 项目拆分业务口径及反馈链路尚未完整验收", "检查覆盖")
            elif cid == "D41":
                ok = "保留空行" in cj.get("expected", "") and not cj.get("ambiguity")
                record("pass" if ok else "fail", "D41: 默认保留排版空行，已移除重复业务疑问", "业务数据")
                record("skip", "D41: 产品拆分前后空行位置对比待实现", "检查覆盖")
        except Exception as exc:
            record("fail", f"{cid}: 内容检查无法执行：{exc}", "业务数据")

    try:
        results.extend(check_contracts(root))
    except Exception as exc:
        record("fail", f"独立断言执行异常：{type(exc).__name__}: {exc}", "检查器")
    record("warn", "尚未在 Excel/WPS 验证实际重算、下拉交互、条件格式显示和打印效果", "实机验证")
    return results


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT, help="待检查的测试集目录")
    parser.add_argument("--quiet", action="store_true", help="仅显示失败和分类汇总")
    parser.add_argument("--json", type=Path, help="保存机器可读的结果")
    args = parser.parse_args()
    results = inspect(args.root.resolve())
    totals = Counter(kind for kind, _, _ in results)
    categories = defaultdict(Counter)
    for kind, message, category in results:
        categories[category][kind] += 1
        if not args.quiet or kind == "fail":
            print(f"[{kind}][{category}] {message}")
    print("\n验收夹具自检（不是产品引擎验收）")
    for category, count in sorted(categories.items()):
        print(f"{category}: 通过 {count['pass']} / 失败 {count['fail']} / "
              f"跳过 {count['skip']} / 提醒 {count['warn']}")
    print(f"合计: 通过 {totals['pass']} / 失败 {totals['fail']} / "
          f"跳过 {totals['skip']} / 提醒 {totals['warn']}")
    if args.json:
        args.json.write_text(json.dumps({
            "scope": "fixture_consistency",
            "totals": dict(totals),
            "categories": {k: dict(v) for k, v in categories.items()},
            "results": [{"status": k, "message": m, "category": c} for k, m, c in results],
        }, ensure_ascii=False, indent=2), encoding="utf-8")
    return 1 if totals["fail"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
