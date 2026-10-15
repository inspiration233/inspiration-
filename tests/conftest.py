# -*- coding: utf-8 -*-
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
FIX = ROOT / "测试集"
sys.path.insert(0, str(ROOT))

from engine.tools.excel_merge.merge import Options  # noqa: E402


def scenario(name):
    d = FIX / name
    assert d.exists(), f"测试集目录不存在: {d}（可运行 测试集/生成测试集.py --skip-large 重建）"
    return d


@pytest.fixture
def out_path(tmp_path):
    return tmp_path / "合并结果.xlsx"


def load(out):
    from openpyxl import load_workbook
    return load_workbook(out)


def rows_of(ws, min_row=2):
    return [r for r in ws.iter_rows(min_row=min_row, values_only=True)]
