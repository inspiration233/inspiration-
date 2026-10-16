# -*- coding: utf-8 -*-
"""V2.x 测试集对齐：场景H 保真全家桶 / 场景I 同模板多单位填报 / 场景J 方向与序列。"""
from openpyxl import load_workbook

from engine.tools.excel_merge.merge import Options, merge, scan

from conftest import load, rows_of, scenario


# ================================================================ 场景H

def test_h_fidelity_full_house(out_path):
    d = scenario("场景H_保真全家桶")
    r = merge([str(d / "保真全家桶.xlsx")], out_path, Options(mode="single"))
    ws = load(out_path)["总表"]      # 单文件组页名=总表
    t = r["totals"]

    # 冻结窗格照原表 / 打印设置(横向A4)
    assert ws.freeze_panes == "A2"
    assert ws.page_setup.orientation == "landscape"
    assert int(ws.page_setup.paperSize) == 9
    # 自定义行高
    assert ws.row_dimensions[2].height == 24
    # 批注
    assert ws["D2"].comment and "返点" in ws["D2"].comment.text
    # 智能表格保留(输出名加序号后缀防跨页重名)
    assert any(n.startswith("销售表") for n in ws.tables)
    tbl = ws.tables[[n for n in ws.tables if n.startswith("销售表")][0]]
    assert tbl.ref == "A1:G13"
    # 块内公式恒等平移 + 跨表公式转值计数
    assert str(ws["E2"].value).startswith("=C2/D2")
    assert t["formulas_to_value"] >= 1                     # =横向表!B2
    assert ws["G2"].value in (None, "")                    # 跨表公式无缓存 → 空
    # 下拉(内联)保留
    dvs = ws.data_validations.dataValidation
    assert any(dv.formula1 == '"达标,未达"' for dv in dvs)
    # 条件格式保留(CellIs + 公式型)
    cf_ranges = [str(rng) for rng in ws.conditional_formatting]
    assert any("E2" in s for s in cf_ranges)
    assert any("F2" in s for s in cf_ranges)


def test_h_mixed_orientation_one_file(out_path):
    """一个文件里 Sheet 纵向 + Sheet 横向，各自转正互不影响。"""
    d = scenario("场景J_方向与序列")
    r = scan([str(d / "混合方向.xlsx")])
    by_sheet = {x["sheet"]: x["orientation"] for x in r["header_detections"]}
    assert by_sheet == {"纵向页": "vertical", "横向页": "horizontal"}

    out = out_path.parent / "mixed.xlsx"
    merge([str(d / "混合方向.xlsx")], out, Options(mode="workbook"))
    wb = load(out)
    v = wb["纵向页"]
    assert [c.value for c in v[1]] == ["客户", "金额", "日期"]
    h = wb["横向页"]                                        # 横向转正: 字段成表头
    assert [c.value for c in h[1]] == ["客户", "金额", "日期"]
    assert h["A2"].value == "上海岚图"                      # 记录1(原B列)转正为行
    assert h["B2"].value == 12000
    out.unlink(missing_ok=True)


# ================================================================ 场景I

def test_i_multi_unit_same_template(out_path):
    d = scenario("场景I_同模板多单位填报")
    r = merge([str(d / f) for f in ["单位甲_季度报表.xlsx", "单位乙_季度报表.xlsx",
                                    "单位丙_季度报表.xlsx"]],
              out_path, Options(mode="single"))
    assert r["groups"][0]["shape"] == "stack"
    t = r["totals"]
    wb = load(out_path)
    assert wb.sheetnames == ["总表"]
    ws = wb["总表"]

    # 共性前置区保留一份(标题/说明)，差异行(制表人行)进元数据列
    assert ws["A1"].value == "2026年第三季度经营情况表"
    assert str(ws["A2"].value).startswith("注：金额单位")
    header = [c.value for c in ws[3]]
    assert header[:5] == ["序号", "项目", "本期金额", "上年同期", "增长率"]
    for label in ("填报单位", "制表人", "填报日期"):
        assert label in header, header

    # 数据堆叠 + 合计剔除 + 序号跨文件连续
    rows = rows_of(ws, min_row=4)
    assert len(rows) == 6 + 7 + 6                          # 甲6 + 乙7 + 丙6(真合计行剔除)
    assert t["summary_rows_removed"] == 1
    assert [row[0] for row in rows] == list(range(1, len(rows) + 1))

    # 元数据列各源填各自值
    unit_col = header.index("填报单位")
    units = [row[unit_col] for row in rows]
    assert set(units) == {"甲分公司", "乙分公司", "丙分公司"}

    # 增长率公式跨文件平移(乙的首行行号 = 甲之后)
    growth_col = header.index("增长率")
    f1, f2 = str(rows[0][growth_col]), str(rows[6][growth_col])
    assert f1.startswith("=C4/D4") and f2.startswith("=C10/D10")   # 甲r4-r9, 乙接续r10

    # 表头只出现一次
    assert not any(row[1] == "项目" for row in rows)


# ================================================================ 场景J

def test_j_orientation_regressions():
    d = scenario("场景J_方向与序列")
    cases = {
        "序列月份_纵向.xlsx": ("vertical", None),   # 月份数据列不误判(值任意, 方向必须纵向)
        "流水号不误判.xlsx": ("vertical", True),
        "交叉表_低置信.xlsx": ("vertical", False),  # 同质数据 → 低置信交人工
    }
    for f, (orient, conf) in cases.items():
        det = scan([str(d / f)])["header_detections"][0]
        assert det["orientation"] == orient, f
        if conf is not None:
            assert det["orientation_confident"] is conf, f


def test_j_orientation_override():
    """低置信表手动指定方向 → 覆盖生效。"""
    d = scenario("场景J_方向与序列")
    p = str(d / "交叉表_低置信.xlsx")
    r = scan([p], Options(orientation_overrides={"交叉表_低置信/模糊": "horizontal"}))
    det = r["header_detections"][0]
    assert det["orientation"] == "horizontal"
    assert det["orientation_confident"] is True
