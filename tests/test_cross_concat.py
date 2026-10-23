# -*- coding: utf-8 -*-
"""原生交叉表(记录沿列)分组横向拼接: 布局保持 + 右侧追加 + 示例列去重。"""
import sys
from pathlib import Path

import pytest

from conftest import load

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from engine.tools.excel_merge.merge import Options, merge, scan  # noqa: E402


# --------------------------------------------------------------- 夹具构造

LABELS = ["长度", "宽度", "重量", "高度", "颜色", "材质", "功率", "电压",
          "电流", "转速", "扭矩", "油耗", "排放", "噪音", "轮胎", "座椅"]


def _cross_book(path, records, preamble_variant=1, extra_labels=(), reorder=False,
                stray=True):
    """构造 前置区+例示列+记录列 的交叉表; records=[(名称, [16个数值])]。

    贴近真实配置表密度: 16个字段标签、记录列全数值(个别None造脏)。
    """
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    labels = (LABELS[::-1] if reorder else LABELS) + list(extra_labels)
    wb = Workbook()
    ws = wb.active
    ws.title = "配置表"
    ws["A1"] = "技术配置表"
    ws.merge_cells("A1:G1")
    ws["A2"] = f"填报说明：版本{preamble_variant}"
    ws.merge_cells("A2:G2")
    ws["A3"] = "单位（盖章）："
    ws["D3"] = "联系人："
    ws["F3"] = "填报时间："
    ws.merge_cells("F3:G3")
    ws["A4"] = "项目"
    ws.cell(4, 2, "例示")
    for r, lbl in enumerate(labels, start=5):
        ws.cell(r, 1, lbl)                   # 字段标签列(首列)
        ws.cell(r, 2, r)                     # 例示列: 数值样例
    col = 3
    for name, values in records:
        ws.cell(4, col, name)
        for r, v in enumerate(values, start=5):
            if v is not None:
                ws.cell(r, col, v)
        col += 1
    ws.cell(4, col, "预留列")                # 只有表头的尾部列
    if stray:
        ws.cell(7, 15, "杂散")               # 标签区右侧孤立格, 不应进入拼接
    for c in range(1, col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 12.5 if c % 2 else 8
    wb.save(path)
    wb.close()


@pytest.fixture
def two_cross_files(tmp_path):
    f1 = tmp_path / "甲_配置.xlsx"
    f2 = tmp_path / "乙_配置.xlsx"
    vals1 = [list(range(100, 116)), [None] + list(range(200, 215))]
    vals2 = [list(range(300, 316)), list(range(400, 415)) + [None]]
    _cross_book(f1, [("方案一", vals1[0]), ("方案二", vals1[1])], preamble_variant=1)
    _cross_book(f2, [("方案三", vals2[0]), ("方案四", vals2[1])], preamble_variant=2)
    return f1, f2


# --------------------------------------------------------------- 用例

def test_cross_concat_appends_columns(two_cross_files, out_path):
    f1, f2 = two_cross_files
    r = merge([str(f1), str(f2)], out_path, Options(mode="group"))
    ws = load(out_path)["配置表"]
    # 前置区/表头原样: 例示列仅一份, 甲的 方案一/方案二 + 预留列, 乙的3列右拼
    assert ws["A1"].value == "技术配置表"
    assert [ws.cell(4, c).value for c in range(1, 9)] == [
        "项目", "例示", "方案一", "方案二", "预留列", "方案三", "方案四", "预留列"]
    assert ws.cell(4, 9).value is None       # 乙的例示列未重复拼入
    # 标签列 + 按标签对齐的数据
    assert ws["A5"].value == "长度" and ws["A20"].value == "座椅"
    assert ws.cell(5, 3).value == 100        # 甲 方案一
    assert ws.cell(5, 6).value == 300        # 乙 方案三
    assert ws.cell(5, 7).value == 400        # 乙 方案四
    assert ws.cell(5, 4).value is None       # 甲 方案二 的脏 None(长度行)
    assert ws.cell(6, 4).value == 200
    assert ws.cell(5, 8).value is None       # 预留列只有表头
    # 合并保留
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert {"A1:G1", "A2:G2", "F3:G3"} <= merged
    # 杂散格与标签区外内容未带入
    assert ws.cell(7, 15).value is None
    # 报告: 交叉形态; 前置区不同 → 采用首份并提示
    grp = next(g for g in r["groups"] if g["sheet"] == "配置表")
    assert grp["shape"] == "cross" and grp["count"] == 7
    assert any("前置区" in w for w in r["warnings"])


def test_cross_concat_skips_identical_example_column(tmp_path, out_path):
    """两份完全相同的交叉表: 普通记录列照拼(允许重复), 例示列只留一份。"""
    f1 = tmp_path / "a.xlsx"
    f2 = tmp_path / "b.xlsx"
    _cross_book(f1, [("方案一", list(range(1, 17)))])
    _cross_book(f2, [("方案一", list(range(1, 17)))])
    r = merge([str(f1), str(f2)], out_path, Options(mode="group"))
    ws = load(out_path)["配置表"]
    assert [ws.cell(4, c).value for c in range(1, 7)] == [
        "项目", "例示", "方案一", "预留列", "方案一", "预留列"]   # 无第二个 例示
    assert not any("内容完全相同" in w for w in r["warnings"])
    assert ws.cell(5, 5).value == 1          # 第二份记录列数据照拼


def test_cross_concat_label_alignment_and_missing_warning(tmp_path, out_path):
    """字段标签顺序不同按标签对齐; 首表没有的标签提示未拼接。"""
    f1 = tmp_path / "a.xlsx"
    f2 = tmp_path / "b.xlsx"
    _cross_book(f1, [("方案一", list(range(10, 26)))])
    _cross_book(f2, [("方案二", list(range(20, 36)) + [99])],
                extra_labels=("功率2",), reorder=True)
    r = merge([str(f1), str(f2)], out_path, Options(mode="group"))
    ws = load(out_path)["配置表"]
    assert ws["A5"].value == "长度" and ws.cell(6, 1).value == "宽度"
    assert ws.cell(5, 3).value == 10         # 甲
    assert ws.cell(4, 5).value == "例示"     # 乙例示列取值随逆序变化 → 不算同一示例列, 照拼
    assert ws.cell(4, 6).value == "方案二"
    assert ws.cell(5, 6).value == 35         # 乙标签逆序, 仍按标签对齐到长度行
    assert ws.cell(6, 6).value == 34
    assert not any(ws.cell(r, 6).value == 99 for r in range(5, 22))  # 功率2未拼入
    assert any("功率2" in w for w in r["warnings"])


def test_cross_concat_thin_group_follows_detected_member(tmp_path, out_path):
    """同组个别成员记录太少(方向信号不足)时, 跟随已识别交叉表的成员横向拼接。"""
    from openpyxl import Workbook
    f1 = tmp_path / "a.xlsx"                 # 密表: 可自动识别
    _cross_book(f1, [("方案一", list(range(1, 17)))])
    f2 = tmp_path / "b.xlsx"                 # 薄表: 仅2个标签, 自动识别必失败
    wb = Workbook()
    ws = wb.active
    ws.title = "配置表"
    ws["A1"] = "技术配置表"
    ws["A4"] = "项目"
    ws.cell(4, 2, "例示")
    for r, lbl in enumerate(["长度", "宽度"], start=5):
        ws.cell(r, 1, lbl)
        ws.cell(r, 2, r)
    ws.cell(4, 3, "方案薄")
    ws.cell(5, 3, 5)
    ws.cell(6, 3, 6)
    wb.save(f2)
    wb.close()
    merge([str(f1), str(f2)], out_path, Options(mode="group"))
    ws = load(out_path)["配置表"]
    assert ws.cell(4, 5).value == "例示"     # 薄表例示只覆盖2个字段, 与首份不同 → 照拼
    assert ws.cell(4, 6).value == "方案薄"   # 薄表记录列右拼
    assert ws.cell(5, 6).value == 5
    assert ws.cell(4, 7).value is None


def test_cross_concat_manual_override_for_text_records(tmp_path, out_path):
    """记录列为文本(自动信号不足)时, 体检手动切横向同样走右拼。"""
    from openpyxl import Workbook

    def text_book(path, name, vals):
        wb = Workbook()
        ws = wb.active
        ws.title = "配置表"
        ws["A1"] = "验收要求"
        ws["A3"] = "单位："
        ws["A4"] = "项目"
        ws.cell(4, 2, "例示")
        for r, (lbl, sample) in enumerate(
                [("发动机", "涡轮增压"), ("变速箱", "自动"), ("驱动", "四驱"),
                 ("座椅", "真皮"), ("空调", "自动")], start=5):
            ws.cell(r, 1, lbl)
            ws.cell(r, 2, sample)
        ws.cell(4, 3, name)
        for r, v in enumerate(vals, start=5):
            ws.cell(r, 3, v)
        wb.save(path)
        wb.close()

    f1, f2 = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    text_book(f1, "豪华型", ["2.0T", "8AT", "四驱", "真皮", "自动"])
    text_book(f2, "舒适型", ["1.5T", "CVT", "前驱", "织物", "手动"])
    merge([str(f1), str(f2)], out_path,
          Options(mode="group", orientation_overrides={"a": "horizontal", "b": "horizontal"}))
    ws = load(out_path)["配置表"]
    assert ws.cell(4, 3).value == "豪华型" and ws.cell(4, 4).value == "舒适型"
    assert ws.cell(5, 3).value == "2.0T" and ws.cell(5, 4).value == "1.5T"
    assert ws.cell(4, 5).value is None       # 例示列未重复


def test_cross_concat_not_deduped_as_identical_sheet(two_cross_files, out_path):
    """完全相同的交叉表不再被'同表去重'整表丢弃(记录列需要右拼)。"""
    f1, _ = two_cross_files
    f2 = f1.with_name("copy_" + f1.name)
    f2.write_bytes(f1.read_bytes())
    r = merge([str(f1), str(f2)], out_path, Options(mode="group"))
    assert not any("内容完全相同" in w for w in r["warnings"])
    ws = load(out_path)["配置表"]
    assert ws.cell(4, 6).value == "方案一" and ws.cell(4, 7).value == "方案二"


def test_titled_vertical_table_not_flipped(tmp_path, out_path):
    """带标题行的'姓名+数值列'纵向表不受例示信号影响, 保持纵向。"""
    from openpyxl import Workbook
    f1 = tmp_path / "人员.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "名单"
    ws["A1"] = "人员名单"
    ws.append(["姓名", "年龄", "工资", "奖金"])
    for name, age, pay, bonus in [("张三", 25, 3000, 500), ("李四", 30, 4000, 800),
                                  ("王五", 28, 3500, 600), ("赵六", 35, 5000, 1000)]:
        ws.append([name, age, pay, bonus])
    wb.save(f1)
    wb.close()
    det = scan([str(f1)])["header_detections"][0]
    assert det["orientation"] == "vertical" and not det["native_cross"]
    merge([str(f1)], out_path, Options(mode="single"))
    ws = load(out_path).active
    assert [c.value for c in ws[2]] == ["姓名", "年龄", "工资", "奖金"]
    assert ws.max_row == 6                   # 标题+表头+4行, 未转置


def test_cross_and_vertical_sheets_merge_independently(tmp_path, out_path):
    """同一对文件: 交叉表横向拼接, 常规表纵向堆叠, 互不干扰。"""
    from openpyxl import load_workbook
    f1, f2 = tmp_path / "甲.xlsx", tmp_path / "乙.xlsx"
    _cross_book(f1, [("方案一", list(range(1, 17)))])
    _cross_book(f2, [("方案二", list(range(2, 18)))])
    for p, name in ((f1, "例1"), (f2, "例2")):
        wb = load_workbook(p)
        ws = wb.create_sheet("明细")
        ws.append(["序号", "数量"])
        ws.append([name, 3])
        ws.append([2, 4])
        wb.save(p)
        wb.close()
    merge([str(f1), str(f2)], out_path, Options(mode="group"))
    cross = load(out_path)["配置表"]
    assert cross.cell(4, 3).value == "方案一" and cross.cell(4, 5).value == "方案二"
    detail = load(out_path)["明细"]
    vals = [r0[0] for r0 in detail.iter_rows(min_row=2, max_col=1, values_only=True)]
    # 文本序号(例N)保留且只留首份来源一组; 数字序号连续重排
    assert vals == ["例1", 1, 2]


def test_pure_transpose_stays_transposed(tmp_path, out_path):
    """纯转置表(字段自r1沿A列)仍走转置路径, 不进横向拼接。"""
    from openpyxl import Workbook

    def book(path, companies):
        wb = Workbook()
        ws = wb.active
        ws.title = "台账"
        data = [["客户"] + companies,
                ["联系人", "张三", "李四", "王五"],
                ["金额", 1, 2, 3],
                ["状态", "结清", "未结", "结清"],
                ["备注", None, "加急", None]]
        for row in data:
            ws.append(row)
        wb.save(path)
        wb.close()

    f1, f2 = tmp_path / "t1.xlsx", tmp_path / "t2.xlsx"
    book(f1, ["甲公司", "乙公司", "丙公司"])
    book(f2, ["丁公司", "戊公司", "己公司"])
    merge([str(f1), str(f2)], out_path, Options(mode="single"))
    ws = load(out_path).active
    assert [c.value for c in ws[1]] == ["客户", "联系人", "金额", "状态", "备注"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 6                     # 2份 × 3条记录(转置堆叠)
