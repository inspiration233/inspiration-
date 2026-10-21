"""No silent corruption from selected regions, rule anchors, or ambiguous fields."""
import pytest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from engine.tools.excel_merge.align import pair_columns
from engine.tools.excel_merge.workbench import merge_preview, merge_export


def source(tmp_path, name, headers, rows):
    wb = Workbook(); ws = wb.active
    ws.title = "明细"; ws.append(headers)
    for row in rows:
        ws.append(row)
    path = tmp_path / f"{name}.xlsx"
    wb.save(path)
    spec = {"path": str(path), "sheet": "明细", "selection": {
        "axis": "rows", "header_start": 1, "header_end": 1, "data_start": 2,
        "data_end": ws.max_row, "field_start": 1, "field_end": len(headers)}}
    return wb, ws, spec


def payload(*specs, **options):
    return {"paths": [s["path"] for s in specs], "selections": list(specs),
            "options": {"auto_group": False, "merge_shape": "stack", **options}}


def exported(tmp_path, data):
    out = tmp_path / "结果.xlsx"
    report = merge_export({**data, "out": str(out)})
    assert report["can_export"], report["issues"]
    return load_workbook(out).active


def test_rule_anchors_follow_deleted_middle_record_and_dxf_survives(tmp_path):
    wb, ws, spec = source(tmp_path, "a", ["单位", "数量"], [["甲", 10], ["合计", 15], ["乙", 20]])
    dv = DataValidation(type="custom", formula1="B2>=0", errorStyle="warning", imeMode="disabled")
    dv.add("B2:B4"); ws.add_data_validation(dv)
    ws.conditional_formatting.add("B2:B4", FormulaRule(formula=["$B2>0"], font=Font(color="FFFF0000")))
    ws.conditional_formatting.add("B2:B4", ColorScaleRule(
        start_type="min", start_color="FFFF0000", end_type="max", end_color="FF00FF00"))
    wb.save(spec["path"])
    actual = exported(tmp_path, payload(spec, remove_summary=True))
    assert actual["A3"].value == "乙"
    rules = [(str(a.sqref), r) for a, rs in actual.conditional_formatting._cf_rules.items() for r in rs]
    assert {(area, tuple(rule.formula)) for area, rule in rules if rule.type == "expression"} == {
        ("B2", ("$B2>0",)), ("B3", ("$B3>0",))}
    assert all(rule.dxf.font.color.rgb == "FFFF0000" for _, rule in rules if rule.type == "expression")
    assert sum(rule.type == "colorScale" for _, rule in rules) == 1
    validations = actual.data_validations.dataValidation
    assert {(str(d.sqref), d.formula1) for d in validations} == {("B2", "B2>=0"), ("B3", "B3>=0")}
    assert all(d.errorStyle == "warning" and d.imeMode == "disabled" for d in validations)


def test_rule_origin_before_selected_data_is_rebased_not_dropped(tmp_path):
    wb, ws, spec = source(tmp_path, "a", ["单位", "数量"], [["排除", 1], ["甲", 10], ["乙", 20]])
    dv = DataValidation(type="custom", formula1="B2>=0"); dv.add("B2:B4"); ws.add_data_validation(dv)
    ws.conditional_formatting.add("B2:B4", FormulaRule(formula=["B2>0"]))
    wb.save(spec["path"]); spec["selection"]["data_start"] = 3
    actual = exported(tmp_path, payload(spec))
    assert [(str(d.sqref), d.formula1) for d in actual.data_validations.dataValidation] == [("B2:B3", "B2>=0")]
    assert [(str(a.sqref), r.formula) for a, rs in actual.conditional_formatting._cf_rules.items()
            for r in rs] == [("B2:B3", ["B2>0"])]


def test_rule_anchors_follow_nonmonotone_column_correspondence(tmp_path):
    _, _, first = source(tmp_path, "a", ["数量", "备注", "单位"], [[1, "x", "模板"]])
    wb, ws, second = source(tmp_path, "b", ["单位", "数量", "备注"], [["乙", 20, "y"]])
    dv = DataValidation(type="custom", formula1='A2<>""'); dv.add("A2:C2"); ws.add_data_validation(dv)
    ws.conditional_formatting.add("A2:C2", FormulaRule(formula=['A2<>""']))
    wb.save(second["path"])
    actual = exported(tmp_path, payload(first, second))
    assert [actual.cell(3, i).value for i in (1, 2, 3)] == [20, "y", "乙"]
    assert {(str(d.sqref), d.formula1) for d in actual.data_validations.dataValidation} == {
        ("C3", 'C3<>""'), ("A3:B3", 'A3<>""')}
    assert {(str(a.sqref), tuple(r.formula)) for a, rs in actual.conditional_formatting._cf_rules.items()
            for r in rs} == {("C3", ('C3<>""',)), ("A3:B3", ('A3<>""',))}


def test_cut_data_merge_is_a_blocking_issue(tmp_path):
    wb, ws, spec = source(tmp_path, "a", ["单位", "任务"], [["甲", "一"], [None, "二"], [None, "三"]])
    ws.merge_cells("A2:A4"); wb.save(spec["path"])
    spec["selection"]["data_start"] = 3
    result = merge_preview(payload(spec))
    assert not result["can_export"]
    assert result["totals"]["data_merges_dropped"] == 1
    assert any("合并区" in i["message"] or "合并单元格" in i["message"] for i in result["issues"])


def test_duplicate_header_positions_can_be_explicitly_swapped_without_losing_values(tmp_path):
    _, _, first = source(tmp_path, "a", ["编号", "数量", "数量"], [["模板", 1, 2]])
    _, _, second = source(tmp_path, "b", ["编号", "数量", "数量"], [["乙", 20, 30]])
    second["field_map"] = {"2": 3, "3": 2}
    actual = exported(tmp_path, payload(first, second))
    assert [actual.cell(2, i).value for i in (1, 2, 3)] == ["模板", 1, 2]
    assert [actual.cell(3, i).value for i in (1, 2, 3)] == ["乙", 30, 20]


@pytest.mark.skip(reason="重名字段已放宽: 同数量重名按位置自动对应,不再硬阻断")
def test_unrelated_or_partial_map_does_not_bypass_duplicate_ambiguity(tmp_path):
    _, _, first = source(tmp_path, "a", ["编号", "数量", "数量"], [["模板", 1, 2]])
    _, _, second = source(tmp_path, "b", ["编号", "数量", "数量"], [["乙", 20, 30]])
    second["field_map"] = {"1": 1}
    result = merge_preview(payload(first, second))
    assert not result["can_export"]
    assert any("重名" in i["message"] for i in result["issues"])


def test_null_map_leaves_target_empty_and_keeps_unmatched_source_separate(tmp_path):
    _, _, first = source(tmp_path, "a", ["编号", "数量"], [["模板", 1]])
    _, _, second = source(tmp_path, "b", ["编号", "数量"], [["乙", 20]])
    second["field_map"] = {"2": None}
    actual = exported(tmp_path, payload(first, second))
    assert [actual.cell(3, i).value for i in (1, 2, 3)] == ["乙", None, 20]


def test_composite_headers_with_different_parents_are_never_leaf_matched():
    assert pair_columns(["计划/数量"], ["实际/数量"]) == {}
    assert pair_columns(["计划/数量"], ["数量"]) == {0: 0}


def test_field_map_cannot_silently_retarget_a_different_group_template(tmp_path):
    _, _, first = source(tmp_path, "a", ["编号", "数量"], [["模板", 1]])
    wb, ws, second = source(tmp_path, "b", ["数量", "编号"], [[20, "乙"]])
    ws.title = "另一组"; wb.save(second["path"]); second["sheet"] = "另一组"
    second["field_map"] = {"1": 2, "2": 1}
    result = merge_preview(payload(first, second, mode="group"))
    assert not result["can_export"]
    assert any("模板不在当前分组" in i["message"] for i in result["issues"])


def test_append_keeps_each_blocks_header_and_data_column_order(tmp_path):
    _, _, first = source(tmp_path, "a", ["编号", "数量"], [["模板", 1]])
    _, _, second = source(tmp_path, "b", ["数量", "编号"], [[20, "乙"]])
    actual = exported(tmp_path, payload(first, second, merge_shape="append"))
    rows = list(actual.values)
    first_header = next(i for i, row in enumerate(rows) if row[:2] == ("编号", "数量"))
    second_header = next(i for i, row in enumerate(rows) if row[:2] == ("数量", "编号"))
    assert rows[first_header + 1][:2] == ("模板", 1)
    assert rows[second_header + 1][:2] == (20, "乙")


def test_horizontal_dependency_transpose_blocks_but_plain_values_remain_available(tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = "明细"
    ws.append(["单位", "甲", "乙"])
    ws.append(["数量", 2, 3])
    path = tmp_path / "横表.xlsx"; wb.save(path)
    spec = {"path": str(path), "sheet": "明细", "selection": {
        "axis": "columns", "header_start": 1, "header_end": 1, "data_start": 2,
        "data_end": 3, "field_start": 1, "field_end": 2}}
    plain = merge_preview(payload(spec))
    assert plain["can_export"], plain["issues"]
    for kind in ("formula", "validation", "conditional_formatting"):
        modified = load_workbook(path); sheet = modified.active
        if kind == "formula":
            sheet["C2"] = "=B2+1"
        elif kind == "validation":
            dv = DataValidation(type="whole", operator="greaterThan", formula1="0")
            dv.add("B2:C2"); sheet.add_data_validation(dv)
        else:
            sheet.conditional_formatting.add("B2:C2", FormulaRule(formula=["B2>0"]))
        modified.save(path)
        result = merge_preview(payload(spec))
        assert not result["can_export"], kind
        assert any("不能可靠转置" in i["message"] for i in result["issues"])
        wb.save(path)
