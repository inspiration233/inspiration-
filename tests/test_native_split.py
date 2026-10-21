# -*- coding: utf-8 -*-
"""Independent semantic and OOXML round-trip checks for native splitting."""
from copy import copy
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET
import warnings

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule, Rule
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName

from engine.tools.excel_merge.merge import Cancelled
from engine.tools.excel_merge.native_split import SplitBlocked, split_preview, split_export


def save(tmp_path, wb, **payload):
    path = tmp_path / "源文件.xlsx"
    wb.save(path)
    return dict(path=str(path), out_dir=str(tmp_path / "out"), **payload)


def apply_choices(payload, preview, code):
    """The UI merges a chosen patch, it never sets a broad ignore-safety flag."""
    payload = dict(payload)
    for issue in preview["issues"]:
        if issue["code"] != code:
            continue
        for key, value in issue["choices"][0]["patch"].items():
            if isinstance(value, dict):
                payload[key] = {**payload.get(key, {}), **value}
            elif isinstance(value, list):
                payload[key] = list(dict.fromkeys(payload.get(key, []) + value))
            else:
                payload[key] = value
    return payload


def result_book(result, label):
    return load_workbook(next(f["path"] for f in result["files"] if f["label"] == label))


def test_row_layout_styles_formulas_rules_and_names(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "年度计划"
    ws.merge_cells("A1:D1")
    ws["A1"] = "2026 年度计划"
    ws["B2"] = 100
    ws.append([])  # the selected header is explicitly row 4
    for col, value in enumerate(["单位", "数量", "预算", "状态"], 1):
        ws.cell(4, col, value)
    for row, values in enumerate([
        ["甲", 2, "=B5*rate", "进行中"],
        ["乙", 3, "=B6*rate", "已完成"],
        ["甲", 4, "=$B7*$B$2", "已完成"],
        [None, None, None, None],
    ], 5):
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
        ws.row_dimensions[row].height = row * 5
    ws.row_dimensions[3].height = 19
    ws.row_dimensions[7].hidden = True
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["D"].hidden = True
    ws["B7"].font = Font(name="宋体", bold=True, color="FFFF0000")
    ws["B7"].fill = PatternFill("solid", fgColor="FFFFFF00")
    ws["B7"].border = Border(bottom=Side(style="double", color="FF000000"))
    ws["B7"].alignment = Alignment(horizontal="right", wrap_text=True)
    ws["B7"].number_format = "0.00"
    ws["B7"].protection = Protection(locked=False)
    ws["B7"].comment = Comment("原备注", "测试")
    ws["A7"].hyperlink = "https://example.com/task"
    ws.freeze_panes = "B5"
    ws.print_area = "A1:D8"
    ws.print_title_rows = "1:4"
    ws.page_setup.orientation = "landscape"
    ws.oddHeader.center.text = "原始页眉"
    ws.auto_filter.ref = "A4:D8"
    dv = DataValidation(type="custom", formula1="B5>=0", allow_blank=True,
                        showErrorMessage=True, error="不能为负", errorTitle="数量")
    dv.add("B5:B8")
    ws.add_data_validation(dv)
    red = Font(color="FFFF0000")
    ws.conditional_formatting.add("D5:D8", FormulaRule(formula=['$B5>1'], font=red))
    ws.conditional_formatting.add("B5:B8", CellIsRule(operator="lessThan", formula=["100"],
                                                    fill=PatternFill("solid", fgColor="FFC6EFCE")))
    wb.defined_names.add(DefinedName("rate", attr_text="'年度计划'!$B$2"))
    wb.defined_names.add(DefinedName("body", attr_text="'年度计划'!$B$5:$B$8"))
    payload = save(tmp_path, wb, selection=dict(axis="rows", header_start=4, header_end=4,
                   data_start=5, data_end=8, field_start=1, field_end=4),
                   group_fields=[dict(kind="field", index=1)])
    before = Path(payload["path"]).read_bytes()
    preview = split_preview(payload)
    assert preview["can_export"], preview["issues"]
    assert [g["records"] for g in preview["groups"]] == [[5, 7], [6], [8]]
    result = split_export(payload)
    out = result_book(result, "甲")
    actual = out["年度计划"]
    assert actual["A1"].value == ws["A1"].value
    assert {str(r) for r in actual.merged_cells.ranges} == {"A1:D1"}
    assert actual["C5"].value == "=B5*rate"
    assert actual["C6"].value == "=$B6*$B$2"
    assert actual["B6"]._style == ws["B7"]._style
    assert actual["B6"].comment.text == "原备注"
    assert actual["A6"].hyperlink.target == "https://example.com/task"
    assert actual["A6"].hyperlink.ref == "A6"
    assert actual.row_dimensions[6].height == 35
    assert actual.row_dimensions[6].hidden
    assert actual.row_dimensions[3].height == 19
    assert actual.column_dimensions["A"].width == 26
    assert actual.column_dimensions["D"].hidden
    assert actual.freeze_panes == "B5"
    assert actual.print_area == "'年度计划'!$A$1:$D$6"
    assert actual.print_title_rows == "$1:$4"
    assert actual.oddHeader.center.text == "原始页眉"
    assert actual.page_setup.orientation == "landscape"
    assert actual.auto_filter.ref == "A4:D6"
    assert out.defined_names["body"].attr_text == "'年度计划'!$B$5:$B$6"
    validations = actual.data_validations.dataValidation
    assert [(str(d.sqref), d.formula1) for d in validations] == [("B5", "B5>=0"), ("B6", "B6>=0")]
    assert all(d.showErrorMessage and d.error == "不能为负" for d in validations)
    rules = [(str(area.sqref), r) for area, items in actual.conditional_formatting._cf_rules.items() for r in items]
    assert {(a, tuple(r.formula)) for a, r in rules} == {
        ("D5", ("$B5>1",)), ("D6", ("$B6>1",)), ("B5", ("100",)), ("B6", ("100",))}
    assert all(r.dxf.font.color.rgb == "FFFF0000" for _, r in rules if r.type == "expression")
    assert all(r.dxf.fill.fgColor.rgb == "FFC6EFCE" for _, r in rules if r.type == "cellIs")
    blank = result_book(result, "（空白）")["年度计划"]
    assert blank.row_dimensions[5].height == 40  # an empty record is retained
    assert Path(payload["path"]).read_bytes() == before


def test_columns_preserve_widths_formulas_rules_and_preceding_columns(tmp_path):
    wb = Workbook()
    ws = wb.active
    for col, unit, number in [(3, "甲", 2), (4, "乙", 3), (5, "甲", 4), (6, "乙", 5)]:
        ws.cell(1, col, unit)
        ws.cell(2, col, number)
        ws.cell(3, col, f"={chr(64 + col)}2*2")
        ws.column_dimensions[chr(64 + col)].width = col * 10
    ws["A1"] = "单位"; ws["A2"] = "数量"; ws["A3"] = "预算"
    ws.column_dimensions["B"].hidden = True
    ws.row_dimensions[3].height = 27
    ws.freeze_panes = "C2"
    ws.print_area = "A1:F3"
    ws.print_title_cols = "A:B"
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="99")
    dv.add("C2:F2"); ws.add_data_validation(dv)
    ws.conditional_formatting.add("C2:F2", FormulaRule(formula=["C2>1"], font=Font(bold=True)))
    payload = save(tmp_path, wb, selection=dict(axis="columns", header_start=1, header_end=2,
                   data_start=3, data_end=6, field_start=1, field_end=3),
                   group_fields=[dict(kind="field", index=1)])
    result = split_export(payload)
    actual = result_book(result, "甲").active
    assert [actual.cell(2, c).value for c in (3, 4)] == [2, 4]
    assert actual["C3"].value == "=C2*2" and actual["D3"].value == "=D2*2"
    assert actual.column_dimensions["C"].width == 30
    assert actual.column_dimensions["D"].width == 50
    assert actual.column_dimensions["B"].hidden
    assert actual.row_dimensions[3].height == 27
    assert actual.freeze_panes == "C2"
    assert actual.print_area == "'Sheet'!$A$1:$D$3"
    assert actual.print_title_cols == "$A:$B"
    assert {(str(a.sqref), tuple(r.formula)) for a, rs in actual.conditional_formatting._cf_rules.items()
            for r in rs} == {("C2", ("C2>1",)), ("D2", ("D2>1",))}
    assert [(str(d.sqref), d.formula1, d.formula2) for d in actual.data_validations.dataValidation] == [
        ("C2", "0", "99"), ("D2", "0", "99")]


@pytest.mark.parametrize("value", [datetime(2026, 3, 31), "2026年3月", "2026-03-31",
                                  "3-5月", "2026年1月至3月", "2025-12至2026-03"])
def test_month_filter_accepts_dates_text_and_ranges(tmp_path, value):
    wb = Workbook(); ws = wb.active
    ws.append(["月份", "事项"]); ws.append([value, "应选中"]); ws.append(["2026-04", "不应选中"])
    payload = save(tmp_path, wb, filters=[dict(field=1, op="month", value="2026-03")])
    assert split_preview(payload)["groups"][0]["records"] == [2]


def test_interval_and_recipients_and_blank_records(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["开始", "结束", "牵头", "参与"])
    ws.append(["2026-02", "2026-04", "甲", "乙"])
    ws.append(["2026-03", "2026-03", "乙", "甲"])
    ws.append(["2026-04", "2026-05", "丙", "丙"])
    payload = save(tmp_path, wb, group_fields=[dict(kind="field", index=3), dict(kind="field", index=4)],
                   group_mode="recipients", filters=[dict(field=1, end_field=2, op="interval", value="2026-03")])
    preview = split_preview(payload)
    assert preview["can_export"]
    assert {x["label"]: x["records"] for x in preview["groups"]} == {"甲": [2, 3], "乙": [2, 3]}
    assert preview["totals"]["output_records"] == 4


def test_unrecognized_month_and_formula_group_are_actionable(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["月份", "单位"]); ws.append(["春季以后", "=1+1"])
    payload = save(tmp_path, wb, filters=[dict(field=1, op="month", value="2026-03")])
    preview = split_preview(payload)
    assert any(x["code"] == "filter_value" and "春季以后" in x["message"] for x in preview["issues"])
    assert not preview["can_export"]
    payload.pop("filters"); payload["group_fields"] = [dict(kind="field", index=2)]
    assert any(x["code"] == "formula_group" for x in split_preview(payload)["issues"])


def test_merged_label_requires_explicit_confirmation_and_clips_anchor(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["分类", "单位", "事项"])
    ws.append(["共享标签", "乙", "事项一"])
    ws.append([None, "甲", "事项二"])
    ws.append([None, "甲", "事项三"])
    ws.merge_cells("A2:A4")
    ws["A2"].fill = PatternFill("solid", fgColor="FFCCEEFF")
    payload = save(tmp_path, wb, group_fields=[dict(kind="field", index=2)])
    preview = split_preview(payload)
    assert not preview["can_export"]
    assert {x["code"] for x in preview["issues"]} == {"merged_record"}
    payload = apply_choices(payload, preview, "merged_record")
    result = split_export(payload)
    actual = result_book(result, "甲").active
    assert actual["A2"].value == "共享标签"
    assert actual["A2"].fill.fgColor.rgb == "FFCCEEFF"
    assert {str(r) for r in actual.merged_cells.ranges} == {"A2:A3"}
    assert [actual.cell(r, 3).value for r in (2, 3)] == ["事项二", "事项三"]


def test_numeric_merged_record_cannot_be_relabelled(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["金额", "单位"]); ws.append([100, "甲"]); ws.append([None, "乙"])
    ws.merge_cells("A2:A3")
    payload = save(tmp_path, wb, group_fields=[dict(kind="field", index=2)],
                   shared_label_ranges=["A2:A3"], allow_warnings=True)
    assert not split_preview(payload)["can_export"]
    with pytest.raises(SplitBlocked):
        split_export(payload)


@pytest.mark.parametrize("formula", ["=SUM(B2:B4)", "=SUM(B:B)", "=SUM(amounts)"])
def test_total_population_needs_scope_decision(tmp_path, formula):
    wb = Workbook(); ws = wb.active
    ws.append(["单位", "数量"]); ws.append(["甲", 2]); ws.append(["乙", 3]); ws.append(["甲", 4])
    ws["A5"] = "总计"; ws["B5"] = formula
    wb.defined_names.add(DefinedName("amounts", attr_text="'Sheet'!$B$2:$B$4"))
    payload = save(tmp_path, wb, selection=dict(axis="rows", header_start=1, header_end=1,
                   data_start=2, data_end=4, field_start=1, field_end=2),
                   group_fields=[dict(kind="field", index=1)], allow_warnings=True)
    preview = split_preview(payload)
    assert not preview["can_export"]
    assert {x["code"] for x in preview["issues"]} == {"formula_scope"}
    payload = apply_choices(payload, preview, "formula_scope")
    actual = result_book(split_export(payload), "甲")
    expected = "=SUM(B2:B3)" if formula == "=SUM(B2:B4)" else formula
    assert actual.active["B4"].value == expected
    assert actual.defined_names["amounts"].attr_text == "'Sheet'!$B$2:$B$3"


def test_scalar_dependency_removed_is_never_blank_or_zero(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["单位", "数量", "公式"]); ws.append(["甲", 2, "=B3"]); ws.append(["乙", 3, "=B2"])
    payload = save(tmp_path, wb, group_fields=[dict(kind="field", index=1)], allow_warnings=True)
    preview = split_preview(payload)
    assert any(x["code"] == "missing_reference" for x in preview["issues"])
    with pytest.raises(SplitBlocked):
        split_export(payload)
    assert not Path(payload["out_dir"]).exists()


def test_hidden_list_config_only_auto_retained_when_all_content_is_used(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["单位", "状态"]); ws.append(["甲", "进行中"]); ws.append(["乙", "完成"])
    cfg = wb.create_sheet("配置"); cfg.sheet_state = "hidden"
    cfg.append(["进行中"]); cfg.append(["完成"])
    private = wb.create_sheet("其他业务"); private.append(["不能发出"])
    wb.defined_names.add(DefinedName("states", attr_text="'配置'!$A$1:$A$2"))
    dv = DataValidation(type="list", formula1="states"); dv.add("B2:B3"); ws.add_data_validation(dv)
    payload = save(tmp_path, wb, group_fields=[dict(kind="field", index=1)])
    preview = split_preview(payload)
    assert preview["can_export"], preview["issues"]
    assert preview["retained_sheets"] == ["Sheet", "配置"]
    assert preview["omitted_sheets"] == ["其他业务"]
    actual = result_book(split_export(payload), "甲")
    assert actual.sheetnames == ["Sheet", "配置"]
    assert actual["配置"].sheet_state == "hidden"
    assert actual.active.data_validations.dataValidation[0].formula1 == "states"
    cfg["Z20"] = "无关的敏感业务"
    wb.save(payload["path"])
    preview = split_preview(payload)
    assert not preview["can_export"]
    assert any(x["code"] == "dependency_sheet" for x in preview["issues"])
    payload = apply_choices(payload, preview, "dependency_sheet")
    assert split_preview(payload)["can_export"]


def test_visible_dependency_and_back_reference_relocated(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["单位", "值", "计算"]); ws.append(["甲", 2, "='费率'!$A$1*B2"])
    ws.append(["乙", 3, "='费率'!$A$1*B3"])
    rate = wb.create_sheet("费率"); rate["A1"] = 100
    payload = save(tmp_path, wb, group_fields=[dict(kind="field", index=1)])
    preview = split_preview(payload)
    assert any(x["code"] == "dependency_sheet" for x in preview["issues"])
    payload = apply_choices(payload, preview, "dependency_sheet")
    actual = result_book(split_export(payload), "乙")
    assert actual.active["C2"].value == "='费率'!$A$1*B2"
    assert actual["费率"]["A1"].value == 100


def test_sheet_and_selected_cell_group_bases(tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = "三月"
    ws.merge_cells("A1:B1"); ws["A1"] = "工程部"
    ws.append(["事项", "数量"]); ws.append(["任务", 2])
    payload = save(tmp_path, wb, selection=dict(axis="rows", header_start=2, header_end=2,
                   data_start=3, data_end=3, field_start=1, field_end=2),
                   group_fields=[dict(kind="cell", ref="B1"), dict(kind="sheet")])
    preview = split_preview(payload)
    assert preview["can_export"]
    assert preview["groups"][0]["label"] == "工程部 · 三月"


def test_cancel_rolls_back_only_this_export_and_existing_files_survive(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["单位"]); ws.append(["甲"]); ws.append(["乙"])
    payload = save(tmp_path, wb, group_fields=[dict(kind="field", index=1)])
    out = Path(payload["out_dir"]); out.mkdir()
    existing = out / "源文件_甲.xlsx"; existing.write_bytes(b"existing")
    state = {"cancel": False}
    def progress(*args):
        state["cancel"] = True
    with pytest.raises(Cancelled):
        split_export(payload, cancel=lambda: state["cancel"], progress=progress)
    assert list(out.iterdir()) == [existing]
    assert existing.read_bytes() == b"existing"
    result = split_export(payload)
    assert len(result["files"]) == 2
    assert any(Path(f["path"]).name == "源文件_甲_2.xlsx" for f in result["files"])
    assert existing.read_bytes() == b"existing"


def test_non_xlsx_and_unsafe_names_and_invalid_axis(tmp_path):
    with pytest.raises(ValueError, match=".xlsx"):
        split_preview(dict(path=str(tmp_path / "legacy.xls")))
    wb = Workbook(); ws = wb.active
    ws.append(["单位"]); ws.append(["../../危险:*?"]); ws.append(["CON"])
    payload = save(tmp_path, wb, group_fields=[dict(kind="field", index=1)])
    result = split_export(payload)
    assert all(Path(f["path"]).parent == Path(payload["out_dir"]) for f in result["files"])
    with pytest.raises(ValueError, match="axis"):
        split_preview(dict(payload, selection={"axis": "diagonal"}))


def _add_extension(path, part, feature=None, encoding="utf-8"):
    with ZipFile(path) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    xml = ET.fromstring(parts[part])
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    container = xml.find("{" + main + "}extLst")
    if container is None:
        container = ET.SubElement(xml, "{" + main + "}extLst")
    if feature is not None:
        extension = ET.SubElement(container, "{" + main + "}ext", uri="{78C0D931-6437-407D-A8EE-F0AAD7539E65}")
        if feature:
            ET.SubElement(extension, "{http://schemas.microsoft.com/office/spreadsheetml/2009/9/main}" + feature)
    parts[part] = ET.tostring(xml, encoding=encoding)
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def test_empty_xml_extension_has_nothing_to_lose(tmp_path):
    wb = Workbook(); wb.active.append(["单位"]); wb.active.append(["甲"])
    payload = save(tmp_path, wb)
    _add_extension(payload["path"], "xl/worksheets/sheet1.xml")
    preview = split_preview(payload)
    assert preview["can_export"], preview
    assert not any("扩展" in notice for notice in preview["warnings"])
    actual = result_book(split_export(payload), preview["groups"][0]["label"])
    assert actual.active["A2"].value == "甲"
    actual.close()


def test_empty_extension_wrapper_read_warning_allows_export(tmp_path):
    wb = Workbook(); wb.active.append(["单位"]); wb.active.append(["甲"])
    payload = save(tmp_path, wb)
    _add_extension(payload["path"], "xl/worksheets/sheet1.xml", "")
    preview = split_preview(payload)
    assert preview["can_export"], preview
    assert any("extension is not supported and will be removed" in x for x in preview["warnings"])
    result = split_export(payload)
    actual = result_book(result, result["groups"][0]["label"])
    assert actual.active["A2"].value == "甲"
    actual.close()


@pytest.mark.parametrize("part,feature,encoding", [("xl/styles.xml", "slicerStyles", "utf-8"),
                                         ("xl/worksheets/sheet1.xml", "conditionalFormattings", "utf-8"),
                                         ("xl/worksheets/sheet1.xml", "futureCellData", "utf-8"),
                                         ("xl/workbook.xml", "futureWorkbookFeature", "utf-8"),
                                         ("xl/workbook.xml", "futureWorkbookFeature", "utf-16")])
def test_extensions_allow_export_with_notice_and_regular_cells_round_trip(tmp_path, part, feature, encoding):
    wb = Workbook(); ws = wb.active
    ws.append(["单位", "数量", "计算"]); ws.append(["甲", 2, "=B2*2"])
    ws["B2"].font = Font(name="宋体", bold=True, color="FF123456")
    ws["B2"].number_format = "0.00"
    ws["B2"].border = Border(bottom=Side(style="thin", color="FF123456"))
    ws.page_setup.orientation = "landscape"; ws.page_margins.left = 0.4
    payload = save(tmp_path, wb)
    _add_extension(payload["path"], part, feature, encoding)
    original = Path(payload["path"]).read_bytes()
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("error", UserWarning)
        preview = split_preview(payload)
        result = split_export(payload)
    assert not caught
    assert preview["can_export"], preview
    assert any("扩展" in notice for notice in preview["warnings"])
    assert any("扩展" in notice for notice in result["warnings"])
    actual = result_book(result, preview["groups"][0]["label"])
    assert actual.active["B2"].value == 2 and actual.active["C2"].value == "=B2*2"
    assert copy(actual.active["B2"].font) == copy(ws["B2"].font)
    assert copy(actual.active["B2"].border) == copy(ws["B2"].border)
    assert actual.active["B2"].number_format == "0.00"
    assert actual.active.page_setup.orientation == "landscape"
    assert actual.active.page_margins.left == 0.4
    assert Path(payload["path"]).read_bytes() == original
    actual.close()


def test_printer_binary_is_notice_and_standard_page_setup_survives(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["单位"]); ws.append(["甲"])
    ws.page_setup.paperSize = "9"; ws.page_setup.orientation = "landscape"
    ws.page_margins.left = 0.6
    payload = save(tmp_path, wb)
    with ZipFile(payload["path"], "a") as archive:
        archive.writestr("xl/printerSettings/printerSettings1.bin", b"printer-specific-settings")
    result = split_export(payload)
    assert result["can_export"] and any("打印机" in x for x in result["warnings"])
    actual = result_book(result, result["groups"][0]["label"])
    assert actual.active.page_setup.paperSize == 9
    assert actual.active.page_setup.orientation == "landscape"
    assert actual.active.page_margins.left == 0.6
    with ZipFile(result["files"][0]["path"]) as archive:
        assert not any("printerSettings" in x for x in archive.namelist())
    actual.close()


def test_style_permission_does_not_allow_unrelated_read_warning(tmp_path, monkeypatch):
    from engine.tools.excel_merge import native_split
    wb = Workbook(); wb.active.append(["单位"]); wb.active.append(["甲"])
    payload = save(tmp_path, wb)
    _add_extension(payload["path"], "xl/styles.xml", "slicerStyles")
    original_load = native_split.load_workbook
    def warned_load(*args, **kwargs):
        warnings.warn("Unreadable formula dependency", UserWarning)
        return original_load(*args, **kwargs)
    monkeypatch.setattr(native_split, "load_workbook", warned_load)
    preview = split_preview(payload)
    assert not preview["can_export"]
    assert any(x["code"] == "read_warning" for x in preview["issues"])


def test_rule_relative_reference_at_last_cell_is_not_silently_changed(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["单位", "数量"]); ws.append(["甲", 2]); ws.append(["甲", 3]); ws.append(["乙", 4])
    # At A2 the rule refers to retained B3; at A3 it refers to removed B4.
    dv = DataValidation(type="custom", formula1="B3>0")
    dv.add("A2:A3"); ws.add_data_validation(dv)
    payload = save(tmp_path, wb, group_fields=[dict(kind="field", index=1)])
    preview = split_preview(payload)
    assert not preview["can_export"]
    assert any(x["code"] == "missing_reference" and "验证" in x["location"] for x in preview["issues"])


def test_internal_hyperlinks_move_and_removed_targets_block(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["单位", "数量"]); ws.append(["甲", 2]); ws.append(["乙", 3]); ws.append(["甲", 4])
    ws["A4"].hyperlink = "#'Sheet'!B4"
    payload = save(tmp_path, wb, group_fields=[dict(kind="field", index=1)])
    actual = result_book(split_export(payload), "甲").active
    assert actual["A3"].hyperlink.target == "#'Sheet'!B3"
    ws["A4"].hyperlink = "#'Sheet'!B3"
    wb.save(payload["path"])
    assert any(x["code"] == "missing_reference" for x in split_preview(payload)["issues"])


def test_content_outside_fields_is_not_deleted_without_decision(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["单位", "数量", None, "另一个表"])
    ws.append(["甲", 2, None, "其他记录"]); ws.append(["乙", 3, None, "其他记录二"])
    payload = save(tmp_path, wb, selection=dict(axis="rows", header_start=1, header_end=1,
                   data_start=2, data_end=3, field_start=1, field_end=2),
                   group_fields=[dict(kind="field", index=1)])
    preview = split_preview(payload)
    assert any(x["code"] == "outside_fields" for x in preview["issues"])
    payload = apply_choices(payload, preview, "outside_fields")
    assert split_preview(payload)["can_export"]


def test_plain_excel_table_is_kept_and_ref_adjusted(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["单位", "数量"]); ws.append(["甲", 2]); ws.append(["乙", 3]); ws.append(["甲", 4])
    ws.add_table(Table(displayName="Tasks", ref="A1:B4"))
    payload = save(tmp_path, wb, group_fields=[dict(kind="field", index=1)])
    actual = result_book(split_export(payload), "甲").active
    assert actual.tables["Tasks"].ref == "A1:B3"
    assert actual.tables["Tasks"].autoFilter.ref == "A1:B3"


def test_statistical_cf_keeps_one_population_across_noncontiguous_source_rows(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["单位", "数量"]); ws.append(["甲", 2]); ws.append(["乙", 3]); ws.append(["甲", 4])
    ws.conditional_formatting.add("B2:B4", ColorScaleRule(
        start_type="min", start_color="FFFF0000", end_type="max", end_color="FF00FF00"))
    ws.conditional_formatting.add("B2:B4", Rule(type="top10", rank=1))
    payload = save(tmp_path, wb, group_fields=[dict(kind="field", index=1)])
    actual = result_book(split_export(payload), "甲").active
    rules = [(str(a.sqref), r) for a, rs in actual.conditional_formatting._cf_rules.items() for r in rs]
    assert [(area, rule.type) for area, rule in rules] == [("B2:B3", "colorScale"), ("B2:B3", "top10")]
    colors = rules[0][1].colorScale.color
    assert [c.rgb for c in colors] == ["FFFF0000", "FF00FF00"]


def test_blank_group_previous_is_explicit_preserves_blank_and_uses_filtered_predecessor(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["单位", "月份", "任务"])
    ws.append(["甲", "2026-02", "前置任务"])
    ws.append([None, "2026-03", "当月任务"])
    ws.append(["乙", "2026-03", "另一任务"])
    payload = save(tmp_path, wb, group_fields=[dict(kind="field", index=1)],
                   filters=[dict(field=2, op="month", value="2026-03")])
    preview = split_preview(payload)
    assert preview["can_export"]
    assert {g["label"] for g in preview["groups"]} == {"（空白）", "乙"}
    notice = next(i for i in preview["issues"] if i["code"] == "blank_groups")
    assert notice["blocking"] is False
    payload.update(notice["choices"][0]["patch"])
    result = split_export(payload)
    actual = result_book(result, "甲").active
    assert actual["A2"].value is None  # attribution must not fill the cell
    assert actual["C2"].value == "当月任务"


def test_output_fields_month_column_composes_with_row_split_and_rule_mapping(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.merge_cells("A1:F1"); ws["A1"] = "年度计划"
    ws.append(["单位", "任务", "1月", "2月", "3月", "全年"])
    for row, values in enumerate([
        ["甲", "任务一", 1, 2, 3],
        [None, "任务二", 4, 5, 6],
        ["乙", "任务三", 7, 8, 9],
    ], 3):
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
        ws.cell(row, 6, f"=SUM(C{row}:E{row})")
        ws.row_dimensions[row].height = row * 7
    ws.column_dimensions["E"].width = 26
    ws.freeze_panes = "C3"; ws.print_area = "A1:F5"
    dv = DataValidation(type="custom", formula1="E3>=0"); dv.add("E3:E5"); ws.add_data_validation(dv)
    ws.conditional_formatting.add("E3:E5", FormulaRule(formula=["E3>0"], font=Font(color="FFFF0000")))
    payload = save(tmp_path, wb, selection=dict(axis="rows", header_start=2, header_end=2,
                   data_start=3, data_end=5, field_start=1, field_end=6),
                   group_fields=[dict(kind="field", index=1)], blank_group="previous",
                   output_fields=[1, 2, 5])
    preview = split_preview(payload)
    assert preview["can_export"], preview["issues"]
    sample = next(g for g in preview["groups"] if g["label"] == "甲")
    assert sample["columns"] == [{"index": 1, "label": "单位"}, {"index": 2, "label": "任务"},
                                 {"index": 5, "label": "3月"}]
    assert sample["sample"] == [["甲", "任务一", 3], [None, "任务二", 6]]
    actual = result_book(split_export(payload), "甲").active
    assert actual.max_column == 3
    assert [actual.cell(2, c).value for c in range(1, 4)] == ["单位", "任务", "3月"]
    assert [actual.cell(3, c).value for c in range(1, 4)] == ["甲", "任务一", 3]
    assert [actual.cell(4, c).value for c in range(1, 4)] == [None, "任务二", 6]
    assert {str(r) for r in actual.merged_cells.ranges} == {"A1:C1"}
    assert actual.column_dimensions["C"].width == 26
    assert actual.row_dimensions[4].height == 28
    assert actual.freeze_panes == "C3"
    assert actual.print_area == "'Sheet'!$A$1:$C$4"
    assert [(str(d.sqref), d.formula1) for d in actual.data_validations.dataValidation] == [("C3:C4", "C3>=0")]
    assert [(str(a.sqref), r.formula) for a, rs in actual.conditional_formatting._cf_rules.items()
            for r in rs] == [("C3:C4", ["C3>0"])]


def test_output_fields_never_silently_drops_a_required_formula_operand(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["单位", "数量", "金额"]); ws.append(["甲", 2, "=B2*100"])
    payload = save(tmp_path, wb, group_fields=[dict(kind="field", index=1)], output_fields=[1, 3])
    preview = split_preview(payload)
    assert not preview["can_export"]
    assert any(i["code"] == "missing_reference" for i in preview["issues"])
