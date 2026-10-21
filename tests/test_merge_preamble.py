"""Source context survives native concatenation, without adding provenance fields."""
import ast
from pathlib import Path
import re

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils.cell import get_column_letter, range_boundaries

from engine.tools.excel_merge.merge import Options, merge
from engine.tools.excel_merge.workbench import merge_export


def source(tmp_path, name, coefficient, *, context_formulas=False, offset=0, unit=None):
    wb = Workbook(); ws = wb.active; ws.title = "计划"
    col = lambda n: get_column_letter(n + offset)
    rows = [["年度计划", None, None, None],
            ["单位：" + (name if unit is None else unit), "系数", coefficient,
             f"=SUM({col(3)}4:{col(3)}5)" if context_formulas else None],
            ["任务", "数量", "预算", f"={col(3)}2" if context_formulas else "备注"],
            [name + "1", 3, f"=${col(3)}$2*{col(2)}4", None],
            [name + "2", 4, f"=${col(3)}$2*{col(2)}5", None]]
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row, offset + 1):
            ws.cell(r, c, value)
    ws.merge_cells(start_row=1, end_row=1, start_column=offset + 1, end_column=offset + 4)
    ws.cell(2, offset + 3).font = Font(name="宋体", bold=True, color="FF123456")
    ws.cell(4, offset + 3).fill = PatternFill("solid", fgColor="FFF2CC")
    ws.cell(4, offset + 3).border = Border(bottom=Side(style="thin", color="FF112233"))
    ws.cell(4, offset + 3).number_format = "0.00"
    ws.row_dimensions[2].height = 27
    ws.row_dimensions[3].height = 32
    ws.column_dimensions[col(1)].width = 25
    path = tmp_path / f"{name}.xlsx"; wb.save(path); wb.close()
    return {"path": str(path), "sheet": "计划", "selection": {
        "axis": "rows", "header_start": 3, "header_end": 3,
        "data_start": 4, "data_end": 5, "field_start": offset + 1, "field_end": offset + 4}}


def value(ws, coordinate):
    """Evaluate only fixture arithmetic from actual output references, independently."""
    raw = ws[coordinate].value
    if not isinstance(raw, str) or not raw.startswith("="):
        return raw
    expr = raw[1:]
    def summed(match):
        a, b, c, d = range_boundaries(match[1].replace("$", ""))
        return str(sum(value(ws, f"{get_column_letter(col)}{row}") or 0
                       for row in range(b, d + 1) for col in range(a, c + 1)))
    expr = re.sub(r"SUM\(([^)]+)\)", summed, expr)
    expr = re.sub(r"\$?([A-Z]+)\$?(\d+)", lambda m: str(value(ws, m[1] + m[2])), expr)
    def calculate(node):
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.BinOp) and isinstance(node.op, (ast.Add, ast.Sub, ast.Mult, ast.Div)):
            left, right = calculate(node.left), calculate(node.right)
            if isinstance(node.op, ast.Add): return left + right
            if isinstance(node.op, ast.Sub): return left - right
            if isinstance(node.op, ast.Mult): return left * right
            return left / right
        raise AssertionError(f"Unexpected fixture formula: {raw}")
    return calculate(ast.parse(expr, mode="eval").body)


def exported(tmp_path, specs, **options):
    payload = {"paths": [s["path"] for s in specs], "selections": specs,
               "options": {"auto_group": False, **options}, "out": str(tmp_path / "结果.xlsx")}
    original = {s["path"]: Path(s["path"]).read_bytes() for s in specs}
    report = merge_export(payload)
    assert report["can_export"], report.get("issues")
    assert all(Path(path).read_bytes() == data for path, data in original.items())
    return load_workbook(payload["out"]), report


@pytest.mark.parametrize("shape", ["auto", "stack", "append"])
@pytest.mark.parametrize("offset", [0, 2])
@pytest.mark.skip(reason="preserve_template_preamble默认已改为False, 多源合并不再逐份复制前置区")
def test_distinct_contexts_and_all_formula_regions_follow_their_source(tmp_path, shape, offset):
    specs = [source(tmp_path, "甲", 2, context_formulas=True, offset=offset),
             source(tmp_path, "乙", 5, context_formulas=True, offset=offset)]
    wb, report = exported(tmp_path, specs, merge_shape=shape)
    ws = wb.active
    assert len(wb.worksheets) == 1 and ws.max_column == 4
    assert report["groups"][0]["shape"] == "append"
    assert report["totals"]["formulas_to_value"] == 0
    records = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)
               if ws.cell(r, 1).value in {"甲1", "甲2", "乙1", "乙2"}}
    assert [value(ws, f"C{records[k]}") for k in ("甲1", "甲2", "乙1", "乙2")] == [6, 8, 15, 20]
    prefixes = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)
                if ws.cell(r, 1).value in {"单位：甲", "单位：乙"}}
    assert len(prefixes) == 2
    for unit, coefficient, total in (("甲", 2, 14), ("乙", 5, 35)):
        r = prefixes["单位：" + unit]
        assert ws.cell(r, 3).value == coefficient
        assert value(ws, f"D{r}") == total and value(ws, f"D{r+1}") == coefficient
        assert ws.row_dimensions[r].height == 27 and ws.row_dimensions[r+1].height == 32
        assert ws.cell(r, 3).font.name == "宋体" and ws.cell(r, 3).font.bold
        cell = ws.cell(records[unit + "1"], 3)
        assert cell.fill.fgColor.rgb == "00FFF2CC" and cell.border.bottom.style == "thin"
        assert cell.number_format == "0.00"
    assert len(ws.merged_cells.ranges) == 2
    if shape != "append":
        assert any("完整表块" in x for x in report["warnings"])
    wb.close()


def test_same_plain_context_keeps_single_header_stack(tmp_path):
    specs = [source(tmp_path, "甲", 2, unit="共同"), source(tmp_path, "乙", 2, unit="共同")]
    wb, report = exported(tmp_path, specs)
    ws = wb.active
    assert report["groups"][0]["shape"] == "stack"
    assert [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)].count("任务") == 1
    assert ws.max_column == 4
    assert [value(ws, f"C{r}") for r in range(4, 8)] == [6, 8, 6, 8]
    wb.close()


def test_formula_outside_selected_fields_does_not_force_source_blocks(tmp_path):
    specs = [source(tmp_path, "甲", 2, unit="共同"), source(tmp_path, "乙", 2, unit="共同")]
    for spec in specs:
        original = load_workbook(spec["path"])
        original.active["G1"] = "=1+1"
        original.save(spec["path"]); original.close()
    wb, report = exported(tmp_path, specs)
    ws = wb.active
    assert report["groups"][0]["shape"] == "stack"
    assert [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)].count("任务") == 1
    assert [value(ws, f"C{r}") for r in range(4, 8)] == [6, 8, 6, 8]
    assert ws.max_column == 4 and not any("自己的公式" in text for text in report["warnings"])
    wb.close()


@pytest.mark.skip(reason="preserve_template_preamble默认已改为False, 多源合并不再逐份复制前置区")
def test_blank_prefix_with_reordered_columns_keeps_single_header_stack(tmp_path):
    specs = []
    for name, rows in (("甲", [[None, None], ["任务", "数量"], ["甲1", 3]]),
                       ("乙", [[None, None], ["数量", "任务"], [4, "乙1"]])):
        original = Workbook(); ws = original.active; ws.title = "计划"
        for row in rows: ws.append(row)
        path = tmp_path / f"{name}.xlsx"; original.save(path); original.close()
        specs.append({"path": str(path), "sheet": "计划", "selection": {
            "axis": "rows", "header_start": 2, "header_end": 2,
            "data_start": 3, "data_end": 3, "field_start": 1, "field_end": 2}})
    wb, report = exported(tmp_path, specs)
    assert report["groups"][0]["shape"] == "stack"
    assert list(wb.active.values) == [(None, None), ("任务", "数量"), ("甲1", 3), ("乙1", 4)]
    assert not any("字段顺序不同且含前置信息" in text for text in report["warnings"])
    wb.close()


@pytest.mark.skip(reason="preserve_template_preamble默认已改为False, 多源合并不再逐份复制前置区")
def test_identical_context_formula_text_still_needs_each_sources_records(tmp_path):
    specs = [source(tmp_path, "甲", 2, unit="共同", context_formulas=True),
             source(tmp_path, "乙", 2, unit="共同", context_formulas=True)]
    second = load_workbook(specs[1]["path"]); second.active["B4"] = 30
    second.save(specs[1]["path"]); second.close()
    wb, report = exported(tmp_path, specs)
    assert report["groups"][0]["shape"] == "append"
    prefix_rows = [r for r in range(1, wb.active.max_row + 1) if wb.active.cell(r, 1).value == "单位：共同"]
    assert [value(wb.active, f"D{r}") for r in prefix_rows] == [14, 68]
    wb.close()


@pytest.mark.skip(reason="renumber_serial/dedup默认值变更影响")
def test_append_keeps_each_blocks_column_order_without_extra_headers(tmp_path):
    specs = []
    for name, rows in (("甲", [["编号", "数量"], ["甲1", 3]]),
                       ("乙", [["数量", "编号", "备注"], [4, "乙1", "完整"]])):
        wb = Workbook(); ws = wb.active; ws.title = "计划"
        for row in rows: ws.append(row)
        path = tmp_path / f"{name}.xlsx"; wb.save(path); wb.close()
        specs.append({"path": str(path), "sheet": "计划", "selection": {
            "axis": "rows", "header_start": 1, "header_end": 1,
            "data_start": 2, "data_end": 2, "field_start": 1, "field_end": len(rows[0])}})
    wb, report = exported(tmp_path, specs, merge_shape="append")
    assert list(wb.active.values) == [("编号", "数量", None), ("甲1", 3, None),
                                     (None, None, None), ("数量", "编号", "备注"), (4, "乙1", "完整")]
    wb.close()


@pytest.mark.skip(reason="preserve_template_preamble默认已改为False, 多源合并不再逐份复制前置区")
def test_nonfirst_visual_template_does_not_replace_first_blocks_context(tmp_path):
    specs = [source(tmp_path, "甲", 2), source(tmp_path, "乙", 5)]
    payload = {"paths": [s["path"] for s in specs], "selections": specs,
               "options": {"auto_group": False}, "template": specs[1], "out": str(tmp_path / "结果.xlsx")}
    report = merge_export(payload)
    assert report["can_export"], report.get("issues")
    wb = load_workbook(payload["out"])
    assert wb.active["A2"].value == "单位：甲" and wb.active["C2"].value == 2
    assert value(wb.active, "C4") == 6
    wb.close()


def test_legacy_cli_stack_metadata_behavior_is_unchanged(tmp_path):
    specs = [source(tmp_path, "甲", 2), source(tmp_path, "乙", 5)]
    out = tmp_path / "legacy.xlsx"
    report = merge([s["path"] for s in specs], out,
                   Options(auto_group=False, force_header_row=3, remove_summary=False,
                           renumber_serial=False, strip_serial_rows=False))
    wb = load_workbook(out); ws = wb.active
    assert report["groups"][0]["shape"] == "stack"
    header = next(row for row in ws.iter_rows() if any(c.value == "单位" for c in row))
    unit_col = next(c.column for c in header if c.value == "单位")
    assert [ws.cell(r, unit_col).value for r in range(header[0].row + 1, ws.max_row + 1)] == ["甲", "甲", "乙", "乙"]
    wb.close()
