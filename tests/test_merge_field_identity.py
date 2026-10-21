"""Explicit correspondences remain attached to their field after external edits."""
import pytest
from types import SimpleNamespace
from openpyxl import Workbook, load_workbook

from engine.tools.excel_merge.align import align_sources, pair_columns
from engine.tools.excel_merge.preview import preview
from engine.tools.excel_merge.workbench import merge_export, merge_preview


def source(path, prefix, plan, actual):
    wb = Workbook(); ws = wb.active
    for row in [["客户", "计划", "实际"], [prefix + "1", plan, actual], [prefix + "2", plan + 1, actual + 1]]:
        ws.append(row)
    wb.save(path); wb.close()
    return preview({"path": str(path)})


def swap_columns(path):
    wb = load_workbook(path)
    for row in wb.active:
        row[1].value, row[2].value = row[2].value, row[1].value
    wb.save(path); wb.close()


def payload(tmp_path):
    first, second = tmp_path / "底表.xlsx", tmp_path / "反馈.xlsx"
    target, current = source(first, "A", 2, 1), source(second, "B", 9, 20)
    specs = [{"path": str(path), "sheet": view["sheet"], "selection": view["selection"],
              "selection_origin": "auto"} for path, view in ((first, target), (second, current))]
    specs[1].update(field_map={"1": 1, "2": 2, "3": 3},
                    field_map_context={"source": current["fields"], "target": target["fields"]})
    return {"paths": [str(first), str(second)], "selections": specs,
            "options": {"auto_group": False}, "out": str(tmp_path / "结果.xlsx")}


@pytest.mark.parametrize("changed", ["source", "template", "both"])
def test_changed_column_order_relocates_explicit_mapping_using_snapshot(tmp_path, changed):
    request = payload(tmp_path)
    if changed in ("source", "both"):
        swap_columns(request["paths"][1])
    if changed in ("template", "both"):
        swap_columns(request["paths"][0])
    report = merge_export(request)
    assert report["can_export"], report
    wb = load_workbook(request["out"])
    try:
        rows = list(wb.active.values)
        columns = {name: index for index, name in enumerate(rows[0])}
        record = next(row for row in rows[1:] if row[0] == "B1")
        assert record[columns["计划"]] == 9
        assert record[columns["实际"]] == 20
    finally:
        wb.close()


def test_changed_field_identity_does_not_reuse_its_old_coordinate(tmp_path):
    request = payload(tmp_path)
    wb = load_workbook(request["paths"][1])
    wb.active["B1"] = "其他字段"
    wb.save(request["paths"][1]); wb.close()
    result = merge_preview(request)
    assert not result["can_export"]
    assert any("表头已变化" in issue["message"] for issue in result["issues"])
    with pytest.raises(ValueError, match="表头已变化"):
        merge_export(request)


def test_legacy_literal_json_field_is_not_an_encoded_hierarchy():
    assert pair_columns(['["计划","数量"]'], ["计划/数量"]) == {}
    first = SimpleNamespace(structure=SimpleNamespace(keys=['["计划","数量"]']))
    second = SimpleNamespace(structure=SimpleNamespace(keys=["计划/数量"]))
    keys, maps = align_sources([first, second], "tolerant")
    assert keys == ['["计划","数量"]', "计划/数量"]
    assert maps == [[0, None], [None, 0]]
