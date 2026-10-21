"""Optional dual-output workflow must publish all its files or none."""
from pathlib import Path
import threading

import pytest
from openpyxl import Workbook, load_workbook

from engine import sidecar
from engine.tools.excel_merge import workflow
from engine.tools.excel_merge.merge import Cancelled
from engine.tools.excel_merge.preview import preview


def payload(tmp_path, total=False):
    def write(path, quantity):
        wb = Workbook()
        ws = wb.active
        ws.title = "明细"
        ws.append(["编号", "月份", "数量"])
        ws.append(["R1", "1月", quantity])
        ws.append(["R2", "2月", 20])
        if total:
            ws.append(["合计", None, "=SUM(C2:C3)"])
        wb.save(path)
        wb.close()
    original, feedback = tmp_path / "current.xlsx", tmp_path / "feedback.xlsx"
    write(original, 10)
    write(feedback, 12)
    selection = {"axis": "rows", "header_start": 1, "header_end": 1,
                 "data_start": 2, "data_end": 3, "field_start": 1, "field_end": 3}
    return {"path": str(original), "sheet": "明细", "selection": selection,
            "feedback": [{"path": str(feedback), "sheet": "明细", "selection": selection}],
            "key_fields": [1], "update_fields": [3], "compare_fields": [3],
            "filtered_output": {"filters": [{"field": 2, "op": "equals", "value": "1月"}]},
            "out": str(tmp_path / "result.xlsx")}


def destinations(p):
    out = Path(p["out"])
    return [out, out.with_name(out.stem + "_筛选结果.xlsx"), out.with_suffix(".changes.json")]


@pytest.mark.parametrize("failure_index", [2, 3])
def test_publish_failure_rolls_back_all_created_outputs(tmp_path, monkeypatch, failure_index):
    p = payload(tmp_path)
    targets = destinations(p)
    original_link = workflow.os.link
    count = 0
    def fail_link(source, destination):
        nonlocal count
        if Path(destination) in targets:
            count += 1
            if count == failure_index:
                raise OSError("模拟最后发布失败")
        return original_link(source, destination)
    monkeypatch.setattr(workflow.os, "link", fail_link)
    with pytest.raises(OSError):
        workflow.workflow_export(p)
    assert not any(path.exists() for path in targets)
    assert not list(tmp_path.glob(".excel-feedback-*"))


@pytest.mark.parametrize("cancel_after", [1, 3])
def test_cancel_during_publication_leaves_no_partial_outputs(tmp_path, monkeypatch, cancel_after):
    p = payload(tmp_path)
    targets = destinations(p)
    original_link = workflow.os.link
    cancelled = threading.Event()
    count = 0
    def publish_then_cancel(source, destination):
        nonlocal count
        result = original_link(source, destination)
        if Path(destination) in targets:
            count += 1
            if count == cancel_after:
                cancelled.set()
        return result
    monkeypatch.setattr(workflow.os, "link", publish_then_cancel)
    with pytest.raises(Cancelled):
        workflow.workflow_export(p, cancel=cancelled.is_set)
    assert not any(path.exists() for path in targets)


@pytest.mark.parametrize("index", [0, 1, 2])
def test_any_existing_destination_is_never_overwritten(tmp_path, index):
    p = payload(tmp_path)
    target = destinations(p)[index]
    target.write_bytes(b"user-owned")
    with pytest.raises(ValueError):
        workflow.workflow_export(p)
    assert target.read_bytes() == b"user-owned"
    assert all(not path.exists() for path in destinations(p) if path != target)


def test_racing_destination_is_preserved_and_our_first_file_removed(tmp_path, monkeypatch):
    p = payload(tmp_path)
    targets = destinations(p)
    original_link = workflow.os.link
    def race(source, destination):
        if Path(destination) == targets[0]:
            targets[1].write_bytes(b"another-process")
        return original_link(source, destination)
    monkeypatch.setattr(workflow.os, "link", race)
    with pytest.raises(FileExistsError):
        workflow.workflow_export(p)
    assert targets[1].read_bytes() == b"another-process"
    assert not targets[0].exists() and not targets[2].exists()


def test_empty_optional_filter_exports_whole_workbook_only(tmp_path):
    p = payload(tmp_path)
    p["filtered_output"]["filters"][0]["value"] = "12月"
    ready = workflow.workflow_preview(p)
    assert ready["can_export"], ready
    assert not ready["filtered_preview"]["groups"]
    result = workflow.workflow_export(p)
    assert result["exported"], result
    targets = destinations(p)
    assert targets[0].exists() and targets[2].exists() and not targets[1].exists()
    assert any("只导出" in warning for warning in result["warnings"])
    wb = load_workbook(targets[0])
    assert wb.active["C2"].value == 12 and wb.active["A3"].value == "R2"
    wb.close()


def test_empty_filter_does_not_hide_other_blocking_errors(tmp_path):
    p = payload(tmp_path)
    p["filtered_output"]["filters"] = [{"field": 2, "op": "month", "value": "不合法的月份"}]
    ready = workflow.workflow_preview(p)
    assert not ready["can_export"]
    assert any(issue.get("code") == "filter_value" for issue in ready["issues"])
    result = workflow.workflow_export(p)
    assert not result["exported"]
    assert not any(path.exists() for path in destinations(p))


def test_nested_filtered_resolution_survives_preview_and_export(tmp_path):
    p = payload(tmp_path, total=True)
    first = workflow.workflow_preview(p)
    issue = next(issue for issue in first["issues"] if issue.get("code") == "formula_scope")
    assert issue["id"].startswith("filtered:") and not first["can_export"]
    patch = issue["choices"][0]["patch"]
    p["filtered_output"].update(patch["filtered_output"])
    ready = workflow.workflow_preview(p)
    assert ready["can_export"], ready
    result = workflow.workflow_export(p)
    assert result["exported"]
    wb = load_workbook(result["filtered_output_path"])
    assert wb.active["C3"].value == "=SUM(C2:C2)"
    wb.close()


def test_stale_or_already_cancelled_queued_export_cannot_create_outputs(tmp_path):
    p = payload(tmp_path)
    event = threading.Event()
    event.set()
    with pytest.raises(Cancelled):
        sidecar._dispatch({**p, "cmd": "reconcile"}, event)
    assert not any(path.exists() for path in destinations(p))
    event.clear()
    fingerprint = sidecar.source_fingerprint(p)
    wb = load_workbook(p["feedback"][0]["path"])
    wb.active["C2"] = 999
    wb.save(p["feedback"][0]["path"])
    wb.close()
    with pytest.raises(ValueError, match="预览后发生变化"):
        sidecar._dispatch({**p, "cmd": "reconcile", "preview_fingerprint": fingerprint}, event)
    assert not any(path.exists() for path in destinations(p))


@pytest.mark.parametrize("axis", ["rows", "columns"])
def test_empty_and_blank_header_preview_still_offer_explicit_coordinates(tmp_path, axis):
    path = tmp_path / "empty.xlsx"
    wb = Workbook()
    wb.save(path)
    wb.close()
    empty = preview({"path": str(path), "selection": {"axis": axis}})
    assert empty["selection"]["data_start"] > empty["selection"]["data_end"]
    assert empty["fields"][0]["empty"]
    wb = load_workbook(path)
    if axis == "rows":
        wb.active["A2"] = "R1"
        wb.active["B2"] = 10
    else:
        wb.active["B1"] = "R1"
        wb.active["B2"] = 10
    wb.save(path)
    wb.close()
    data = preview({"path": str(path), "selection": {"axis": axis, "header_start": 1, "header_end": 1,
                    "data_start": 2, "data_end": 2, "field_start": 1, "field_end": 2}})
    assert all(field["empty"] for field in data["fields"])
