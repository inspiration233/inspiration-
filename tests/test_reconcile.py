"""Business invariants for reconciliation, using independent temporary workbooks."""
from copy import copy, deepcopy
import json
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import warnings

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from engine.tools.excel_merge.merge import Cancelled
from engine.tools.excel_merge.reconcile import reconcile_export, reconcile_preview


def _add_extension(path, part, feature=None, encoding="utf-8"):
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    xml = ET.fromstring(parts[part])
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    container = xml.find("{" + main + "}extLst")
    if container is None:
        container = ET.SubElement(xml, "{" + main + "}extLst")
    if feature is not None:
        extension = ET.SubElement(container, "{" + main + "}ext", uri="{78C0D931-6437-407D-A8EE-F0AAD7539E65}")
        if feature:
            ET.SubElement(extension, "{http://schemas.microsoft.com/office/spreadsheetml/2009/9/main}" + feature)
    parts[part] = ET.tostring(xml, encoding=encoding)
    with ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def book(path, records, headers=("编号", "单位", "计划", "完成", "状态"), axis="rows",
         native=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "计划"
    table = [list(headers)] + [list(r) for r in records]
    if axis == "columns":
        table = list(map(list, zip(*table)))
    for row in table:
        ws.append(row)
    if native:
        ws.freeze_panes = "C2"
        ws.column_dimensions["C"].width = 27
        ws.row_dimensions[2].height = 35
        ws["C2"].font = Font(name="宋体", size=13, bold=True, italic=True, color="FF112233")
        ws["C2"].number_format = "0.00"
        ws["C2"].fill = PatternFill("solid", fgColor="FFFFCC")
        ws.page_setup.orientation = "landscape"
        ws.page_margins.left = 0.4
        ws.print_title_rows = "1:1"
        ws.print_area = f"A1:E{len(records)+1}"
        validation = DataValidation(type="list", formula1='"进行中,已完成,已取消"')
        validation.add(f"E2:E{len(records)+1}")
        ws.add_data_validation(validation)
        ws.conditional_formatting.add(
            f"D2:D{len(records)+1}", CellIsRule(operator="greaterThan", formula=["5"],
                                              fill=PatternFill("solid", fgColor="C6EFCE")))
        other = wb.create_sheet("说明")
        other["A1"] = "保留此页"
    wb.save(path)
    wb.close()
    return path


def selection(records, fields=5, axis="rows"):
    return {"axis": axis, "header_start": 1, "header_end": 1,
            "data_start": 2, "data_end": records + 1, "field_start": 1, "field_end": fields}


def payload(tmp_path, current_records, feedback_records, axis="rows", native=False):
    current = book(tmp_path / "current.xlsx", current_records, axis=axis, native=native)
    feedback = book(tmp_path / "feedback.xlsx", feedback_records, axis=axis)
    return {"path": str(current), "sheet": "计划", "selection": selection(len(current_records), axis=axis),
            "feedback": [{"path": str(feedback), "sheet": "计划",
                          "selection": selection(len(feedback_records), axis=axis)}],
            "key_fields": [1], "update_fields": [3, 4, 5], "compare_fields": [3, 5],
            "result_form": "direct",      # 旧用例覆盖直接替换路径; 对比输出有专测
            "out": str(tmp_path / "result.xlsx")}


def resolve(p, preview, kind, value):
    p.setdefault("resolutions", {})
    for issue in preview["issues"]:
        if issue["kind"] == kind:
            p["resolutions"][issue["id"]] = value


@pytest.mark.parametrize("part,feature,encoding", [("xl/styles.xml", "slicerStyles", "utf-8"),
                                         ("xl/worksheets/sheet1.xml", "conditionalFormattings", "utf-8"),
                                         ("xl/worksheets/sheet1.xml", "futureCellData", "utf-8"),
                                         ("xl/workbook.xml", "futureWorkbookFeature", "utf-8"),
                                         ("xl/workbook.xml", "futureWorkbookFeature", "utf-16")])
def test_extensions_report_notice_without_blocking_updates(tmp_path, part, feature, encoding):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")],
                [("T1", "甲", 12, 3, "进行中")], native=True)
    current = load_workbook(p["path"])
    current["计划"]["F1"] = "公式保留"
    current["计划"]["F2"] = "=C2*2"
    current.save(p["path"]); current.close()
    _add_extension(p["path"], part, feature, encoding)
    original = Path(p["path"]).read_bytes()
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("error", UserWarning)
        preview = reconcile_preview(p)
        result = reconcile_export(p)
    assert not caught
    assert preview["can_export"] and result["exported"], result
    assert any("扩展" in notice for notice in preview["warnings"])
    assert result["warnings"] == preview["warnings"]
    actual = load_workbook(p["out"])
    ws = actual["计划"]
    assert ws["C2"].value == 12 and ws["D2"].value == 3
    assert ws["F2"].value == "=C2*2"
    assert ws["C2"].font.name == "宋体" and ws["C2"].font.bold and ws["C2"].font.italic
    assert str(ws["C2"].fill.fgColor.rgb)[-6:] == "FF0000" and ws["C2"].number_format == "0.00"
    assert ws.page_setup.orientation == "landscape" and ws.page_margins.left == 0.4
    assert len(ws.data_validations.dataValidation) == 1 and len(ws.conditional_formatting) == 1
    assert Path(p["path"]).read_bytes() == original
    log = json.loads(Path(result["changes_path"]).read_text(encoding="utf-8"))
    assert log["warnings"] == result["warnings"]
    actual.close()


def test_empty_extensions_pass_and_extension_removal_warning_does_not_block(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")], [("T1", "甲", 12, 3, "进行中")])
    _add_extension(p["path"], "xl/worksheets/sheet1.xml")
    preview = reconcile_preview(p)
    assert preview["can_export"] and not preview["warnings"], preview
    _add_extension(p["path"], "xl/worksheets/sheet1.xml", "")
    preview = reconcile_preview(p)
    assert preview["can_export"], preview
    assert any("extension is not supported and will be removed" in x for x in preview["warnings"])
    assert reconcile_export(p)["exported"]
    actual = load_workbook(p["out"])
    assert actual["计划"]["C2"].value == 12
    actual.close()


def test_printer_binary_is_notice_and_keeps_standard_page_setup(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")],
                [("T1", "甲", 12, 3, "进行中")], native=True)
    with ZipFile(p["path"], "a") as archive:
        archive.writestr("xl/printerSettings/printerSettings1.bin", b"printer-specific-settings")
    result = reconcile_export(p)
    assert result["exported"] and any("打印机" in x for x in result["warnings"]), result
    actual = load_workbook(p["out"])
    assert actual["计划"].page_setup.orientation == "landscape"
    assert actual["计划"].page_margins.left == 0.4
    with ZipFile(p["out"]) as archive:
        assert not any("printerSettings" in x for x in archive.namelist())
    actual.close()


def test_style_permission_does_not_allow_unrelated_read_warning(tmp_path, monkeypatch):
    from engine.tools.excel_merge import reconcile
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")], [("T1", "甲", 12, 3, "进行中")])
    _add_extension(p["path"], "xl/styles.xml", "slicerStyles")
    original_load = reconcile.load_workbook
    def warned_load(*args, **kwargs):
        warnings.warn("Unreadable formula dependency", UserWarning)
        return original_load(*args, **kwargs)
    monkeypatch.setattr(reconcile, "load_workbook", warned_load)
    preview = reconcile_preview(p)
    assert not preview["can_export"]
    assert any("Unreadable formula dependency" in issue["message"] for issue in preview["issues"])


def test_updates_missing_and_native_preservation(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中"), ("T2", "乙", 20, 2, "进行中")],
                [("T1", "甲", 12, 8, "已取消")], native=True)
    before = Path(p["path"]).read_bytes()
    preview = reconcile_preview(p)
    assert preview["can_export"], preview
    assert preview["missing"] == [{"key": ["T2"], "position": 3, "status": "not_received",
                                   "action": "keep_current", "message": "未交回，疑似被删除；已保留当前内容"}]
    assert {c["field"] for c in preview["changes"]} == {3, 4, 5}
    assert next(c for c in preview["changes"] if c["field"] == 5)["kind"] == "cancel"
    result = reconcile_export(p)
    assert result["exported"], result
    assert Path(p["path"]).read_bytes() == before
    wb = load_workbook(p["out"])
    ws = wb["计划"]
    assert [ws.cell(2, c).value for c in range(1, 6)] == ["T1", "甲", 12, 8, "已取消"]
    assert ws["C3"].value == 20 and ws["D3"].value == 2
    assert str(ws["C2"].fill.fgColor.rgb)[-6:] == "FF0000"
    assert ws["C2"].font.name == "宋体" and ws["C2"].font.bold and ws["C2"].font.italic
    # 改动默认全部红填充: D2(0→8)同为红; 未改动格(如 C3)不标红
    assert str(ws["D2"].fill.fgColor.rgb)[-6:] == "FF0000"
    assert ws["C3"].fill.fill_type != "solid" or str(ws["C3"].fill.fgColor.rgb or "")[-6:] != "FF0000"
    assert ws["C2"].number_format == "0.00" and ws.column_dimensions["C"].width == 27
    assert ws.row_dimensions[2].height == 35 and ws.freeze_panes == "C2"
    assert ws.page_setup.orientation == "landscape" and ws.page_margins.left == 0.4
    assert ws.print_title_rows == "$1:$1"
    assert str(ws.data_validations.dataValidation[0].sqref) == "E2:E3"
    assert len(ws.conditional_formatting) == 1 and wb["说明"]["A1"].value == "保留此页"
    wb.close()
    audit = json.loads(Path(result["changes_path"]).read_text(encoding="utf-8"))
    assert audit["changes"] == result["changes"]


def test_default_requires_update_fields_and_no_compare_means_no_red(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")], [("T1", "甲", 11, 8, "进行中")])
    p.pop("update_fields")
    assert not reconcile_preview(p)["can_export"]
    p["update_fields"] = [3, 4]
    p.pop("compare_fields")
    result = reconcile_export(p)
    assert result["exported"]
    wb = load_workbook(p["out"])
    assert wb.active["C2"].font.color.type == "theme"
    wb.close()


def test_column_axis_and_composite_keys(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中"), ("T1", "乙", 20, 0, "进行中")],
                [("T1", "乙", 25, 5, "已完成")], axis="columns")
    p["key_fields"] = [1, 2]
    result = reconcile_export(p)
    assert result["exported"], result
    wb = load_workbook(p["out"])
    ws = wb.active
    assert ws["B3"].value == 10 and ws["C3"].value == 25 and ws["C4"].value == 5
    assert str(ws["C3"].fill.fgColor.rgb)[-6:] == "FF0000"
    wb.close()


def test_duplicate_identity_in_base_table_blocks(tmp_path):
    # 底表主键重复且内容不同: 自动跳过会让反馈被当成新记录追加, 必须修表
    current = [("T1", "甲", 10, 0, "进行中"), ("T1", "乙", 20, 0, "进行中")]
    feedback = [("T1", "甲", 11, 1, "进行中"), ("T2", "乙", 21, 1, "进行中")]
    p = payload(tmp_path, current, feedback)
    preview = reconcile_preview(p)
    assert not preview["can_export"]
    assert any(i["kind"] == "duplicate_key" for i in preview["issues"])
    assert not reconcile_export(p)["exported"]
    assert not Path(p["out"]).exists()


def test_merged_key_column_matches_by_block(tmp_path):
    # 名称类主键纵向合并: 合并块=一条记录, 自动按块匹配(块内第k行↔反馈同名块第k行), 不再阻断
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中"), ("T1", "乙", 20, 1, "进行中"),
                           ("T2", "丙", 30, 2, "进行中"), ("T2", "丁", 40, 3, "进行中")],
                [("T1", "甲", 11, 1, "进行中")])
    wb = load_workbook(p["path"])
    wb.active.merge_cells("A2:A3")
    wb.active.merge_cells("A4:A5")
    wb.save(p["path"])
    wb.close()
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    assert any("合并块 = 一条记录" in w for w in preview["warnings"])
    assert len(preview["changes"]) == 2     # 计划 10→11 与 完成 0→1
    assert all(c["position"] == 2 for c in preview["changes"])   # 都落在 T1 块的锚点行
    assert reconcile_export(p)["exported"]


def test_merged_name_key_auto_matches_by_block(tmp_path):
    # 组名(名称)列做主键且纵向合并: 自动按合并块匹配, 更新落在块锚点行
    p = payload(tmp_path, [(1, "甲", 10, 0, "进行中"), (2, "甲", 20, 1, "进行中"),
                           (3, "乙", 30, 2, "进行中"), (4, "乙", 40, 3, "进行中")],
                [(1, "甲", 11, 1, "进行中")])
    p["key_fields"] = [2]
    p["update_fields"] = [3, 4]
    wb = load_workbook(p["path"])
    wb.active.merge_cells("B2:B3")
    wb.active.merge_cells("B4:B5")
    wb.save(p["path"])
    wb.close()
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    assert any("合并块 = 一条记录" in w for w in preview["warnings"])
    change = next(c for c in preview["changes"])
    assert change["key"][0] == "甲" and change["after"] == 11 and change["position"] == 2


def test_group_mode_matches_per_group_ordinal(tmp_path):
    # 分发-回填链路: 底稿序号全局连续(甲1-4/乙5-8), 反馈序号按组内从1开始。
    # 分组模式下按(组名, 组内第几条)对应, 序号数值不要求一致。
    base = [("例1", "甲", 1, 0), (1, "甲", 10, 0), (2, "甲", 20, 0), (3, "甲", 30, 0),
            (4, "甲", 40, 0), (5, "乙", 50, 0), (6, "乙", 60, 0), (7, "乙", 70, 0), (8, "乙", 80, 0)]
    feedback = [("例1", "甲", 1, 0), (1, "甲", 10, 0), (2, "甲", 20, 0), (3, "甲", 30, 0),
                (4, "甲", 40, 0), (1, "乙", 50, 0), (2, "乙", 66, 1), (3, "乙", 70, 0), (4, "乙", 80, 0)]
    p = payload(tmp_path, base, feedback)
    p["key_fields"] = [1]
    p["update_fields"] = [3]
    p["group_field"] = 2
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    assert len(preview["changes"]) == 1
    change = preview["changes"][0]
    assert change["key"] == ["乙", 2] and change["after"] == 66
    assert change["position"] == 8          # 底稿里乙组第2条 = 序号6所在行
    assert preview["totals"] ["additions"] == 0 and preview["totals"]["missing"] == 0
    assert any("分组匹配" in w for w in preview["warnings"])
    assert reconcile_export(p)["exported"]
    wb = load_workbook(p["out"])
    assert wb.active["C8"].value == 66      # 值落在正确的行
    wb.close()


def test_group_deletion_identified_by_serial_key(tmp_path):
    # 按主键对位: 交回序号 1,2,4 → 缺口=3, 精确点名"缺少序号3", 且序号4对位不滑
    base = [(1, "甲", 10, 0, "进行中"), (2, "甲", 20, 0, "进行中"),
            (3, "甲", 30, 0, "进行中"), (4, "甲", 40, 0, "进行中")]
    feedback = [(1, "甲", 11, 0, "进行中"), (2, "甲", 21, 0, "进行中"),
                (4, "甲", 44, 0, "进行中")]                      # 序号 3 被删除, 未重排
    p = payload(tmp_path, base, feedback)
    p["key_fields"] = [1]; p["update_fields"] = [3]; p["group_field"] = 2
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    assert any("缺少序号 3" in w and "组内第 3 条" in w for w in preview["warnings"])
    assert not any("整体错位" in w for w in preview["warnings"])
    # 序号 4 的值落在底稿第 4 条上(不滑位)
    change = next(c for c in preview["changes"] if c["after"] == 44)
    assert change["key"] == ["甲", 4]
    assert any("被删除" in m["message"] for m in preview["missing"])


def test_group_deletion_identified_by_original_serial(tmp_path):
    # 拆分默认保留底稿原序号(组内不一定是 1..N): 底稿组内序号 2,5,9,
    # 反馈删了序号 5 那条 → 按"序号值=底稿序号值"对位, 点名缺 5 且序号 9 不滑位
    base = [(2, "甲", 10, 0, "进行中"), (5, "甲", 20, 0, "进行中"),
            (9, "甲", 30, 0, "进行中")]
    feedback = [(2, "甲", 11, 0, "进行中"), (9, "甲", 99, 0, "进行中")]
    p = payload(tmp_path, base, feedback)
    p["key_fields"] = [1]; p["update_fields"] = [3]; p["group_field"] = 2
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    assert any("缺少序号 5" in w for w in preview["warnings"])
    assert not any("整体错位" in w for w in preview["warnings"])
    # 序号 9 的值落在底稿第 3 条上(不滑位)
    change = next(c for c in preview["changes"] if c["after"] == 99)
    assert change["key"] == ["甲", 3]
    assert any("被删除" in m["message"] for m in preview["missing"])


def test_flat_mode_identical_key_sets_across_feedback_warns(tmp_path):
    # 非分组模式: 两份反馈的主键值完全相同(拆分产物序号被每组从 1 重排的典型特征)
    # → 点名提示改用分组回填或重新拆分, 而不是让用户对着一堆冲突记录
    base = tmp_path / "base.xlsx"
    fb1 = tmp_path / "fb1.xlsx"
    fb2 = tmp_path / "fb2.xlsx"
    book(base, [(1, "甲", 10, 0, "进行中"), (2, "甲", 20, 0, "进行中"), (3, "甲", 30, 0, "进行中")])
    book(fb1, [(1, "甲", 11, 0, "进行中"), (2, "甲", 21, 0, "进行中"), (3, "甲", 31, 0, "进行中")])
    book(fb2, [(1, "甲", 12, 0, "进行中"), (2, "甲", 22, 0, "进行中"), (3, "甲", 32, 0, "进行中")])
    p = {"path": str(base), "sheet": "计划", "selection": selection(2),
         "feedback": [{"path": str(fb1), "sheet": "计划", "selection": selection(2)},
                      {"path": str(fb2), "sheet": "计划", "selection": selection(2)}],
         "key_fields": [1], "update_fields": [3], "result_form": "direct",
         "resolutions": {}, "out": str(tmp_path / "out.xlsx")}
    preview = reconcile_preview(p)
    assert any("主键值完全相同" in w and "分组" in w for w in preview["warnings"])


def test_group_feedback_renumbered_with_deletion_not_pretending_precise(tmp_path):
    # 反馈端删行后又把序号拉成 1..n: 与"删了末尾"同貌, 不能冒充精确点名——
    # 应退数量警示并点名疑似重新编号, 请人工核对整组
    base = [(2, "甲", 10, 0, "进行中"), (5, "甲", 20, 0, "进行中"),
            (9, "甲", 30, 0, "进行中")]
    feedback = [(1, "甲", 11, 0, "进行中"), (2, "甲", 99, 0, "进行中")]   # 删了序号5, 重排成1,2
    p = payload(tmp_path, base, feedback)
    p["key_fields"] = [1]; p["update_fields"] = [3]; p["group_field"] = 2
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    assert any("疑似被重新编号" in w and "无法定位" in w for w in preview["warnings"])
    assert not any("缺少序号" in w for w in preview["warnings"])


def test_group_feedback_renumbered_full_set_aligns_by_ordinal(tmp_path):
    # 反馈端重排但没删行(交回序号恰为 1..N): 重排值 k = 组内第 k 条, 对位完全正确
    base = [(2, "甲", 10, 0, "进行中"), (5, "甲", 20, 0, "进行中"),
            (9, "甲", 30, 0, "进行中")]
    feedback = [(1, "甲", 10, 0, "进行中"), (2, "甲", 20, 0, "进行中"),
                (3, "甲", 55, 0, "进行中")]                              # 全量交回, 重排
    p = payload(tmp_path, base, feedback)
    p["key_fields"] = [1]; p["update_fields"] = [3]; p["group_field"] = 2
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    assert not any("组「" in w and ("少" in w or "缺少" in w) for w in preview["warnings"])
    change = next(c for c in preview["changes"] if c["after"] == 55)
    assert change["position"] == 4          # 底稿序号 9 所在行
    assert preview["totals"]["missing"] == 0


def test_multi_sheet_reconcile_both_sheets_updated(tmp_path):
    # 多表回填: 底稿两张表同名存在于反馈, 各自独立回填, 输出合成一个文件
    base = tmp_path / "base.xlsx"
    fb = tmp_path / "fb.xlsx"
    book(base, [(1, "甲", 10, 0, "进行中"), (2, "乙", 20, 0, "进行中")])
    book(fb, [(1, "甲", 11, 0, "进行中"), (2, "乙", 20, 0, "进行中")])
    # 第二张表(结构与总表不同)也造同名表
    from openpyxl import load_workbook
    for path, mark in ((base, "底稿"), (fb, "反馈")):
        wb = load_workbook(path)
        ws = wb.create_sheet("附表")
        ws.append(["编号", "备注"])
        ws.append(["A1", mark]); ws.append(["A2", mark])
        wb.save(path); wb.close()
    p = {"path": str(base), "sheet": "计划", "selection": selection(2),
         "feedback": [{"path": str(fb), "sheet": "计划", "selection": selection(2)}],
         "key_fields": [1], "update_fields": [3], "result_form": "direct",
         "resolutions": {}, "out": str(tmp_path / "out.xlsx")}
    p["sheets"] = [
        {"sheet": "计划", "selection": selection(2), "key_fields": [1], "update_fields": [3]},
        {"sheet": "附表", "selection": {"axis": "rows", "header_start": 1, "header_end": 1,
                                        "data_start": 2, "data_end": 3, "field_start": 1, "field_end": 2},
         "key_fields": [1], "update_fields": [2]},
    ]
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    assert all(c.get("sheet") for c in preview["changes"])
    assert any(c["sheet"] == "计划" and c["after"] == 11 for c in preview["changes"])
    assert any(c["sheet"] == "附表" and c["after"] == "反馈" for c in preview["changes"])
    result = reconcile_export(p)
    assert result["exported"], result
    wb = load_workbook(p["out"])
    assert wb["计划"]["C2"].value == 11 and wb["附表"]["B2"].value == "反馈"
    wb.close()


def _fill_hex(cell):
    fill = cell.fill
    return str(fill.fgColor.rgb or "")[-6:] if fill and fill.fill_type == "solid" else None


def test_compare_output_inserts_original_rows(tmp_path):
    # 保留对比: 改动行下插完整原值行(新值红/原值橙), 删除行整行橙, 新增行整行红
    base = [("例1", "甲", 1, 0), (1, "甲", 10, 5), (2, "甲", 20, 6), (3, "甲", 30, 7), (1, "乙", 50, 8)]
    feedback = [("例1", "甲", 1, 0), (1, "甲", 11, 5), (2, "甲", 20, 6), (1, "乙", 50, 9), (9, "乙", 90, 99)]
    p = payload(tmp_path, base, feedback)
    p["key_fields"] = [1]; p["update_fields"] = [3, 4]; p["group_field"] = 2
    p["result_form"] = "compare"
    result = reconcile_export(p)
    assert result["exported"], result
    wb = load_workbook(p["out"]); ws = wb["计划"]
    # 甲1(原第3行)改了C: 新值行@3 红, 原值行@4 橙, 其余值完整
    assert ws.cell(3, 3).value == 11 and _fill_hex(ws.cell(3, 3)) in ("FF0000", "E74C3C", "C0392B", "D63031")
    assert ws.cell(4, 3).value == 10 and _fill_hex(ws.cell(4, 3)) in ("F39C12", "E67E22", "FFC000", "F5A623")
    assert ws.cell(4, 1).value == 1 and ws.cell(4, 4).value == 5   # 原值行完整
    assert _fill_hex(ws.cell(3, 4)) is None                        # 未变格无填充
    # 甲3 疑似被删除: 上方插空值行(整行红, 反馈现状=不存在) + 原值行(整行橙)
    r_blank, r_del = 6, 7
    assert ws.cell(r_blank, 1).value is None and ws.cell(r_blank, 3).value is None
    assert _fill_hex(ws.cell(r_blank, 3)) in ("FF0000", "E74C3C", "C0392B", "D63031")
    assert ws.cell(r_del, 1).value == 3
    assert _fill_hex(ws.cell(r_del, 3)) in ("F39C12", "E67E22", "FFC000", "F5A623")
    # 新增行(乙9)整行红: 按值定位
    new_r = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 3).value == 90)
    assert _fill_hex(ws.cell(new_r, 3)) in ("FF0000", "E74C3C", "C0392B", "D63031")
    wb.close()


def test_compare_output_columns_axis(tmp_path):
    base = [(1, "甲", 10, 0, "进行中"), (2, "乙", 20, 0, "进行中")]
    feedback = [(1, "甲", 11, 0, "进行中"), (2, "乙", 20, 0, "进行中")]
    p = payload(tmp_path, base, feedback, axis="columns")
    p["result_form"] = "compare"
    result = reconcile_export(p)
    assert result["exported"], result
    wb = load_workbook(p["out"]); ws = wb.active
    # 第2列(记录1)改了C行(字段3): 其右插入原值列
    assert ws.cell(3, 2).value == 11 and _fill_hex(ws.cell(3, 2)) is not None
    assert ws.cell(3, 3).value == 10 and _fill_hex(ws.cell(3, 3)) is not None
    wb.close()


def test_diff_categories_fill_and_clear(tmp_path):
    # 差异分类: 底稿空→反馈值=填写; 底稿值→反馈空=清空; 值→值=改值
    base = [(1, "甲", 10, 5, "进行中"), (2, "乙", None, 6, "进行中")]
    feedback = [(1, "甲", None, 7, "进行中"), (2, "乙", 99, None, "进行中")]
    p = payload(tmp_path, base, feedback)
    p["update_fields"] = [3, 4]
    preview = reconcile_preview(p)
    kinds = {(c["key"][0], c["field"]): c["kind"] for c in preview["changes"]}
    assert kinds[(1, 3)] == "clear" and kinds[(1, 4)] == "update"
    assert kinds[(2, 3)] == "fill" and kinds[(2, 4)] == "clear"
    assert preview["totals"] == {**preview["totals"], "updates": 1, "fills": 1, "clears": 2}


def test_group_deletion_conservative_when_serials_unusable(tmp_path):
    # 序号体系被打乱(越界值): 退回按位对位+数量差警示+错位提醒
    base = [(1, "甲", 10, 0, "进行中"), (2, "甲", 20, 0, "进行中"),
            (3, "甲", 30, 0, "进行中"), (4, "甲", 40, 0, "进行中")]
    feedback = [(2, "甲", 21, 0, "进行中"), (3, "甲", 31, 0, "进行中"),
                (9, "甲", 99, 0, "进行中")]                      # 序号 9 越界
    p = payload(tmp_path, base, feedback)
    p["key_fields"] = [1]; p["update_fields"] = [3]; p["group_field"] = 2
    preview = reconcile_preview(p)
    assert any("交回 3 条" in w and "底稿有 4 条" in w for w in preview["warnings"])
    assert any("整体错位" in w for w in preview["warnings"])


def test_group_full_coverage_no_deletion_warning(tmp_path):
    base = [(1, "甲", 10, 0, "进行中"), (2, "甲", 20, 0, "进行中")]
    feedback = [(1, "甲", 11, 0, "进行中"), (2, "甲", 21, 0, "进行中")]
    p = payload(tmp_path, base, feedback)
    p["key_fields"] = [1]; p["update_fields"] = [3]; p["group_field"] = 2
    preview = reconcile_preview(p)
    assert not any("删除" in w for w in preview["warnings"])


def test_group_mode_reports_unknown_group_as_new_records(tmp_path):
    base = [(1, "甲", 10, 0), (2, "甲", 20, 0)]
    feedback = [(1, "甲", 11, 0), (1, "丙", 99, 0)]
    p = payload(tmp_path, base, feedback)
    p["key_fields"] = [1]
    p["update_fields"] = [3]
    p["group_field"] = 2
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    assert any("底表中没有的分组：丙" in w for w in preview["warnings"])
    assert any(c["kind"] == "new" and c["key"] == ["丙", 1] for c in preview["changes"])
    assert next(c for c in preview["changes"] if c["kind"] != "new")["after"] == 11


def test_group_mode_requires_single_key_field(tmp_path):
    p = payload(tmp_path, [(1, "甲", 10, 0)], [(1, "甲", 11, 0)])
    p["key_fields"] = [1, 2]
    p["update_fields"] = [3]
    p["group_field"] = 2
    preview = reconcile_preview(p)
    assert not preview["can_export"]
    assert any(i["kind"] == "field_selection" and "一个序号列" in i["message"] for i in preview["issues"])


def test_conflicting_feedback_for_absent_record_says_ignored(tmp_path):
    # 底表没有该主键时, 反馈内多条矛盾反馈被整组忽略, 不再说"已保留当前内容"
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")],
                [("T2", "乙", 15, 1, "进行中"), ("T2", "丙", 16, 2, "进行中")])
    preview = reconcile_preview(p)
    assert not any("已保留当前内容" in w for w in preview["warnings"])
    assert any("已全部忽略" in w for w in preview["warnings"])


def test_duplicate_identity_in_feedback_keeps_current(tmp_path):
    # 同一反馈文件内主键重复且值不同: 视同冲突, 默认保留当前内容
    current = [("T1", "甲", 10, 0, "进行中"), ("T2", "乙", 20, 0, "进行中")]
    feedback = [("T1", "甲", 11, 1, "进行中"), ("T1", "乙", 99, 1, "进行中")]
    p = payload(tmp_path, current, feedback)
    preview = reconcile_preview(p)
    assert preview["can_export"]
    assert any("不同反馈" in w for w in preview["warnings"])
    result = reconcile_export(p)
    assert result["exported"]
    wb = load_workbook(p["out"])
    assert wb.active["C2"].value == 10      # T1 保留当前
    wb.close()


def test_duplicate_identity_differing_only_in_irrelevant_columns_merges(tmp_path):
    # 重复主键仅无关列不同(回填关注字段一致): 自动合并, 不阻断
    current = [("T1", "甲", 10, 0, "进行中"), ("T2", "乙", 20, 0, "进行中")]
    feedback = [("T1", "甲", 11, 1, "进行中"), ("T1", "丙", 11, 1, "进行中")]  # 仅姓名列(非回填字段)不同
    p = payload(tmp_path, current, feedback)
    preview = reconcile_preview(p)
    assert preview["can_export"]
    assert any("内容一致" in w for w in preview["warnings"])


@pytest.mark.parametrize("role", ["current", "feedback"])
def test_empty_identity_skipped_with_warning_not_blocked(tmp_path, role):
    current = [("T1", "甲", 10, 0, "进行中"), ("T2", "乙", 20, 0, "进行中")]
    feedback = [("T1", "甲", 11, 1, "进行中"), ("T2", "乙", 21, 1, "进行中")]
    (current if role == "current" else feedback)[1] = (None, "乙", 20, 0, "进行中")
    p = payload(tmp_path, current, feedback)
    preview = reconcile_preview(p)
    # 空主键本身不再作为阻断问题, 只汇总提示
    assert not any(i["kind"] == "empty_key" for i in preview["issues"])
    assert any("主键为空" in w for w in preview["warnings"])
    if role == "feedback":
        assert preview["can_export"]
        assert reconcile_export(p)["exported"]
    else:
        # 底表身份缺失后, 反馈中的 T2 成为"是否追加新记录"的真实询问(见 addition 决策)
        assert all(i["kind"] in {"addition"} for i in preview["issues"])


@pytest.mark.parametrize("role", ["feedback"])
def test_identical_duplicate_identity_merged_automatically(tmp_path, role):
    current = [("T1", "甲", 10, 0, "进行中"), ("T2", "乙", 20, 0, "进行中")]
    feedback = [("T1", "甲", 11, 1, "进行中"), ("T2", "乙", 21, 1, "进行中")]
    records = current if role == "current" else feedback
    records.append(tuple(records[0]))          # 完全相同的一条重复记录
    p = payload(tmp_path, current, feedback)
    preview = reconcile_preview(p)
    assert preview["can_export"]
    assert any("内容一致" in w for w in preview["warnings"])


def test_conflicting_feedback_defaults_to_keep_current(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")], [("T1", "甲", 11, 1, "进行中")])
    second = book(tmp_path / "other.xlsx", [("T1", "甲", 12, 2, "进行中")])
    p["feedback"].append({"path": str(second), "selection": selection(1)})
    preview = reconcile_preview(p)
    # 冲突不再阻断: 默认保留当前内容, 冲突明细进变更清单与提示
    assert preview["can_export"]
    assert not any(i["kind"] == "conflict" for i in preview["issues"])
    assert any("不同反馈" in w for w in preview["warnings"])
    conflicts = [c for c in preview["changes"] if c["kind"] == "conflict"]
    assert len(conflicts) == 1 and len(conflicts[0]["candidates"]) == 2
    result = reconcile_export(p)
    assert result["exported"]
    wb = load_workbook(p["out"])
    assert wb.active["C2"].value == 10          # 保留当前
    wb.close()


def test_duplicate_import_is_idempotent(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")], [("T1", "甲", 11, 1, "进行中")])
    single = reconcile_preview(p)
    p["feedback"].append(deepcopy(p["feedback"][0]))
    repeated = reconcile_preview(p)
    assert repeated["can_export"] and repeated["changes"] == single["changes"]
    assert repeated["totals"]["duplicate_imports"] == 1


def test_addition_auto_appends_and_extends_rules_and_template(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")],
                [("T1", "甲", 10, 0, "进行中"), ("T2", "乙", 15, 1, "进行中")], native=True)
    preview = reconcile_preview(p)
    # 安全场景下新记录默认自动追加, 不再确认
    assert preview["can_export"] and not any(i["kind"] == "addition" for i in preview["issues"])
    result = reconcile_export(p)
    assert result["exported"], result
    wb = load_workbook(p["out"])
    ws = wb["计划"]
    assert ws["A3"].value == "T2" and ws["B3"].value == "乙" and ws["C3"].value == 15
    assert ws["C3"].font.bold and ws["C3"].number_format == "0.00"
    assert str(ws.data_validations.dataValidation[0].sqref) == "E2:E3"
    assert str(next(iter(ws.conditional_formatting)).sqref) == "D2:D3"
    assert str(ws.print_area) == "'计划'!$A$1:$E$3"
    wb.close()
    assert any(c["kind"] == "new" for c in result["changes"])


def test_addition_with_footer_auto_skips(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中"), ("页脚", None, None, None, None)],
                [("T2", "乙", 15, 1, "进行中")])
    p["selection"]["data_end"] = 2
    preview = reconcile_preview(p)
    # 跳过是唯一安全动作: 自动执行并聚合提示, 不再逐条要求人工确认
    assert preview["can_export"], preview["issues"]
    assert not any(i["kind"] == "addition" for i in preview["issues"])
    assert preview["totals"]["additions_skipped"] == 1
    assert any("自动跳过" in w for w in preview["warnings"])
    assert reconcile_export(p)["exported"]
    wb = load_workbook(p["out"])
    assert wb.active["A3"].value == "页脚" and wb.active.max_row == 3
    wb.close()


def test_merged_shadow_target_auto_skips_and_anchor_still_writes(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中"), ("T2", "乙", 20, 0, "进行中")],
                [("T1", "甲", 11, 1, "进行中"), ("T2", "乙", 21, 2, "进行中")])
    wb = load_workbook(p["path"])
    wb.active.merge_cells("C2:C3")   # T1 的计划列成为锚点, T2 的 C3 成为遮盖区
    wb.save(p["path"])
    wb.close()
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    assert not any(i["kind"] == "merged_target" for i in preview["issues"])
    assert any("非左上角" in w for w in preview["warnings"])
    fields = [c["field"] for c in preview["changes"]]
    assert 3 in fields                      # T1 锚点 C2 正常回填
    assert preview["totals"]["additions_skipped"] == 0


def test_new_records_do_not_duplicate_mapping_warnings(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")],
                [("T1", "甲", 11, 1, "进行中"), ("T2", "乙", 15, 2, "进行中")])
    # 底表与反馈表第 3、4 列都叫"计划": 标题歧义按相同列位置自动对应
    book(tmp_path / "current.xlsx", [("T1", "甲", 10, 0, "进行中")],
         ("编号", "单位", "计划", "计划", "状态"))
    book(tmp_path / "feedback.xlsx",
         [("T1", "甲", 11, 1, "进行中"), ("T2", "乙", 15, 2, "进行中")],
         ("编号", "单位", "计划", "计划", "状态"))
    p["update_fields"] = [3]
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    # 每个来源字段映射只提示一次(主键+回填集一次, 新记录全字段一次), 不随记录数刷屏
    assert sum("列位置" in w for w in preview["warnings"]) == 2


def test_reordered_columns_and_explicit_mapping(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")], [("T1", "甲", 11, 1, "进行中")])
    feedback = Path(p["feedback"][0]["path"])
    book(feedback, [(8, "T1", "已完成", 14, "甲")], ("完成", "编号", "状态", "计划", "单位"))
    preview = reconcile_preview(p)
    assert preview["can_export"]
    assert next(c for c in preview["changes"] if c["field"] == 3)["after"] == 14
    book(feedback, [(8, "T1", "已完成", 14, "甲")], ("完成", "编号", "状态", "改名计划", "单位"))
    assert not reconcile_preview(p)["can_export"]
    p["feedback"][0]["field_map"] = {"3": 4}
    assert reconcile_export(p)["exported"]


def test_duplicate_header_maps_by_position_when_layout_matches(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")], [("T1", "甲", 11, 1, "进行中")])
    book(Path(p["feedback"][0]["path"]), [("T1", "甲", 11, 1, "进行中")],
         ("编号", "单位", "计划", "计划", "状态"))
    p["update_fields"] = [3]
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    assert next(c for c in preview["changes"] if c["field"] == 3)["after"] == 11
    assert any("列位置" in w for w in preview["warnings"])


def test_duplicate_header_without_position_match_still_blocks(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")], [("T1", "甲", 11, 1, "进行中")])
    book(Path(p["feedback"][0]["path"]), [("T1", "甲", 11, 1, "进行中")],
         ("编号", "单位", "完成", "计划", "计划"))
    p["update_fields"] = [3]
    preview = reconcile_preview(p)
    assert not preview["can_export"]
    p["feedback"][0]["field_map"] = {"3": 4}
    assert reconcile_preview(p)["can_export"]


def test_baseline_and_relocated_formula_semantics(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, "=C2*2", "进行中"), ("T2", "乙", 20, "=C3*2", "进行中")],
                [("T2", "乙", 25, "=C2*2", "进行中")])
    baseline = book(tmp_path / "baseline.xlsx",
                    [("T2", "乙", 18, "=C2*2", "进行中"), ("T1", "甲", 10, "=C3*2", "进行中")])
    p["baseline"] = {"path": str(baseline), "selection": selection(2)}
    p["compare_fields"] = [3, 4]
    preview = reconcile_preview(p)
    assert preview["can_export"], preview
    assert len(preview["changes"]) == 1
    assert preview["changes"][0]["before"] == 18 and preview["changes"][0]["after"] == 25
    assert reconcile_export(p)["exported"]
    wb = load_workbook(p["out"])
    assert wb.active["D3"].value == "=C3*2"
    wb.close()


def test_existing_red_fill_is_not_evidence_of_change(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")], [("T1", "甲", 10, 0, "进行中")])
    wb = load_workbook(p["path"])
    wb.active["C2"].font = Font(color="FF112233", bold=True)
    wb.active["C2"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    wb.save(p["path"])
    wb.close()
    assert reconcile_preview(p)["changes"] == []


@pytest.mark.parametrize("object_name", ["xl/embeddings/object1.bin",
                                      "xl/vbaProject.bin", "xl/slicers/slicer1.xml"])
def test_unsupported_objects_are_rejected_without_silent_loss(tmp_path, object_name):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")], [("T1", "甲", 11, 1, "进行中")])
    with ZipFile(p["path"], "a") as archive:
        archive.writestr(object_name, b"unsupported")
    result = reconcile_export(p)
    assert not result["exported"] and not result["can_export"]
    assert not Path(p["out"]).exists()


def test_no_overwrite_non_xlsx_and_cancellation(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")], [("T1", "甲", 11, 1, "进行中")])
    p["out"] = p["path"]
    assert not reconcile_export(p)["exported"]
    p["out"] = str(tmp_path / "new.xlsx")
    with pytest.raises(Cancelled):
        reconcile_export(p, cancel=lambda: True)
    assert not Path(p["out"]).exists()
    p["path"] = str(tmp_path / "unsupported.xls")
    assert not reconcile_preview(p)["can_export"]


def test_multilevel_header_paths_distinguish_same_leaf_titles(tmp_path):
    current = tmp_path / "current.xlsx"
    feedback = tmp_path / "feedback.xlsx"
    for path, planned, actual in ((current, 10, 1), (feedback, 12, 5)):
        wb = Workbook()
        ws = wb.active
        ws.append(["任务", "计划", "进展"])
        ws.append(["编号", "数量", "数量"])
        ws.append(["T1", planned, actual])
        wb.save(path)
        wb.close()
    p = {"path": str(current), "selection": {"axis": "rows", "header_start": 1, "header_end": 2,
         "data_start": 3, "data_end": 3, "field_start": 1, "field_end": 3},
         "feedback": [{"path": str(feedback)}], "key_fields": [1], "update_fields": [2, 3],
         "compare_fields": [2], "out": str(tmp_path / "result.xlsx")}
    preview = reconcile_preview(p)
    assert preview["can_export"], preview
    assert preview["fields"][1]["path"] == ["计划", "数量"]
    assert reconcile_export(p)["exported"]


def test_feedback_selection_defaults_to_own_extent(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中"), ("T2", "乙", 20, 0, "进行中")],
                [("T2", "乙", 22, 1, "进行中")])
    p["feedback"][0].pop("selection")
    assert reconcile_preview(p)["can_export"]
    assert reconcile_export(p)["exported"]


@pytest.mark.parametrize("axis", ["rows", "columns"])
def test_new_tasks_preserve_context_and_relocate_formula_for_both_axes(tmp_path, axis):
    formula_old = "=C2*2" if axis == "rows" else "=B3*2"
    formula_new = "=C3*2" if axis == "rows" else "=C3*2"
    p = payload(tmp_path, [("T1", "甲", 10, formula_old, "进行中")],
                [("T1", "甲", 10, formula_old, "进行中"), ("T2", "乙", 20, formula_new, "进行中")],
                axis=axis)
    preview = reconcile_preview(p)
    resolve(p, preview, "addition", "append")
    result = reconcile_export(p)
    assert result["exported"], result
    wb = load_workbook(p["out"])
    ws = wb.active
    assert ws.cell(3, 2).value == "乙" if axis == "rows" else ws.cell(2, 3).value == "乙"
    assert ws.cell(3, 4).value == "=C3*2" if axis == "rows" else ws.cell(4, 3).value == "=C3*2"
    wb.close()


def test_merged_unit_can_participate_in_composite_identity(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中"), ("T2", "甲", 20, 0, "进行中")],
                [("T2", "甲", 22, 1, "进行中")])
    wb = load_workbook(p["path"])
    wb.active.merge_cells("B2:B3")
    wb.save(p["path"])
    wb.close()
    p["key_fields"] = [1, 2]
    result = reconcile_export(p)
    assert result["exported"], result
    wb = load_workbook(p["out"])
    assert wb.active["C3"].value == 22 and "B2:B3" in wb.active.merged_cells
    wb.close()


def test_formula_in_permuted_fields_blocks_instead_of_summing_wrong_column(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, "=C2*2", "进行中")],
                [("T1", "甲", 10, "=C2*2", "进行中")])
    book(Path(p["feedback"][0]["path"]), [("T1", "甲", "=D2*2", 12, "进行中")],
         ("编号", "单位", "完成", "计划", "状态"))
    result = reconcile_export(p)
    assert not result["exported"] and not result["can_export"]
    assert any("公式" in i["message"] for i in result["issues"])


def test_explicit_null_mapping_disables_automatic_matching(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")], [("T1", "甲", 11, 1, "进行中")])
    p["feedback"][0]["field_map"] = {"3": None}
    preview = reconcile_preview(p)
    assert not preview["can_export"]
    assert any(i["kind"] == "field_mapping" and i.get("target_field") == 3 for i in preview["issues"])


@pytest.mark.parametrize("axis", ["rows", "columns"])
def test_second_import_reverted_to_explicit_baseline_restores_only_fill(tmp_path, axis):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")],
                [("T1", "甲", 12, 5, "进行中")], axis=axis)
    planned = "C2" if axis == "rows" else "B3"
    completed = "D2" if axis == "rows" else "B4"
    baseline_path = p["path"]
    wb = load_workbook(baseline_path)
    wb.active[planned].font = Font(name="宋体", size=13, bold=True, italic=True)
    wb.active[planned].fill = PatternFill(fill_type="solid", fgColor="FFFFCC00")
    wb.active[completed].font = Font(color="FF775577")
    wb.save(baseline_path)
    wb.close()
    p["compare_fields"] = [3]
    p["baseline"] = {"path": baseline_path, "selection": selection(1, axis=axis)}
    first = reconcile_export(p)
    assert first["exported"], first
    wb = load_workbook(first["out"])
    assert str(wb.active[planned].fill.fgColor.rgb)[-6:] == "FF0000"
    # A user changes other font attributes between monthly imports.
    font = copy(wb.active[planned].font)
    font.sz = 16
    font.strike = True
    wb.active[planned].font = font
    wb.save(first["out"])
    wb.close()

    returned = book(tmp_path / "returned.xlsx", [("T1", "甲", 10, 8, "进行中")], axis=axis)
    p["path"] = first["out"]
    p["feedback"] = [{"path": str(returned), "selection": selection(1, axis=axis)}]
    p["out"] = str(tmp_path / "second.xlsx")
    preview = reconcile_preview(p)
    assert preview["can_export"]
    assert not any(change["field"] == 3 for change in preview["changes"])
    second = reconcile_export(p)
    assert second["exported"], second
    wb = load_workbook(second["out"])
    cell = wb.active[planned]
    assert cell.value == 10
    # 值改回与基线一致 → 填充恢复基线的黄色; 字体保持用户后来改的样子
    assert cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb == "FFFFCC00"
    assert cell.font.name == "宋体" and cell.font.sz == 16
    assert cell.font.bold and cell.font.italic and cell.font.strike
    assert wb.active[completed].value == 8 and wb.active[completed].font.color.rgb == "FF775577"
    wb.close()


def test_explicit_baseline_original_red_remains_red_without_content_change(tmp_path):
    p = payload(tmp_path, [("T1", "甲", 10, 0, "进行中")], [("T1", "甲", 10, 0, "进行中")])
    wb = load_workbook(p["path"])
    wb.active["C2"].font = Font(name="宋体", bold=True)
    wb.active["C2"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    wb.save(p["path"])
    wb.close()
    p["baseline"] = {"path": p["path"], "selection": selection(1)}
    preview = reconcile_preview(p)
    assert preview["can_export"] and preview["changes"] == []
    assert reconcile_export(p)["exported"]
    wb = load_workbook(p["out"])
    assert str(wb.active["C2"].fill.fgColor.rgb)[-6:] == "FF0000" and wb.active["C2"].font.bold
    wb.close()


def _roles_payload(tmp_path, base_records):
    base = tmp_path / "base.xlsx"
    fb = tmp_path / "fb.xlsx"
    headers = ("序号", "项目名称", "单位", "金额", "状态")
    book(base, base_records, headers=headers)
    book(fb, base_records[:1], headers=headers)
    sel = {"axis": "rows", "header_start": 1, "header_end": 1, "data_start": 2,
           "data_end": 1 + len(base_records), "field_start": 1, "field_end": 5}
    return {"path": str(base), "sheet": "计划", "selection": sel,
            "feedback": [{"path": str(fb), "sheet": "计划", "selection": dict(sel)}],
            "key_fields": [], "update_fields": [], "result_form": "direct",
            "out": str(tmp_path / "o.xlsx")}


def test_suggested_roles_multi_row_group(tmp_path):
    # 多行组表: 单位列高频重复+序号组内唯一 → 预填推荐 分组=单位列, 主键=序号列
    records = [(1, "a", "甲", 10, "x"), (2, "b", "甲", 20, "x"),
               (3, "c", "乙", 30, "x"), (4, "d", "乙", 40, "x"), (5, "e", "乙", 50, "x")]
    preview = reconcile_preview(_roles_payload(tmp_path, records))
    assert preview["suggested_roles"] == {"group_field": 3, "key_fields": [1]}


def test_suggested_roles_one_row_per_group_skips_group(tmp_path):
    # 每组一行的表: 分组列没有对位价值, 只推荐序号主键
    records = [(1, "s1", "甲", 10, "x"), (2, "s2", "乙", 20, "x"),
               (3, "s3", "丙", 30, "x"), (4, "s4", "丁", 40, "x")]
    preview = reconcile_preview(_roles_payload(tmp_path, records))
    assert preview["suggested_roles"] == {"key_fields": [1]}


def test_serial_ordinal_accepts_integral_float():
    # WPS 常把整数存成 float(如 2.0): 序号判读应认
    from engine.tools.excel_merge.reconcile import _serial_ordinal
    assert _serial_ordinal(2.0) == 2
    assert _serial_ordinal(17.0) == 17
    assert _serial_ordinal(2.5) is None
    assert _serial_ordinal("例1") is None


def test_suggested_roles_extended_group_titles(tmp_path):
    # 语义词表覆盖常见业务组名(区域/供应商/门店等), 但仍须通过数据形态验证
    records = [(1, "a", "华东", 10, "x"), (2, "b", "华东", 20, "x"),
               (3, "c", "华南", 30, "x"), (4, "d", "华南", 40, "x"),
               (5, "e", "华北", 50, "x"), (6, "f", "华北", 60, "x")]
    base = tmp_path / "base.xlsx"
    fb = tmp_path / "fb.xlsx"
    headers = ("序号", "名称", "区域", "金额", "状态")
    book(base, records, headers=headers)
    book(fb, records[:1], headers=headers)
    sel = {"axis": "rows", "header_start": 1, "header_end": 1, "data_start": 2,
           "data_end": 7, "field_start": 1, "field_end": 5}
    preview = reconcile_preview({"path": str(base), "sheet": "计划", "selection": sel,
                                 "feedback": [{"path": str(fb), "sheet": "计划", "selection": dict(sel)}],
                                 "key_fields": [], "update_fields": [],
                                 "result_form": "direct", "out": str(tmp_path / "o.xlsx")})
    assert preview["suggested_roles"] == {"group_field": 3, "key_fields": [1]}


def test_suggested_roles_rejects_unique_titled_group_column(tmp_path):
    # 标题叫"单位"但内容每行唯一(不是分组形态) → 不推荐分组, 内容验证淘汰标题候选
    records = [(1, "a", "甲公司", 10, "x"), (2, "b", "乙公司", 20, "x"),
               (3, "c", "丙公司", 30, "x"), (4, "d", "丁公司", 40, "x"),
               (5, "e", "戊公司", 50, "x"), (6, "f", "己公司", 60, "x")]
    base = tmp_path / "base.xlsx"
    fb = tmp_path / "fb.xlsx"
    headers = ("序号", "名称", "单位", "金额", "状态")
    book(base, records, headers=headers)
    book(fb, records[:1], headers=headers)
    sel = {"axis": "rows", "header_start": 1, "header_end": 1, "data_start": 2,
           "data_end": 7, "field_start": 1, "field_end": 5}
    preview = reconcile_preview({"path": str(base), "sheet": "计划", "selection": sel,
                                 "feedback": [{"path": str(fb), "sheet": "计划", "selection": dict(sel)}],
                                 "key_fields": [], "update_fields": [],
                                 "result_form": "direct", "out": str(tmp_path / "o.xlsx")})
    assert preview["suggested_roles"] == {"key_fields": [1]}


def _flat_payload(tmp_path, base_records, fb_records, headers=("序号", "名称", "状态")):
    base = tmp_path / "base.xlsx"
    fb = tmp_path / "fb.xlsx"
    book(base, base_records, headers=headers)
    book(fb, fb_records, headers=headers)
    sel = {"axis": "rows", "header_start": 1, "header_end": 1, "data_start": 2,
           "data_end": 1 + len(base_records if base_records else fb_records), "field_start": 1, "field_end": len(headers)}
    fb_sel = dict(sel, data_end=1 + len(fb_records))
    return {"path": str(base), "sheet": "计划", "selection": sel,
            "feedback": [{"path": str(fb), "sheet": "计划", "selection": fb_sel}],
            "key_fields": [1], "update_fields": [3], "result_form": "direct",
            "out": str(tmp_path / "o.xlsx")}


def test_flat_missing_inside_returned_range_flagged_deleted(tmp_path):
    # 非分组: 反馈交回 1,2,4,5,6(缺3, 夹在交回区间内) → 疑似删除而非普通未反馈
    base = [(i, f"n{i}", "x") for i in range(1, 7)]
    fb = [(1, "n1", "y"), (2, "n2", "y"), (4, "n4", "y"), (5, "n5", "y"), (6, "n6", "y")]
    preview = reconcile_preview(_flat_payload(tmp_path, base, fb))
    assert preview["can_export"], preview["issues"]
    assert any("疑似被反馈方删除" in w for w in preview["warnings"])
    flagged = [m for m in preview["missing"] if "可能已被删除" in m["message"]]
    assert len(flagged) == 1 and flagged[0]["key"] == [3], flagged
    plain = [m for m in preview["missing"] if m["message"] == "未交回，疑似被删除；已保留当前内容"]
    assert not plain


def test_flat_missing_tail_stays_not_received(tmp_path):
    # 缺的是尾部(5,6 在交回区间之外): 与"没填到"无法区分, 仍报未反馈
    base = [(i, f"n{i}", "x") for i in range(1, 7)]
    fb = [(1, "n1", "y"), (2, "n2", "y"), (3, "n3", "y"), (4, "n4", "y")]
    preview = reconcile_preview(_flat_payload(tmp_path, base, fb))
    assert not any("疑似被反馈方删除" in w for w in preview["warnings"])
    assert all(m["message"] == "未交回，疑似被删除；已保留当前内容" for m in preview["missing"])


def test_flat_text_code_gap_flagged_deleted(tmp_path):
    # 文本编号 T001..T006 交回缺 T003(等长编号, 字典序=编号序) → 疑似删除
    codes = [f"T00{i}" for i in range(1, 7)]
    base = [(c, f"n{c}", "x") for c in codes]
    fb = [(c, f"n{c}", "y") for c in codes if c != "T003"]
    preview = reconcile_preview(_flat_payload(tmp_path, base, fb))
    flagged = [m for m in preview["missing"] if "可能已被删除" in m["message"]]
    assert len(flagged) == 1 and flagged[0]["key"] == ["T003"], flagged


def test_compare_output_flags_flat_deleted_row_orange(tmp_path):
    # 非分组: 序号 3 疑似删除(夹在交回区间内) → 保留对比输出中该行整行橙
    base = [(i, f"n{i}", "x") for i in range(1, 6)]
    fb = [(1, "n1", "y"), (2, "n2", "y"), (4, "n4", "y"), (5, "n5", "y")]
    p = _flat_payload(tmp_path, base, fb)
    p["result_form"] = "compare"
    result = reconcile_export(p)
    assert result["exported"], result
    wb = load_workbook(p["out"]); ws = wb["计划"]
    r_del = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == 3)
    assert _fill_hex(ws.cell(r_del, 2)) in ("F39C12", "E67E22", "FFC000", "F5A623")
    assert _fill_hex(ws.cell(r_del, 3)) in ("F39C12", "E67E22", "FFC000", "F5A623")
    wb.close()


def test_group_renumbered_deletion_located_by_content(tmp_path):
    # 序号被重排+删中间: 序号信息已丢, 但名称等非回填列没变 → 按内容对位
    # 精确点名"缺少序号 5", 其余改动落正确行(不滑位), 不再退回"疑似重排"警示
    base = [(2, "甲", "a", 10, "x"), (5, "甲", "b", 20, "x"), (9, "甲", "c", 30, "x")]
    feedback = [(1, "甲", "a", 11, "x"), (2, "甲", "c", 99, "x")]   # 删了b, 序号重排成1,2
    p = payload(tmp_path, base, feedback)
    p["key_fields"] = [1]; p["update_fields"] = [4]; p["group_field"] = 2
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    assert any("缺少序号 5" in w and "内容比对" in w for w in preview["warnings"])
    assert not any("疑似被重新编号" in w for w in preview["warnings"])
    change = next(c for c in preview["changes"] if c["after"] == 99)
    assert change["key"] == ["甲", 3]                    # c 的改动落底稿第 3 条, 不滑位


def test_group_inserted_row_becomes_addition(tmp_path):
    # 组内插行: 新行内容对不上 → 按新增处理, 原有行按内容对位, 无误删警告
    base = [(2, "甲", "a", 10, "x"), (5, "甲", "b", 20, "x"), (9, "甲", "c", 30, "x")]
    feedback = [(1, "甲", "a", 10, "x"), (2, "甲", "new", 77, "x"),
                (3, "甲", "b", 20, "x"), (4, "甲", "c", 30, "x")]   # 中间插了 new
    p = payload(tmp_path, base, feedback)
    p["key_fields"] = [1]; p["update_fields"] = [4]; p["group_field"] = 2
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    assert preview["totals"]["additions"] == 1, preview["totals"]
    assert not any("缺少序号" in w for w in preview["warnings"])
    assert any("1 行内容与该组底稿记录都对不上" in w for w in preview["warnings"])


def test_unit_value_changed_hinted_as_delete_plus_addition(tmp_path):
    # 反馈里单位被改: 乙组消失+丙组出现 → 关联提示"可能单位被改动",
    # 乙按疑似删除保留, 丙作为新记录补入(结果照常输出)
    base = [(1, "甲", "a", 10, "x"), (1, "乙", "b", 20, "x")]
    fb = tmp_path / "改单位.xlsx"
    book(fb, [(1, "甲", "a", 11, "x"), (1, "丙", "b", 99, "x")])
    p = payload(tmp_path, base, [(1, "甲", "a", 11, "x")])
    p["feedback"] = [{"path": str(fb), "sheet": "计划",
                      "selection": selection(2)}]
    p["key_fields"] = [1]; p["update_fields"] = [4]; p["group_field"] = 2
    preview = reconcile_preview(p)
    assert preview["can_export"], preview["issues"]
    assert any("单位被改动" in w for w in preview["warnings"])
    assert preview["totals"]["additions"] == 1 and preview["totals"]["missing"] == 1
    assert any("疑似被删除" in m["message"] for m in preview["missing"])


def test_marker_color_avoids_existing_red(tmp_path):
    # 底表数据区已有纯红填充: 标记红必须避让(不再是FF0000), 原有红保留,
    # 预览响应带 marker_colors 且警告说明换色
    from openpyxl.styles import PatternFill
    base = tmp_path / "base.xlsx"
    fb = tmp_path / "fb.xlsx"
    book(base, [(1, "a", "x"), (2, "b", "x"), (3, "c", "x")])
    book(fb, [(1, "a", "y"), (2, "b", "x"), (3, "c", "x")])
    wb = load_workbook(base)
    wb.active.cell(4, 3).fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    wb.save(base); wb.close()
    p = payload(tmp_path, [], [])      # 占位, 逐项重填
    p = {"path": str(base), "sheet": "计划", "selection": selection(3),
         "feedback": [{"path": str(fb), "sheet": "计划", "selection": selection(3)}],
         "key_fields": [1], "update_fields": [3], "result_form": "compare",
         "resolutions": {}, "out": str(tmp_path / "r.xlsx")}
    preview = reconcile_preview(p)
    assert preview["marker_colors"][0] != "FFFF0000", preview["marker_colors"]
    assert any("自动改用 #" in w for w in preview["warnings"])
    result = reconcile_export(p)
    assert result["exported"], result
    wb = load_workbook(p["out"]); ws = wb["计划"]
    assert _fill_hex(ws.cell(2, 3)) == "8B0000"            # 新值行=避让后的深红
    assert _fill_hex(ws.cell(3, 3)) in ("F39C12", "E67E22", "FFC000", "F5A623")  # 原值橙
    assert _fill_hex(ws.cell(5, 3)) == "FF0000"            # 底表原有红原样保留
    wb.close()


def test_mapping_survives_header_range_drift(tmp_path):
    # 反馈的表头行数识别与底稿差一行(大表常见): 完整表头路径对不上时,
    # 按"最细字段行"标题回退匹配(两侧唯一才启用), 并警告说明——不阻断
    from openpyxl import Workbook
    base = tmp_path / "base.xlsx"
    fb = tmp_path / "fb.xlsx"

    def make(path):
        wb = Workbook(); ws = wb.active; ws.title = "计划"
        ws["A1"] = "全年计划"; ws.merge_cells("A1:F1")
        for c, h in enumerate(["序号", "名称", "单位", "金额", "状态", "备注"], 1):
            ws.cell(2, c, h); ws.cell(3, c, h)
        for i in range(1, 5):
            for c, v in enumerate([i, f"名{i}", "甲", i * 10, "进行中", "r"], 1):
                ws.cell(3 + i, c, v)
        wb.save(path); wb.close()

    make(base); make(fb)
    sel = {"axis": "rows", "header_start": 2, "header_end": 3,
           "data_start": 4, "data_end": 7, "field_start": 1, "field_end": 6}
    p = {"path": str(base), "sheet": "计划", "selection": dict(sel),
         "feedback": [{"path": str(fb), "sheet": "计划",
                       "selection": dict(sel, header_end=2)}],   # 表头识别少一行
         "key_fields": [1], "update_fields": [4, 5], "result_form": "direct",
         "resolutions": {}, "out": str(tmp_path / "r.xlsx")}
    preview = reconcile_preview(p)
    assert preview["can_export"], [i["message"] for i in preview["issues"]]
    assert any("表头行数与底稿略有差异" in w for w in preview["warnings"])
    # 反向(底稿识别少一行, 反馈完整)同样可导出
    p2 = dict(p, selection=dict(sel, header_end=2),
              feedback=[{"path": str(fb), "sheet": "计划", "selection": dict(sel)}])
    assert reconcile_preview(p2)["can_export"]
    # 无漂移时不触发回退警告
    p3 = dict(p, feedback=[{"path": str(fb), "sheet": "计划", "selection": dict(sel)}])
    rep3 = reconcile_preview(p3)
    assert rep3["can_export"] and not any("表头行数" in w for w in rep3["warnings"])


def test_split_meta_absolute_mapping_and_fallback(tmp_path):
    # 拆分产物回填: 带溯源属性 → 按绝对列号精确对应(重名/漂移免疫);
    # 表头被改 → 快照不符点名警告并回退标题匹配
    import re as _re
    from zipfile import ZipFile as ZF2
    from engine.tools.excel_merge.native_split import split_export
    src = tmp_path / "src.xlsx"
    headers = ("序号", "序号", "单位", "金额", "状态")     # 双列同名序号
    wb2 = Workbook(); ws2 = wb2.active; ws2.title = "计划"
    ws2.append(list(headers))
    rows = [(1, 1, "甲", 10, "x"), (2, 2, "甲", 20, "x"), (3, 3, "乙", 30, "x"), (4, 4, "乙", 40, "x")]
    for r in rows:
        ws2.append(list(r))
    wb2.save(src); wb2.close()
    sel = {"axis": "rows", "header_start": 1, "header_end": 1, "data_start": 2,
           "data_end": 5, "field_start": 1, "field_end": 5}
    split = split_export({"path": str(src), "sheet": "计划", "selection": sel,
                          "group_fields": [{"kind": "field", "index": 3}],
                          "group_mode": "combination", "blank_group": "separate",
                          "filters": [], "resolutions": {}, "out_dir": str(tmp_path / "out"),
                          "include_example_rows": False, "retain_all_sheets": False})
    fb = Path(split["files"][0]["path"])
    # 反馈不带 selection(引擎 fallback 自适应)——重名列靠 meta 而非标题
    p = {"path": str(src), "sheet": "计划", "selection": dict(sel),
         "feedback": [{"path": str(fb)}],
         "key_fields": [1], "update_fields": [4], "result_form": "direct",
         "resolutions": {}, "out": str(tmp_path / "r.xlsx")}
    preview = reconcile_preview(p)
    assert preview["can_export"], [i["message"] for i in preview["issues"]]
    assert any("按导出时的列对应表精确对应" in w for w in preview["warnings"])
    # 改表头: 快照不符 → 点名回退
    wb3 = load_workbook(fb); wb3.active.cell(1, 4).value = "被改的金额"; wb3.save(fb); wb3.close()
    preview2 = reconcile_preview(p)
    assert any("第 4 列 的表头与导出时不一致" in w for w in preview2["warnings"])


def test_field_mappings_echoed_for_ui(tmp_path):
    # 预览响应携带每份反馈的全字段映射(含重名列按位置对上的)——
    # 界面字段对应表据此回显, 重名列不再误显示"不对应"
    headers = ("序号", "序号", "单位", "金额", "状态")     # 两列同名"序号"
    records = [(1, 1, "甲", 10, "x"), (2, 2, "甲", 20, "x"), (3, 3, "乙", 30, "x")]
    base = tmp_path / "base.xlsx"; fb = tmp_path / "fb.xlsx"
    book(base, records, headers=headers)
    book(fb, records[:2], headers=headers)
    p = {"path": str(base), "sheet": "计划", "selection": selection(3),
         "feedback": [{"path": str(fb), "sheet": "计划", "selection": selection(2)}],
         "key_fields": [1], "update_fields": [4], "group_field": 3,
         "result_form": "direct", "resolutions": {}, "out": str(tmp_path / "o.xlsx")}
    preview = reconcile_preview(p)
    assert preview["can_export"], [i["message"] for i in preview["issues"]]
    assert len(preview["field_mappings"]) == 1
    mapping = preview["field_mappings"][0]["mapping"]
    assert mapping.get("1") == 1 and mapping.get("2") == 2   # JSON 键为字符串; 重名列按位置对上并回显
    assert len(mapping) >= 5


def test_reconcile_output_carries_column_meta(tmp_path):
    # 回填导出的结果文件同样写入结构溯源属性——下一轮以它为底稿/反馈时,
    # 识别与列对应按本轮结构走(重名列大表不重新失配)
    base = [(1, "a", "甲", 10, "x"), (2, "b", "甲", 20, "x"), (3, "c", "乙", 30, "x")]
    p = payload(tmp_path, base, base[:2])
    p["key_fields"] = [1]; p["update_fields"] = [4]; p["group_field"] = 3
    p["out"] = str(tmp_path / "round1.xlsx")
    result = reconcile_export(p)
    assert result["exported"], result
    from engine.tools.excel_merge.preview import _split_meta
    meta = _split_meta(p["out"])
    assert meta and meta["columns"].get("1") == "编号" and len(meta["columns"]) == 5
    # 以第一轮结果为底稿, 其副本作反馈: meta 绝对对应, 零 mapping issue
    import shutil
    fb2 = tmp_path / "round2_fb.xlsx"
    shutil.copy(p["out"], fb2)
    p2 = dict(p, path=p["out"], feedback=[{"path": str(fb2), "sheet": "计划", "selection": selection(3)}],
              out=str(tmp_path / "round2.xlsx"))
    preview2 = reconcile_preview(p2)
    assert preview2["can_export"], [i["message"] for i in preview2["issues"]]
    assert any("按导出时的列对应表精确对应" in w for w in preview2["warnings"])
