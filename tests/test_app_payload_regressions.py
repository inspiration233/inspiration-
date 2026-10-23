# -*- coding: utf-8 -*-
"""App 真实 payload 形态回归: 每张工作表自动附带回显 selection(origin=auto)。

历史bug: apply_selection 在auto模式下用"原始标签键"(同名不去重)覆盖读取器的
去重键, 重名表头按位对齐被破坏(非模板源几乎全部列丢失); 合并单元格坐标又被
按转置语义改写, 吞掉真实数据格。本文件锁住 App 形态与引擎直调结果一致。
"""
import sys
from pathlib import Path

import pytest

from conftest import load

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from engine.tools.excel_merge.merge import Options, merge  # noqa: E402
from engine.tools.excel_merge.preview import preview  # noqa: E402
from engine.tools.excel_merge.selection import apply_selection  # noqa: E402


def _app_payload(paths, options):
    """按 App buildPayload 的形态构造: 每张可见工作表一条 auto 选区。"""
    selections = []
    for p in paths:
        report = preview({"path": str(p)})
        for sheet in report["sheets"]:
            rep = preview({"path": str(p), "sheet": sheet["name"]})
            item = {"path": str(p), "sheet": sheet["name"],
                    "selection_origin": rep.get("selection") and "auto" or "auto"}
            if str(p).lower().endswith(".xlsx"):
                item["selection"] = rep["selection"]
            selections.append(item)
    return {"paths": [str(p) for p in paths], "options": options, "selections": selections}


def test_auto_echoed_selection_matches_direct_engine(tmp_path, out_path):
    """同名表头(XXX×18)+交叉表: App 形态(带auto回显选区)与直调引擎逐格一致。"""
    src = Path("测试集/特殊示例")
    f1 = src / "_诊断_修复后_特殊示例_1.xlsx"
    f2 = src / "_诊断_修复后_特殊示例_2.xlsx"
    if not f1.exists():    # 兼容文件在工作根目录的布置
        f1, f2 = Path("_诊断_修复后_特殊示例_1.xlsx"), Path("_诊断_修复后_特殊示例_2.xlsx")
    options = {"mode": "group", "align": "loose", "renumber_serial": True,
               "strip_serial_rows": True, "min_fill_ratio": 0.2}
    payload = _app_payload([f1, f2], options)
    assert payload["selections"], "App 形态必须带上每张表的回显选区"

    with_app = tmp_path / "app.xlsx"
    direct = tmp_path / "direct.xlsx"
    r1 = merge(payload["paths"], with_app, Options(**options,
                                                   selections=payload["selections"]))
    r2 = merge(payload["paths"], direct, Options(**options))

    wa, wb_ = load(with_app), load(direct)
    assert wa.sheetnames == wb_.sheetnames
    for name in wa.sheetnames:
        a, b = wa[name], wb_[name]
        cells = sum(1 for r in range(1, max(a.max_row, b.max_row) + 1)
                    for c in range(1, max(a.max_column, b.max_column) + 1)
                    if (a.cell(r, c).value or "") != (b.cell(r, c).value or ""))
        assert cells == 0, f"{name}: App形态与直调不一致 ({cells} 格)"
    # 非模板源的字段不能丢: 例4 行应有满列数据
    ws = wa["总表"]
    row13 = sum(1 for c in range(1, ws.max_column + 1)
                if ws.cell(13, c).value not in (None, ""))
    assert row13 >= 15, f"例4 行仅 {row13} 格, 重名表头对齐又失效了"


def test_manual_selection_on_native_cross_rejected(tmp_path):
    """横向配置表不支持手动框选: 明确报错而不是按转置语义吞数据。"""
    from engine.tools.excel_merge import reader as rd
    f = Path("测试集/特殊示例/_诊断_修复后_特殊示例_1.xlsx")
    if not f.exists():
        f = Path("_诊断_修复后_特殊示例_1.xlsx")
    tables, r = rd.list_tables(str(f), False)
    src = next(s for s, _ in tables if s.sheet_name == "车辆技术配置要求")
    try:
        apply_selection(src, {"selection_origin": "manual",
                              "selection": {"axis": "rows", "header_start": 4,
                                            "header_end": 4, "data_start": 5,
                                            "data_end": 20, "field_start": 1,
                                            "field_end": 7}})
        assert False, "原生交叉表手动框选应被拒绝"
    except ValueError as e:
        assert "横向配置表" in str(e)
    finally:
        r.close()


def test_duplicate_headers_align_by_position_with_auto_selection(tmp_path, out_path):
    """重名表头(全部同名)两文件: auto回显选区不得破坏按位对齐。"""
    from openpyxl import Workbook
    f1, f2 = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    for path, tag in ((f1, "甲"), (f2, "乙")):
        wb = Workbook()
        ws = wb.active
        ws.title = "明细"
        ws.append(["编号", "数值", "数值", "数值", "数值"])
        for i in range(1, 21):
            ws.append([i, f"{tag}{i}", i * 2, f"{tag}x{i}", i * 3])
        wb.save(path)
        wb.close()
    payload = _app_payload([f1, f2], {"mode": "group", "align": "loose"})
    r = merge(payload["paths"], out_path, Options(**payload["options"],
                                                  selections=payload["selections"]))
    ws = load(out_path)["明细"]
    assert ws.max_row == 41                       # 表头1行 + 甲20行 + 乙20行
    values = [ws.cell(7, 2).value, ws.cell(7, 3).value, ws.cell(7, 5).value]
    assert values == ["甲6", 12, 18]               # 第6条数据: 每列都按位对上, 无串列
    b_rows = [ws.cell(r, 2).value for r in range(22, 42)]
    assert b_rows[:3] == ["乙1", "乙2", "乙3"]     # 乙的行数据完整进入


def test_example_rows_keep_first_source_only(tmp_path, out_path):
    """"例"开头序号的行=模板示例: 只保留首份来源那组, 后续来源跳过并提示。"""
    from openpyxl import Workbook

    def book(path, examples, tag):
        wb = Workbook()
        ws = wb.active
        ws.title = "总表"
        ws.append(["序号", "名称", "数量"])
        for i, name in enumerate(examples, 1):
            ws.append([f"例{name}", f"{tag}示例{i}", i])
        for i in range(1, 6):
            ws.append([i, f"{tag}数据{i}", i * 10])
        wb.save(path)
        wb.close()

    f1, f2 = tmp_path / "甲.xlsx", tmp_path / "乙.xlsx"
    book(f1, ["1", "2"], "甲")
    book(f2, ["3", "4"], "乙")
    r = merge([str(f1), str(f2)], out_path, Options(mode="group", renumber_serial=True))
    ws = load(out_path)["总表"]
    vals = [r0[0] for r0 in ws.iter_rows(min_row=2, max_col=1, values_only=True)]
    assert vals == ["例1", "例2", 1, 2, 3, 4, 5, 6, 7, 8, 9, 10]  # 甲示例+甲5行+乙5行, 无例3/例4
    assert any("示例行" in w and "3 行" not in w or "示例行" in w for w in r["warnings"])
    assert any("乙" in w and "示例行" in w for w in r["warnings"])
    # 单文件: 自己的示例行保留, 无跳过提示
    out2 = tmp_path / "single.xlsx"
    r2 = merge([str(f2)], out2, Options(mode="group", renumber_serial=True))
    vals2 = [r0[0] for r0 in load(out2)["总表"].iter_rows(min_row=2, max_col=1, values_only=True)]
    assert vals2 == ["例3", "例4", 1, 2, 3, 4, 5]
    assert not any("示例行" in w for w in r2["warnings"])
    # 首份没有示例行 → 第二份的示例成为"首组"保留
    f3 = tmp_path / "丙.xlsx"
    book(f3, [], "丙")
    out3 = tmp_path / "second.xlsx"
    merge([str(f3), str(f2)], out3, Options(mode="group", renumber_serial=True))
    vals3 = [r0[0] for r0 in load(out3)["总表"].iter_rows(min_row=2, max_col=1, values_only=True)]
    assert vals3 == [1, 2, 3, 4, 5, "例3", "例4", 6, 7, 8, 9, 10]


def test_serial_formulas_renumbered_not_broken_by_row_moves(tmp_path, out_path):
    """=ROW()-N 公式序号: 行挪到更靠上位置后重排为连续数字, 不再出现负数序号。"""
    from openpyxl import Workbook
    from conftest import load as _load
    from engine.tools.excel_merge.merge import merge, Options
    f1, f2 = tmp_path / "短前置.xlsx", tmp_path / "长前置.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "总表"
    ws["A1"] = "表"; ws.append(["序号", "名称", "数量", "备注"])
    ws.append(["=ROW()-2", "甲1", 1, "a"]); ws.append(["=ROW()-2", "甲2", 2, "b"])
    wb.save(f1); wb.close()
    wb = Workbook(); ws = wb.active; ws.title = "总表"
    for i, text in enumerate(["表", "说明一", "说明二", "说明三", "说明四"], start=1):
        ws.cell(i, 1, text)
    for c, h in enumerate(["序号", "名称", "数量", "备注"], start=1):
        ws.cell(6, c, h)
    ws.cell(7, 1, "=ROW()-6"); ws.cell(7, 2, "乙1"); ws.cell(7, 3, 3); ws.cell(7, 4, "c")
    ws.cell(8, 1, "=ROW()-6"); ws.cell(8, 2, "乙2"); ws.cell(8, 3, 4); ws.cell(8, 4, "d")
    wb.save(f2); wb.close()
    merge([str(f1), str(f2)], out_path, Options(mode="single", renumber_serial=True,
                                                min_fill_ratio=0.2))
    ws = _load(out_path)["总表"]
    serials = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)
               if ws.cell(r, 1).value not in (None, "")]
    assert serials == ["表", "序号", 1, 2, 3, 4], serials   # 公式序号重排为连续数字
