"""Generated packages must be structurally valid even when openpyxl tolerates them."""
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from engine.tools.excel_merge.package_check import validate_xlsx
from engine.tools.excel_merge import workbench


MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RID = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"


def source(path):
    wb = Workbook(); ws = wb.active
    ws.append(["编号", "数量"]); ws.append(["A", 3]); ws.append(["B", 4])
    ws["A1"].font = Font(bold=True)
    ws["A2"].hyperlink = "https://example.com/reference"
    wb.save(path); wb.close()
    return path


def edit_xml(path, part, edit):
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = ET.fromstring(parts[part]); edit(root)
    parts[part] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, blob in parts.items():
            archive.writestr(name, blob)


def bad_font(path):
    edit_xml(path, "xl/styles.xml", lambda root:
             root.find(f"{{{MAIN}}}cellXfs/{{{MAIN}}}xf").set("fontId", "999"))


def test_valid_package_with_external_link_and_unknown_extension_is_allowed(tmp_path):
    path = source(tmp_path / "正常.xlsx")
    def extension(root):
        ext = ET.SubElement(ET.SubElement(root, f"{{{MAIN}}}extLst"), f"{{{MAIN}}}ext", uri="unknown")
        ET.SubElement(ext, "{urn:example:future}customFeature")
    edit_xml(path, "xl/workbook.xml", extension)
    validate_xlsx(path)


@pytest.mark.parametrize("kind", ["font", "row", "view", "relationship", "target"])
def test_structural_reference_errors_are_detected_without_reading_cell_contents(tmp_path, kind):
    path = source(tmp_path / "损坏.xlsx")
    if kind == "font":
        bad_font(path)
    elif kind == "row":
        edit_xml(path, "xl/worksheets/sheet1.xml", lambda root:
                 root.find(f"{{{MAIN}}}sheetData/{{{MAIN}}}row").set("s", "999"))
    elif kind == "view":
        edit_xml(path, "xl/worksheets/sheet1.xml", lambda root:
                 root.find(f"{{{MAIN}}}sheetViews/{{{MAIN}}}sheetView").set("workbookViewId", "7"))
    elif kind == "relationship":
        edit_xml(path, "xl/worksheets/sheet1.xml", lambda root:
                 ET.SubElement(root, f"{{{MAIN}}}pageSetup", {f"{{{RID}}}id": "rIdMissing"}))
    else:
        edit_xml(path, "xl/worksheets/_rels/sheet1.xml.rels", lambda root:
                 ET.SubElement(root, f"{{{REL}}}Relationship", Id="rIdMissing", Type=RID + "/printerSettings",
                               Target="../printerSettings/missing.bin"))
    if kind in ("font", "row", "view"):
        # This was the previous acceptance check; it misses these corruptions.
        book = load_workbook(path, read_only=True)
        assert book.active["B2"].value == 3
        book.close()
    with pytest.raises(ValueError):
        validate_xlsx(path)


def test_failed_output_integrity_does_not_publish_a_result(tmp_path, monkeypatch):
    path = source(tmp_path / "原始.xlsx")
    original = path.read_bytes()
    real_merge = workbench.merge
    def corrupted_merge(paths, out, options, **kwargs):
        report = real_merge(paths, out, options, **kwargs)
        bad_font(out)
        return report
    monkeypatch.setattr(workbench, "merge", corrupted_merge)
    payload = {"paths": [str(path)], "out": str(tmp_path / "结果.xlsx")}
    preview = workbench.merge_preview(payload)
    assert not preview["can_export"]
    assert any(issue["id"] == "output-package" for issue in preview["issues"])
    with pytest.raises(ValueError, match="结构检查未通过"):
        workbench.merge_export(payload)
    assert not (tmp_path / "结果.xlsx").exists()
    assert path.read_bytes() == original
