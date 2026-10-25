# -*- coding: utf-8 -*-
"""拆分增强: 图片/图表保留+平移, WPS内嵌图片部件回注, 示例行, 数值筛选, 保留全部工作表。"""
import io
import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from conftest import load

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from engine.tools.excel_merge.native_split import split_preview, split_export  # noqa: E402


def _book(path, rows=12):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PILImage
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["序号", "分组", "数值"])
    for i in range(1, rows):
        ws.append([i, "甲" if i % 2 else "乙", i * 10])
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=rows - 1), titles_from_data=True)
    ws.add_chart(chart, "E3")                      # 行3(乙) → 甲组删, 乙组留
    buf = io.BytesIO()
    PILImage.new("RGB", (12, 12), "green").save(buf, "PNG")
    ws.add_image(XLImage(io.BytesIO(buf.getvalue())), "E9")   # 行9(乙) → 乙组留
    wb.save(path)
    wb.close()


def _selection(payload_sheet="数据"):
    return {"axis": "rows", "header_start": 1, "header_end": 1,
            "data_start": 2, "data_end": 11, "field_start": 1, "field_end": 3}


def test_images_and_charts_preserved_and_shifted(tmp_path):
    src = tmp_path / "源.xlsx"
    _book(src)
    payload = {"path": str(src), "sheet": "数据", "selection": _selection(),
               "group_fields": [{"kind": "field", "index": 2}],
               "group_mode": "combination", "blank_group": "separate",
               "filters": [], "resolutions": {}, "out_dir": str(tmp_path / "out"),
               "include_example_rows": False, "retain_all_sheets": False}
    report = split_export(payload)
    by_label = {f["label"]: Path(f["path"]) for f in report["files"]}
    assert set(by_label) == {"甲", "乙"}
    # 乙组: 行3,5,7,9,11 保留 → 映射 3→2,5→3,7→4,9→5,11→6
    wb = load(by_label["乙"])
    ws = wb["数据"]
    assert len(ws._charts) == 1 and len(ws._images) == 1
    assert ws._charts[0].anchor._from.row == 1            # E3 → E2
    assert ws._images[0].anchor._from.row == 4            # E9 → E5(0基=4)
    assert ws._charts[0].series[0].val.numRef.f.endswith("$C$2:$C$6")
    assert [ws.cell(r, 2).value for r in range(2, 7)] == ["乙"] * 5
    wb.close()
    # 甲组: 锚点行(3/9)全被删 → 图表与图片丢弃, 仅警告不阻断
    wb_a = load(by_label["甲"])
    assert not wb_a["数据"]._charts and not wb_a["数据"]._images
    wb_a.close()
    assert any("图片/图表" in w for w in report["warnings"])


def test_example_rows_carried_into_every_output(tmp_path):
    from openpyxl import Workbook
    src = tmp_path / "源.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["序号", "分组", "数量"])
    ws.append(["例1", "甲", "示例"])
    ws.append(["例2", "乙", "示例"])
    for i in range(1, 7):
        ws.append([i, "甲" if i % 2 else "乙", i])
    wb.save(src)
    wb.close()
    payload = {"path": str(src), "sheet": "数据", "selection": _selection(),
               "group_fields": [{"kind": "field", "index": 2}],
               "group_mode": "combination", "blank_group": "separate",
               "filters": [], "resolutions": {}, "out_dir": str(tmp_path / "out"),
               "include_example_rows": True, "retain_all_sheets": False}
    report = split_export(payload)
    assert any("示例行" in w for w in report["warnings"])
    for item in report["files"]:
        wb = load(item["path"])
        ws = wb["数据"]
        serials = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
        assert serials[:2] == ["例1", "例2"], item["label"]
        wb.close()
    # 关闭开关: 示例行按普通记录分组, 只出现在对应组
    payload2 = dict(payload, out_dir=str(tmp_path / "out2"), include_example_rows=False)
    report2 = split_export(payload2)
    got = {}
    for item in report2["files"]:
        wb = load(item["path"])
        ws = wb["数据"]
        got[item["label"]] = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
        wb.close()
    assert "例1" in got["甲"] and "例2" in got["乙"]


def test_numeric_and_oneof_filters(tmp_path):
    from openpyxl import Workbook
    src = tmp_path / "源.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["序号", "分组", "数值"])
    for i in range(1, 11):
        ws.append([i, "甲" if i % 2 else "乙", i])
    wb.save(src)
    wb.close()
    base = {"path": str(src), "sheet": "数据", "selection": _selection(),
            "group_fields": [], "group_mode": "combination", "blank_group": "separate",
            "resolutions": {}, "include_example_rows": False, "retain_all_sheets": False}
    def run(filters, out):
        payload = dict(base, filters=filters, out_dir=str(tmp_path / out))
        rep = split_export(payload)
        wb2 = load(rep["files"][0]["path"])
        vals = [wb2["数据"].cell(r, 3).value for r in range(2, wb2["数据"].max_row + 1)]
        wb2.close()
        return vals
    assert run([{"field": 3, "op": "gt", "value": "7"}], "o1") == [8, 9, 10]
    assert run([{"field": 3, "op": "le", "value": 3}], "o2") == [1, 2, 3]
    assert run([{"field": 2, "op": "oneof", "value": "甲、丙"}], "o3") == [1, 3, 5, 7, 9]
    # 非数字比较值 → 阻断问题而不是误判
    rep = split_preview(dict(base, filters=[{"field": 3, "op": "gt", "value": "abc"}]))
    assert not rep["can_export"] and any("数字" in i["message"] for i in rep["issues"])


def test_retain_all_sheets(tmp_path):
    from openpyxl import Workbook
    src = tmp_path / "源.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["序号", "分组"])
    for i in range(1, 7):
        ws.append([i, "甲" if i % 2 else "乙"])
    other = wb.create_sheet("说明页")
    other["A1"] = "附注"
    wb.save(src)
    wb.close()
    payload = {"path": str(src), "sheet": "数据",
               "selection": dict(_selection(), data_end=7),
               "group_fields": [{"kind": "field", "index": 2}],
               "group_mode": "combination", "blank_group": "separate",
               "filters": [], "resolutions": {}, "out_dir": str(tmp_path / "out"),
               "include_example_rows": False, "retain_all_sheets": True}
    report = split_export(payload)
    assert len(report["files"]) == 2
    for item in report["files"]:
        wb = load(item["path"])
        assert wb.sheetnames == ["数据", "说明页"]
        wb.close()


def _wps_source(path):
    """构造带 WPS cellimages 部件与 DISPIMG 公式的源文件(部件手术)。"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["序号", "分组", "图片"])
    for i in range(1, 7):
        ws.append([i, "甲" if i % 2 else "乙", '=DISPIMG("ID_1",1)' if i == 1 else None])
    buf = io.BytesIO()
    wb.save(buf)
    cellimages = (b'<etc:cellImages xmlns:etc="urn:wps-office:etc" '
                  b'xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing">'
                  b'<etc:cellImage><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="1" name="ID_1"/>'
                  b'</xdr:nvPicPr></xdr:pic></etc:cellImage></etc:cellImages>')
    rels = (b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            b'<Relationship Id="rIdCi1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image99.png"/></Relationships>')
    wrels_add = (b'<Relationship Id="rIdCellImages" '
                 b'Type="http://schemas.microsoft.com/office/2007/relationships/xl/cellimages" '
                 b'Target="cellimages.xml"/>')
    ct_add = (b'<Override PartName="/xl/cellimages.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>')
    with ZipFile(io.BytesIO(buf.getvalue())) as zin, ZipFile(path, "w", ZIP_DEFLATED) as zout:
        for name in zin.namelist():
            data = zin.read(name)
            if name == "[Content_Types].xml":
                data = data.replace(b"</Types>", ct_add + b"</Types>")
            elif name == "xl/_rels/workbook.xml.rels":
                data = data.replace(b"</Relationships>", wrels_add + b"</Relationships>")
            zout.writestr(name, data)
        zout.writestr("xl/cellimages.xml", cellimages)
        zout.writestr("xl/_rels/cellimages.xml.rels", rels)
        zout.writestr("xl/media/image99.png", b"\x89PNG-fake-bytes")


def test_wps_cellimages_reinjected(tmp_path):
    src = tmp_path / "源.xlsx"
    _wps_source(src)
    payload = {"path": str(src), "sheet": "数据",
               "selection": dict(_selection(), data_end=7),
               "group_fields": [{"kind": "field", "index": 2}],
               "group_mode": "combination", "blank_group": "separate",
               "filters": [], "resolutions": {}, "out_dir": str(tmp_path / "out"),
               "include_example_rows": False, "retain_all_sheets": False}
    rep = split_export(payload)
    assert any("DISPIMG" in w for w in rep["warnings"])
    by_label = {f["label"]: f["path"] for f in rep["files"]}
    assert set(by_label) == {"甲", "乙"}
    for label, out in by_label.items():
        with ZipFile(out) as z:      # 部件在每张结果里都完整(与分组无关)
            names = z.namelist()
            assert "xl/cellimages.xml" in names
            assert "xl/_rels/cellimages.xml.rels" in names
            assert "xl/media/image99.png" in names
            assert b"rIdCellImages" in z.read("xl/_rels/workbook.xml.rels")
            assert b"/xl/cellimages.xml" in z.read("[Content_Types].xml")
    wb = load(by_label["甲"])        # DISPIMG 公式行(i=1)属于甲组
    disp = [c.value for row in wb["数据"].iter_rows() for c in row
            if isinstance(c.value, str) and "DISPIMG" in c.value]
    wb.close()
    assert disp == ['=DISPIMG("ID_1",1)']


def test_cross_record_formula_falls_back_to_cached_value(tmp_path):
    """引用其它记录的公式在拆分后引用删空: 有缓存值→自动转原计算值不再阻断。"""
    from openpyxl import Workbook
    src = tmp_path / "源.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws.append(["单位", "数量", "引用"])
    ws.append(["甲", 2, "=B3"])          # 引用乙的行
    ws.append(["乙", 3, "=B2"])
    wb.save(src)
    wb.close()
    # openpyxl 不写缓存值: 注入 <v> 模拟真实 Excel/WPS 文件
    from zipfile import ZipFile as ZF
    patched = tmp_path / "带缓存.xlsx"
    with ZF(src) as zin, ZF(patched, "w") as zout:
        for name in zin.namelist():
            data = zin.read(name)
            if name == "xl/worksheets/sheet1.xml":
                import re as _re
                data = (_re.sub(rb"<f>B3</f>(<v\s*/>|<v></v>)?", b"<f>B3</f><v>30</v>", data)
                        if b"<f>B3</f>" in data else data)
                data = (_re.sub(rb"<f>B2</f>(<v\s*/>|<v></v>)?", b"<f>B2</f><v>20</v>", data)
                        if b"<f>B2</f>" in data else data)
            zout.writestr(name, data)
    payload = {"path": str(patched), "sheet": "Sheet",
               "selection": {"axis": "rows", "header_start": 1, "header_end": 1,
                             "data_start": 2, "data_end": 3, "field_start": 1, "field_end": 3},
               "group_fields": [{"kind": "field", "index": 1}],
               "group_mode": "combination", "blank_group": "separate",
               "filters": [], "resolutions": {}, "out_dir": str(tmp_path / "out"),
               "include_example_rows": False, "retain_all_sheets": False}
    report = split_export(payload)
    assert report["can_export"]
    # openpyxl 生成的文件无缓存值: 公式原样保留(引用被移除的也会保留为公式文本)
    assert any("移除" in str(i.get("message","")) for i in report["issues"]) or any("保留" in w for w in report["warnings"])
    by_label = {f["label"]: Path(f["path"]) for f in report["files"]}
    out = load(by_label["甲"])
    assert out["Sheet"]["C2"].value == 30 and not isinstance(out["Sheet"]["C2"].value, str)
    out.close()
    out = load(by_label["乙"])
    assert out["Sheet"]["C2"].value == 20
    out.close()
    # 无缓存值的文件(openpyxl直接生成): 公式原样保留, 不再阻断
    payload2 = dict(payload, path=str(src), out_dir=str(tmp_path / "out2"))
    rep2 = split_preview(payload2)
    assert rep2["can_export"]
    assert any(x["code"] == "missing_reference" and not x.get("resolved") for x in rep2["issues"])


def test_dependency_issue_not_duplicated_and_formula_stays_live(tmp_path):
    """公式引用其它可见页: 只弹一条可选择的依赖问题; 选择保留后公式保持公式。"""
    from openpyxl import Workbook
    src = tmp_path / "源.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["单位", "数量", "引用"])
    ws.append(["甲", 1, "=说明页!A1"])
    ws.append(["乙", 2, "=说明页!A1"])
    other = wb.create_sheet("说明页")
    other["A1"] = 100
    wb.save(src)
    wb.close()
    payload = {"path": str(src), "sheet": "数据",
               "selection": {"axis": "rows", "header_start": 1, "header_end": 1,
                             "data_start": 2, "data_end": 3, "field_start": 1, "field_end": 3},
               "group_fields": [{"kind": "field", "index": 1}],
               "group_mode": "combination", "blank_group": "separate",
               "filters": [], "resolutions": {}, "out_dir": str(tmp_path / "out"),
               "include_example_rows": False, "retain_all_sheets": False}
    preview = split_preview(payload)
    codes = [x["code"] for x in preview["issues"]]
    assert codes.count("dependency_sheet") == 1
    assert "missing_sheet" not in codes          # 不再重复弹无选项的同源阻断项
    assert not preview["can_export"]
    # 选择"允许保留"后: 公式保持活性(不转值), 两页都在结果里
    payload["retain_sheets"] = ["说明页"]
    result = split_export(payload)
    assert result["can_export"]
    out = load(result["files"][0]["path"])
    assert out.sheetnames == ["数据", "说明页"]
    cell = out["数据"]["C2"]
    assert isinstance(cell.value, str) and cell.value.startswith("=说明页!")
    out.close()


def test_serial_renumbered_from_one_per_group(tmp_path):
    """开启 renumber_serial 时序号每组从1连续重排; 文本序号保留。"""
    from openpyxl import Workbook
    from zipfile import ZipFile as ZF
    src0 = tmp_path / "s0.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "数据"
    ws.append(["序号", "单位", "数量"])
    for i, u in enumerate(["甲", "乙", "乙", "乙"], start=2):
        ws.append([f"=ROW()-1", u, i])
    wb.save(src0); wb.close()
    src = tmp_path / "带缓存.xlsx"          # 注入缓存值 1,2,3,4
    with ZF(src0) as zin, ZF(src, "w") as zout:
        for name in zin.namelist():
            data = zin.read(name)
            if name == "xl/worksheets/sheet1.xml":
                import re as _re
                for i in range(2, 6):
                    data = _re.sub((r"<f>ROW\(\)-1</f>(<v\s*/>|<v></v>)?").encode(),
                                   f"<f>ROW()-1</f><v>{i - 1}</v>".encode(), data, count=1)
            zout.writestr(name, data)
    # 逐格精确注入(上面一次性替换不精确, 重新做)
    with ZF(src0) as zin, ZF(src, "w") as zout:
        for name in zin.namelist():
            data = zin.read(name)
            if name == "xl/worksheets/sheet1.xml":
                text = data.decode()
                for i in range(2, 6):
                    text = text.replace(f'<c r="A{i}"><f>ROW()-1</f><v></v></c>',
                                        f'<c r="A{i}"><f>ROW()-1</f><v>{i - 1}</v></c>', 1)
                data = text.encode()
            zout.writestr(name, data)
    payload = {"path": str(src), "sheet": "数据",
               "selection": {"axis": "rows", "header_start": 1, "header_end": 1,
                             "data_start": 2, "data_end": 5, "field_start": 1, "field_end": 3},
               "group_fields": [{"kind": "field", "index": 2}],
               "group_mode": "combination", "blank_group": "separate",
               "filters": [], "resolutions": {}, "out_dir": str(tmp_path / "out"),
               "include_example_rows": False, "retain_all_sheets": False,
               "renumber_serial": True}
    report = split_export(payload)
    assert any("连续重排" in w for w in report["warnings"])
    by_label = {f["label"]: Path(f["path"]) for f in report["files"]}
    out = load(by_label["乙"])
    serials = [out["数据"].cell(r, 1).value for r in range(2, 5)]
    assert serials == [1, 2, 3], serials      # 公式序号按本组从1重排, 无空洞无负数
    out.close()
    out = load(by_label["甲"])
    assert out["数据"]["A2"].value == 1
    out.close()


def test_serial_kept_original_by_default(tmp_path):
    """默认不重排: 数值序号保留底稿原值, 公式序号转为拆分前的显示值——回收后按原序号回填对位。"""
    from openpyxl import Workbook
    from zipfile import ZipFile as ZF
    src0 = tmp_path / "k0.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "数据"
    ws.append(["序号", "单位", "数量"])
    for i, u in enumerate(["甲", "乙", "乙", "乙"], start=2):
        ws.append([f"=ROW()-1", u, i])
    wb.save(src0); wb.close()
    src = tmp_path / "带缓存.xlsx"          # 注入缓存值 1,2,3,4
    with ZF(src0) as zin, ZF(src, "w") as zout:
        for name in zin.namelist():
            data = zin.read(name)
            if name == "xl/worksheets/sheet1.xml":
                text = data.decode()
                for i in range(2, 6):
                    text = text.replace(f'<c r="A{i}"><f>ROW()-1</f><v></v></c>',
                                        f'<c r="A{i}"><f>ROW()-1</f><v>{i - 1}</v></c>', 1)
                data = text.encode()
            zout.writestr(name, data)
    payload = {"path": str(src), "sheet": "数据",
               "selection": {"axis": "rows", "header_start": 1, "header_end": 1,
                             "data_start": 2, "data_end": 5, "field_start": 1, "field_end": 3},
               "group_fields": [{"kind": "field", "index": 2}],
               "group_mode": "combination", "blank_group": "separate",
               "filters": [], "resolutions": {}, "out_dir": str(tmp_path / "out"),
               "include_example_rows": False, "retain_all_sheets": False}
    report = split_export(payload)
    assert not any("连续重排" in w for w in report["warnings"])
    assert any("转为拆分前的显示值" in w for w in report["warnings"])
    by_label = {f["label"]: Path(f["path"]) for f in report["files"]}
    out = load(by_label["乙"])
    serials = [out["数据"].cell(r, 1).value for r in range(2, 5)]
    assert serials == [2, 3, 4], serials    # 底稿原序号原样保留, 回填按值即能对位
    out.close()
    out = load(by_label["甲"])
    assert out["数据"]["A2"].value == 1
    out.close()


def test_horizontal_serial_formula_frozen_to_cached(tmp_path):
    """横向表对称: 序号沿记录列方向的公式(=COLUMN()-N)拆分后同样漂移, 默认转为拆分前显示值。"""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from zipfile import ZipFile as ZF
    src0 = tmp_path / "h0.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "数据"
    ws["A1"], ws["A2"], ws["A3"] = "序号", "单位", "数量"
    for c, unit in enumerate(["甲", "乙", "乙", "乙"], start=2):
        ws.cell(1, c).value = "=COLUMN()-1"
        ws.cell(2, c).value = unit
        ws.cell(3, c).value = c * 10
    wb.save(src0); wb.close()
    src = tmp_path / "带缓存.xlsx"          # 注入缓存值 1..4
    with ZF(src0) as zin, ZF(src, "w") as zout:
        for name in zin.namelist():
            data = zin.read(name)
            if name == "xl/worksheets/sheet1.xml":
                text = data.decode()
                for c in range(2, 6):
                    col = get_column_letter(c)
                    text = text.replace(f'<c r="{col}1"><f>COLUMN()-1</f><v></v></c>',
                                        f'<c r="{col}1"><f>COLUMN()-1</f><v>{c - 1}</v></c>', 1)
                data = text.encode()
            zout.writestr(name, data)
    payload = {"path": str(src), "sheet": "数据",
               "selection": {"axis": "columns", "header_start": 1, "header_end": 1,
                             "data_start": 2, "data_end": 5, "field_start": 1, "field_end": 3},
               "group_fields": [{"kind": "field", "index": 2}],
               "group_mode": "combination", "blank_group": "separate",
               "filters": [], "resolutions": {}, "out_dir": str(tmp_path / "out"),
               "include_example_rows": False, "retain_all_sheets": False}
    report = split_export(payload)
    assert any("转为拆分前的显示值" in w for w in report["warnings"])
    by_label = {f["label"]: Path(f["path"]) for f in report["files"]}
    out = load(by_label["乙"])
    serials = [out["数据"].cell(1, c).value for c in range(2, 5)]
    assert serials == [2, 3, 4], serials    # 底稿原序号原样保留
    out.close()


def test_split_export_writes_column_meta(tmp_path):
    """拆分产物写入 TOOLBOX_SPLIT 溯源属性: 列号→表头快照+结构区间。"""
    import json
    import re as _re
    from zipfile import ZipFile as ZF2
    src = tmp_path / "m.xlsx"
    _book_into(src)
    payload = {"path": str(src), "sheet": "数据",
               "selection": {"axis": "rows", "header_start": 1, "header_end": 1,
                             "data_start": 2, "data_end": 11, "field_start": 1, "field_end": 3},
               "group_fields": [{"kind": "field", "index": 2}],
               "group_mode": "combination", "blank_group": "separate",
               "filters": [], "resolutions": {}, "out_dir": str(tmp_path / "out"),
               "include_example_rows": False, "retain_all_sheets": False}
    report = split_export(payload)
    out = report["files"][0]["path"]
    with ZF2(out) as z:
        assert "docProps/custom.xml" in z.namelist()
        raw = z.read("docProps/custom.xml").decode("utf-8")
    m = _re.search(r'name="TOOLBOX_SPLIT"[^>]*><vt:lpwstr[^>]*>(.*?)</vt:lpwstr>', raw, _re.S)
    meta = json.loads(m.group(1))
    assert meta["columns"]["1"] == "序号" and meta["columns"]["2"] == "分组"
    assert meta["selection"]["header_start"] == 1 and meta["sheet"] == "数据"


def _book_into(path):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "数据"
    ws.append(["序号", "分组", "数值"])
    for i in range(1, 11):
        ws.append([i, "甲" if i % 2 else "乙", i * 10])
    wb.save(path); wb.close()


def test_csv_split_end_to_end(tmp_path):
    """CSV 拆分: GBK 编码识别→拆分→产物为xlsx且带溯源属性(base_file=原始csv名)。"""
    from openpyxl import load_workbook
    from zipfile import ZipFile as ZF3
    import re as _re3
    import json as _json3
    from engine.tools.excel_merge.preview import preview
    from engine.tools.excel_merge.native_split import split_export
    src = tmp_path / "单位数据.csv"
    rows = ["序号,单位名称,金额,状态"] + [f"{i},单位{i % 3},{i * 10},进行中" for i in range(1, 10)]
    src.write_text("\r\n".join(rows), encoding="gbk", newline="")
    report = preview({"path": str(src)})
    assert report["fields"][0]["label"] == "序号" and report["fields"][1]["label"] == "单位名称"
    rep = split_export({"path": str(src), "sheet": "数据", "selection": report["selection"],
                        "group_fields": [{"kind": "field", "index": 2}],
                        "group_mode": "combination", "blank_group": "separate",
                        "filters": [], "resolutions": {}, "out_dir": str(tmp_path / "out"),
                        "include_example_rows": False, "retain_all_sheets": False})
    assert len(rep["files"]) == 3
    out = Path(rep["files"][0]["path"])
    assert out.suffix == ".xlsx"
    wb = load_workbook(out)
    data = [(wb.active.cell(r, 1).value, wb.active.cell(r, 2).value) for r in range(1, wb.active.max_row + 1)]
    assert data[0] == ("序号", "单位名称") and len([d for d in data[1:] if d[0]]) == 3
    wb.close()
    with ZF3(out) as z:
        raw = z.read("docProps/custom.xml").decode("utf-8")
    meta = _json3.loads(_re3.search(r'name="TOOLBOX_SPLIT"[^>]*><vt:lpwstr[^>]*>(.*?)</vt:lpwstr>',
                                    raw, _re3.S).group(1))
    assert meta["base_file"] == "单位数据.csv" and meta["columns"]["2"] == "单位名称"
    assert any("CSV" in w for w in rep["warnings"])


def test_xls_split_preserves_values_and_merges(tmp_path):
    """XLS 拆分(第一步): 值/合并单元格/数字正确转入, 产物为xlsx带属性, 警告说明样式尽力。"""
    import xlwt
    from openpyxl import load_workbook
    from engine.tools.excel_merge.preview import preview
    from engine.tools.excel_merge.native_split import split_export
    src = tmp_path / "legacy.xls"
    book = xlwt.Workbook()
    sh = book.add_sheet("计划")
    heads = ["序号", "单位", "金额"]
    for c, h in enumerate(heads):
        sh.write(0, c, h)
    sh.write(1, 0, 1); sh.write(1, 1, "甲"); sh.write(1, 2, 10.5)
    sh.write(2, 0, 2); sh.write(2, 1, "甲"); sh.write(2, 2, 20.5)
    sh.write(3, 0, 3); sh.write(3, 1, "乙"); sh.write(3, 2, 30.5)
    sh.write_merge(4, 5, 1, 1, "合并区")          # 纵向合并(数据区内)
    book.save(str(src))
    report = preview({"path": str(src)})
    assert report["fields"][1]["label"] == "单位"
    rep = split_export({"path": str(src), "sheet": "计划", "selection": report["selection"],
                        "group_fields": [{"kind": "field", "index": 2}],
                        "group_mode": "combination", "blank_group": "separate",
                        "filters": [], "resolutions": {}, "out_dir": str(tmp_path / "out"),
                        "include_example_rows": False, "retain_all_sheets": False})
    assert len(rep["files"]) >= 2 and any("XLS" in w for w in rep["warnings"])
    for f in rep["files"]:
        assert Path(f["path"]).suffix == ".xlsx"
    wb = load_workbook(rep["files"][0]["path"])
    values = [wb.active.cell(r, 1).value for r in range(1, wb.active.max_row + 1)]
    assert 1 in values and values[0] == "序号"
    wb.close()
    # 合并区随所属组("合并区"值成组)进入对应产物
    found = False
    for f in rep["files"]:
        w = load_workbook(f["path"])
        if any(str(m).startswith("B") for m in w.active.merged_cells.ranges):
            found = True
        w.close()
    assert found
