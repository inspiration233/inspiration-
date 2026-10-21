"""OOXML package references that an openpyxl reload alone does not validate."""
from copy import deepcopy
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.views import BookView
from openpyxl.worksheet.views import Pane, Selection, SheetView, SheetViewList

from engine.tools.excel_merge.native_book import merge_workbook


NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


@pytest.mark.parametrize("pane_state", [None, "split", "frozen", "frozenSplit"])
def test_native_export_rehomes_first_sheet_view_and_preserves_its_settings(tmp_path, pane_state):
    source = tmp_path / "multiple_windows.xlsx"
    output = tmp_path / "combined.xlsx"
    book = Workbook()
    book.views = [BookView(), BookView(), BookView()]
    sheet = book.active
    sheet.title = "Source"
    sheet.append(["ID", "Value", "Double"])
    sheet.append(["X001", 42, "=B2*2"])
    pane = None
    if pane_state:
        pane = Pane(
            xSplit=2000 if pane_state == "split" else 2,
            ySplit=1200 if pane_state == "split" else 6,
            topLeftCell="C7", activePane="bottomRight", state=pane_state,
        )
    first = SheetView(
        workbookViewId=2, zoomScale=85, zoomScaleNormal=95,
        showGridLines=False, showRowColHeaders=False, topLeftCell="C7",
        pane=pane,
        selection=[Selection(pane="bottomRight" if pane else None,
                             activeCell="F9", sqref="F9:G10")],
    )
    sheet.views = SheetViewList(sheetView=[first, SheetView(workbookViewId=0, zoomScale=110)])
    book.save(source)
    book.close()

    original = load_workbook(source)
    expected = deepcopy(original.active.sheet_view)
    expected.workbookViewId = 0
    original.close()
    report = merge_workbook(
        {"paths": [str(source)], "selections": [{"path": str(source), "sheet": "Source"}]},
        output,
    )
    assert report["can_export"], report

    with ZipFile(output) as archive:
        assert archive.testzip() is None
        workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
        view_count = len(workbook_xml.find("s:bookViews", NS))
        assert view_count == 1
        for name in archive.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                worksheet_xml = ET.fromstring(archive.read(name))
                views = worksheet_xml.find("s:sheetViews", NS)
                assert len(views) == 1
                assert all(0 <= int(view.get("workbookViewId")) < view_count for view in views)

    actual = load_workbook(output)
    try:
        assert actual.active.sheet_view == expected
        assert len(actual.active.views.sheetView) == 1
        assert list(actual.active.values) == [("ID", "Value", "Double"), ("X001", 42, "=B2*2")]
    finally:
        actual.close()
