"""Automatic source extents follow file changes; manual crops remain intentional."""
import pytest
from openpyxl import Workbook, load_workbook

from engine.tools.excel_merge.preview import preview
from engine.tools.excel_merge.workbench import merge_export


@pytest.mark.parametrize("origin", ["auto", "manual"])
def test_added_records_after_source_preview_follow_selection_origin(tmp_path, origin):
    path = tmp_path / "来源.xlsx"
    wb = Workbook(); ws = wb.active
    ws.append(["客户", "数量", "金额"])
    ws.append(["甲", 2, 20]); ws.append(["乙", 3, 30])
    wb.save(path); wb.close()
    initial = preview({"path": str(path)})
    wb = load_workbook(path)
    wb.active.append(["新增", 4, 40])
    wb.save(path); wb.close()
    spec = {"path": str(path), "sheet": initial["sheet"],
            "selection": initial["selection"], "selection_origin": origin}
    refreshed = preview(spec)
    assert refreshed["selection"]["data_end"] == (4 if origin == "auto" else 3)
    out = tmp_path / "结果.xlsx"
    report = merge_export({"paths": [str(path)], "selections": [spec], "out": str(out)})
    assert report["can_export"], report
    result = load_workbook(out)
    try:
        actual = list(result.active.values)
        expected = [("客户", "数量", "金额"), ("甲", 2, 20), ("乙", 3, 30)]
        if origin == "auto":
            expected.append(("新增", 4, 40))
        assert actual == expected
    finally:
        result.close()


def test_auto_range_shrinks_after_records_are_removed(tmp_path):
    path = tmp_path / "来源.xlsx"
    wb = Workbook(); ws = wb.active
    ws.append(["客户", "数量", "金额"])
    for index in range(5):
        ws.append([f"任务{index}", index + 1, (index + 1) * 10])
    wb.save(path); wb.close()
    initial = preview({"path": str(path)})
    wb = load_workbook(path)
    wb.active.delete_rows(4, 3)
    wb.save(path); wb.close()
    spec = {"path": str(path), "sheet": initial["sheet"],
            "selection": initial["selection"], "selection_origin": "auto"}
    out = tmp_path / "结果.xlsx"
    report = merge_export({"paths": [str(path)], "selections": [spec], "out": str(out)})
    assert report["can_export"], report
    result = load_workbook(out)
    try:
        assert list(result.active.values) == [("客户", "数量", "金额"), ("任务0", 1, 10), ("任务1", 2, 20)]
    finally:
        result.close()
