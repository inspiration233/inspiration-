"""Cross-workbook native-copy regression tests with independently built XLSX."""
from pathlib import Path
import warnings
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, Color
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.table import TableStyle, TableStyleElement
from openpyxl.cell.text import InlineFont
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.writer.theme import theme_xml

from engine.tools.excel_merge.merge import Cancelled
from engine.tools.excel_merge.native_book import merge_workbook
from engine.tools.excel_merge import native_book


def fixture(path, amount=10, color="FF123456", with_table=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws["A1"] = "保留的标题"
    ws.merge_cells("A1:D1")
    ws["A3"] = "编号"
    ws["B3"] = "数量"
    ws["A4"] = "R1"
    ws["B4"] = "='数据页'!B2"
    ws["C4"] = "=Rate"
    ws["D4"] = "=LocalRate"
    ws["B4"].font = Font(name="宋体", bold=True, italic=True, color=color, size=14)
    ws["B4"].fill = PatternFill("solid", fgColor="FFABCDEF")
    ws["B4"].number_format = "#,##0.00"
    ws["B4"].alignment = Alignment(horizontal="center", wrap_text=True)
    ws["B4"].border = Border(bottom=Side(style="thin", color="FF000000"))
    ws.column_dimensions["B"].width = 31
    ws.row_dimensions[4].height = 28
    ws.freeze_panes = "B4"
    ws.print_title_rows = "1:3"
    ws.print_area = "A1:D4"
    ws.page_setup.orientation = "landscape"
    ws.page_margins.left = 0.31
    ws.oddHeader.center.text = "公司 &P"
    validation = DataValidation(type="list", formula1="Rate")
    validation.add("B4:B9")
    ws.add_data_validation(validation)
    ws.conditional_formatting.add("B4:B9", FormulaRule(
        formula=["B4>Rate"], fill=PatternFill("solid", fgColor="FFC6EFCE")))
    ws["E4"].hyperlink = "#'数据页'!A1"
    ws["E4"] = "跳转"
    ws.defined_names.add(DefinedName("LocalRate", attr_text="'数据页'!$B$2", localSheetId=0))
    data = wb.create_sheet("数据页")
    data.append(["编号", "数量"])
    data.append(["R1", amount])
    data.sheet_state = "hidden"
    wb.defined_names.add(DefinedName("Rate", attr_text="'数据页'!$B$2"))
    if with_table:
        table = Table(displayName="Items", ref="A1:B2")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        data.add_table(table)
        ws["F4"] = "=SUM(Items[数量])"
    wb.create_sheet("无关页")["A1"] = "不应带入"
    wb.save(path)
    wb.close()
    return path


def test_independent_native_pages_dependency_names_and_styles(tmp_path):
    first = fixture(tmp_path / "first.xlsx", 10, "FF123456", with_table=True)
    second = fixture(tmp_path / "second.xlsx", 20, "FFAB1234", with_table=True)
    original = {p: p.read_bytes() for p in (first, second)}
    out = tmp_path / "merged.xlsx"
    report = merge_workbook({"paths": [str(first), str(second)], "selections": [
        {"path": str(first), "sheet": "汇总", "data_start": 4, "data_end": 4},
        {"path": str(second), "sheet": "汇总"}]}, out)
    assert report["can_export"], report
    assert report["totals"]["sheets_merged"] == 4 and report["totals"]["dependency_sheets"] == 2
    assert all(p.read_bytes() == content for p, content in original.items())
    wb = load_workbook(out)
    assert wb.sheetnames == ["汇总", "数据页", "汇总_2", "数据页_2"]
    ws = wb["汇总_2"]
    assert ws["A1"].value == "保留的标题" and "A1:D1" in ws.merged_cells
    assert ws["B4"].value == "='数据页_2'!B2"
    assert ws["C4"].value == "=Rate_2" and ws["D4"].value == "=LocalRate"
    assert ws["F4"].value == "=SUM(Items_2[数量])"
    assert ws["E4"].hyperlink.target == "#'数据页_2'!A1"
    assert wb.defined_names["Rate_2"].attr_text == "'数据页_2'!$B$2"
    assert ws.defined_names["LocalRate"].attr_text == "'数据页_2'!$B$2"
    assert list(wb["数据页_2"].tables) == ["Items_2"]
    assert wb["数据页_2"].sheet_state == "hidden"
    assert ws["B4"].font.color.rgb == "FFAB1234" and wb["汇总"]["B4"].font.color.rgb == "FF123456"
    assert ws["B4"].font.name == "宋体" and ws["B4"].font.bold and ws["B4"].font.italic
    assert ws["B4"].fill.fgColor.rgb == "FFABCDEF" and ws["B4"].number_format == "#,##0.00"
    assert ws["B4"].alignment.wrap_text and ws["B4"].border.bottom.style == "thin"
    assert ws.row_dimensions[4].height == 28 and ws.column_dimensions["B"].width == 31
    assert ws.freeze_panes == "B4" and ws.print_title_rows == "$1:$3"
    assert ws.print_area == "'汇总_2'!$A$1:$D$4" and ws.page_margins.left == 0.31
    assert ws.oddHeader.center.text == "公司 &P"
    assert ws.data_validations.dataValidation[0].formula1 == "Rate_2"
    rule = ws.conditional_formatting[next(iter(ws.conditional_formatting))][0]
    assert rule.formula == ["B4>Rate_2"] and rule.dxf.fill.fgColor.rgb == "FFC6EFCE"
    wb.close()


def test_transitive_and_three_dimensional_dependencies_preserve_order(tmp_path):
    paths = []
    for n in (1, 2):
        path = tmp_path / f"book{n}.xlsx"
        wb = Workbook()
        wb.active.title = "开始"
        wb.active["A1"] = "=SUM('一:三'!A1)"
        for name in ("一", "二", "三"):
            wb.create_sheet(name)["A1"] = n
        wb["三"]["B1"] = "='尾页'!A1"
        wb.create_sheet("尾页")["A1"] = n * 10
        wb.create_sheet("无关")
        wb.save(path)
        wb.close()
        paths.append(path)
    out = tmp_path / "result.xlsx"
    report = merge_workbook({"paths": list(map(str, paths)), "selections": [
        {"path": str(p), "sheet": "开始"} for p in paths]}, out)
    assert report["can_export"], report
    wb = load_workbook(out)
    assert wb.sheetnames == ["开始", "一", "二", "三", "尾页", "开始_2", "一_2", "二_2", "三_2", "尾页_2"]
    assert wb["开始_2"]["A1"].value == "=SUM('一_2:三_2'!A1)"
    assert wb["三_2"]["B1"].value == "='尾页_2'!A1"
    wb.close()


def test_source_theme_colors_and_fonts_materialized_per_workbook(tmp_path):
    paths = []
    for number, rgb, face in ((1, "AA1122", "Arial"), (2, "11BB33", "Courier New")):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "主题样式"
        ws["A1"].font = Font(name="Calibri", scheme="minor", color=Color(theme=4))
        ws["A1"].fill = PatternFill("solid", fgColor=Color(theme=4))
        ws["B1"] = CellRichText("普通", TextBlock(InlineFont(rFont="Calibri", scheme="minor",
                                                         color=Color(theme=4), b=True), "加粗"))
        custom = theme_xml.replace('<a:srgbClr val="4F81BD"/>', f'<a:srgbClr val="{rgb}"/>')
        custom = custom.replace('<a:latin typeface="Calibri"/>', f'<a:latin typeface="{face}"/>')
        wb.loaded_theme = custom.encode("utf-8")
        path = tmp_path / f"theme{number}.xlsx"
        wb.save(path)
        wb.close()
        paths.append(path)
    out = tmp_path / "theme.xlsx"
    report = merge_workbook({"paths": list(map(str, paths))}, out)
    assert report["can_export"], report
    wb = load_workbook(out, rich_text=True)
    assert wb.worksheets[0]["A1"].font.color.rgb == "FFAA1122"
    assert wb.worksheets[1]["A1"].fill.fgColor.rgb == "FF11BB33"
    assert wb.worksheets[0]["A1"].font.name == "Arial"
    assert wb.worksheets[1]["A1"].font.name == "Courier New"
    assert wb.worksheets[1]["B1"].value[1].font.color.rgb == "FF11BB33"
    assert wb.worksheets[1]["B1"].value[1].font.rFont == "Courier New"
    wb.close()


def test_name_collision_between_global_name_and_table_isolated(tmp_path):
    first = fixture(tmp_path / "first.xlsx", with_table=True)
    wb = Workbook()
    wb.active["A1"] = "=Items"
    wb.defined_names.add(DefinedName("Items", attr_text="12"))
    second = tmp_path / "second.xlsx"
    wb.save(second)
    wb.close()
    out = tmp_path / "result.xlsx"
    report = merge_workbook({"paths": [str(first), str(second)]}, out)
    assert report["can_export"], report
    wb = load_workbook(out)
    assert wb["Sheet"]["A1"].value == "=Items_2" and wb.defined_names["Items_2"].attr_text == "12"
    wb.close()


@pytest.mark.parametrize("unsafe", ["xl/embeddings/object.bin", "xl/vbaProject.bin",
                                  "xl/externalLinks/externalLink1.xml", "xl/connections.xml",
                                  "xl/pivotTables/pivotTable1.xml", "xl/richData/rdrichvalue.xml"])
def test_unsupported_objects_block_before_write(tmp_path, unsafe):
    source = fixture(tmp_path / "input.xlsx")
    with ZipFile(source, "a") as archive:
        archive.writestr(unsafe, b"unsupported")
    out = tmp_path / "result.xlsx"
    report = merge_workbook({"paths": [str(source)]}, out)
    assert not report["can_export"] and report["issues"]
    assert not out.exists()


@pytest.mark.parametrize("formula", ["=INDIRECT(\"'数据页'!A1\")", "='不存在'!A1",
                                   "='[outside.xlsx]Sheet1'!A1"])
def test_unprovable_or_broken_dependencies_block(tmp_path, formula):
    source = fixture(tmp_path / "input.xlsx")
    wb = load_workbook(source)
    wb["汇总"]["G4"] = formula
    wb.save(source)
    wb.close()
    out = tmp_path / "result.xlsx"
    report = merge_workbook({"paths": [str(source)], "selections": [{"path": str(source), "sheet": "汇总"}]}, out)
    assert not report["can_export"] and not out.exists()


def test_duplicate_input_is_not_duplicated_and_cancellation_is_respected(tmp_path):
    source = fixture(tmp_path / "input.xlsx")
    out = tmp_path / "result.xlsx"
    p = {"paths": [str(source), str(source)]}
    with pytest.raises(Cancelled):
        merge_workbook(p, out, cancel=lambda: True)
    assert not out.exists()
    report = merge_workbook(p, out)
    assert report["can_export"] and report["totals"]["sheets_merged"] == 3
    assert len(report["files"]) == 1


def test_input_not_overwritten_and_nonxlsx_rejected(tmp_path):
    source = fixture(tmp_path / "input.xlsx")
    original = source.read_bytes()
    assert not merge_workbook({"paths": [str(source)]}, source)["can_export"]
    assert source.read_bytes() == original
    report = merge_workbook({"paths": [str(tmp_path / "old.xls")]}, tmp_path / "out.xlsx")
    assert not report["can_export"] and report["issues"]


def test_table_differential_style_indices_are_registered_per_source(tmp_path):
    sources = []
    for n, rgb in ((1, "FFFF0000"), (2, "FF0000FF")):
        path = fixture(tmp_path / f"input{n}.xlsx", with_table=True)
        wb = load_workbook(path)
        index = wb._differential_styles.add(DifferentialStyle(
            fill=PatternFill("solid", fgColor=rgb)))
        wb["数据页"].tables["Items"].dataDxfId = index
        wb.save(path)
        wb.close()
        sources.append(str(path))
    out = tmp_path / "output.xlsx"
    report = merge_workbook({"paths": sources}, out)
    assert report["can_export"], report
    wb = load_workbook(out)
    first = wb["数据页"].tables["Items"].dataDxfId
    second = wb["数据页_2"].tables["Items_2"].dataDxfId
    assert wb._differential_styles.styles[first].fill.fgColor.rgb == "FFFF0000"
    assert wb._differential_styles.styles[second].fill.fgColor.rgb == "FF0000FF"
    wb.close()


def test_custom_theme_with_excel_table_warns_and_exports_contents(tmp_path):
    source = fixture(tmp_path / "input.xlsx", with_table=True)
    wb = load_workbook(source)
    wb.loaded_theme = theme_xml.replace("4F81BD", "AA1122").encode("utf-8")
    wb.save(source)
    wb.close()
    out = tmp_path / "output.xlsx"
    report = merge_workbook({"paths": [str(source)]}, out)
    assert report["can_export"] and out.exists(), report
    assert any("主题" in message and "配色" in message for message in report["warnings"])
    assert_basic_cells(out)
    wb = load_workbook(out)
    assert wb["数据页"].tables["Items"].tableStyleInfo.name == "TableStyleMedium2"
    assert wb["汇总"]["F4"].value == "=SUM(Items[数量])"
    wb.close()


def rewrite_package(path, updates):
    """Modify only specified parts while retaining a valid real XLSX package."""
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    for name, update in updates.items():
        entries[name] = update(entries.get(name)) if callable(update) else update
    with ZipFile(path, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)


def extension_xml(raw, feature=None,
                  feature_namespace="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"):
    root = ET.fromstring(raw)
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    extension_list = ET.SubElement(root, f"{{{ns}}}extLst")
    if feature:
        extension = ET.SubElement(extension_list, f"{{{ns}}}ext", uri="{78C0D931-6437-407D-A8EE-F0AAD7539E65}")
        node = ET.SubElement(extension, f"{{{feature_namespace}}}{feature}")
        if feature == "calcFeatures":
            ET.SubElement(node, f"{{{feature_namespace}}}calcFeature", name="microsoft.com:RD")
        elif feature == "absPath":
            node.set("url", "C:\\Shared\\")
    return ET.tostring(root, encoding="utf-8")


def assert_basic_cells(path):
    assert path.exists()
    wb = load_workbook(path)
    ws = wb["汇总"]
    assert ws["A1"].value == "保留的标题" and "A1:D1" in ws.merged_cells
    assert ws["A4"].value == "R1" and wb["数据页"]["B2"].value == 10
    assert ws["B4"].value == "='数据页'!B2" and ws["C4"].value == "=Rate"
    assert ws["B4"].font.name == "宋体" and ws["B4"].font.bold and ws["B4"].font.italic
    assert ws["B4"].font.color.rgb == "FF123456" and ws["B4"].font.sz == 14
    assert ws["B4"].fill.fgColor.rgb == "FFABCDEF" and ws["B4"].number_format == "#,##0.00"
    assert ws.freeze_panes == "B4" and ws.column_dimensions["B"].width == 31
    assert ws.page_setup.orientation == "landscape" and ws.page_margins.left == 0.31
    assert ws.data_validations.dataValidation[0].formula1 == "Rate"
    rule = ws.conditional_formatting[next(iter(ws.conditional_formatting))][0]
    assert rule.formula == ["B4>Rate"] and rule.dxf.fill.fgColor.rgb == "FFC6EFCE"
    wb.close()


def test_styles_extension_warns_and_really_exports_xlsx(tmp_path):
    source = fixture(tmp_path / "input.xlsx")
    rewrite_package(source, {"xl/styles.xml": lambda raw: extension_xml(raw, "dxfs")})
    before = source.read_bytes()
    out = tmp_path / "output.xlsx"
    report = merge_workbook({"paths": [str(source)]}, out)
    assert report["can_export"] and not report["issues"], report
    assert any("扩展样式" in message for message in report["warnings"])
    assert source.read_bytes() == before
    assert_basic_cells(out)


@pytest.mark.parametrize("part", ["xl/styles.xml", "xl/worksheets/sheet1.xml"])
def test_empty_extension_container_exports_without_extension_notice(tmp_path, part):
    source = fixture(tmp_path / "input.xlsx")
    rewrite_package(source, {part: extension_xml})
    out = tmp_path / "output.xlsx"
    report = merge_workbook({"paths": [str(source)]}, out)
    assert report["can_export"], report
    assert not any("扩展" in message or "extension" in message.lower() for message in report["warnings"])
    assert_basic_cells(out)


def test_worksheet_formatting_extension_warns_and_exports(tmp_path):
    source = fixture(tmp_path / "input.xlsx")
    rewrite_package(source, {"xl/worksheets/sheet1.xml":
                            lambda raw: extension_xml(raw, "conditionalFormattings")})
    out = tmp_path / "output.xlsx"
    report = merge_workbook({"paths": [str(source)]}, out)
    assert report["can_export"], report
    assert any("扩展条件格式" in message for message in report["warnings"])
    assert any("读取信息" in message and "UserWarning" in message for message in report["warnings"])
    assert_basic_cells(out)


def test_unknown_worksheet_extension_warns_and_really_exports(tmp_path):
    source = fixture(tmp_path / "input.xlsx")
    rewrite_package(source, {"xl/worksheets/sheet1.xml":
                            lambda raw: extension_xml(raw, "businessRecords")})
    out = tmp_path / "output.xlsx"
    report = merge_workbook({"paths": [str(source)]}, out)
    assert report["can_export"] and not report["issues"], report
    assert any("businessRecords" in message for message in report["warnings"])
    assert_basic_cells(out)


@pytest.mark.parametrize("feature,namespace", [
    ("calcFeatures", "http://schemas.microsoft.com/office/spreadsheetml/2018/calcfeatures"),
    ("absPath", "http://schemas.microsoft.com/office/spreadsheetml/2010/11/main"),
    ("vendorMetadata", "urn:vendor:spreadsheet-extension"),
])
def test_workbook_extensions_warn_and_really_export(tmp_path, feature, namespace):
    source = fixture(tmp_path / "input.xlsx")
    rewrite_package(source, {"xl/workbook.xml": lambda raw: extension_xml(raw, feature, namespace)})
    before = source.read_bytes()
    out = tmp_path / "output.xlsx"
    report = merge_workbook({"paths": [str(source)]}, out)
    assert report["can_export"] and not report["issues"], report
    assert any(feature in message and "input.xlsx" in message for message in report["warnings"])
    assert source.read_bytes() == before
    assert_basic_cells(out)


def test_utf16_workbook_extension_does_not_lose_warning(tmp_path):
    source = fixture(tmp_path / "input.xlsx")

    def encoded_extension(raw):
        root = ET.fromstring(extension_xml(raw, "vendorMetadata", "urn:vendor:extension"))
        return ET.tostring(root, encoding="utf-16", xml_declaration=True)

    rewrite_package(source, {"xl/workbook.xml": encoded_extension})
    out = tmp_path / "output.xlsx"
    report = merge_workbook({"paths": [str(source)]}, out)
    assert report["can_export"] and not report["issues"], report
    assert any("vendorMetadata" in message for message in report["warnings"])
    assert_basic_cells(out)


def test_malformed_workbook_xml_still_blocks_before_export(tmp_path):
    source = fixture(tmp_path / "input.xlsx")
    rewrite_package(source, {"xl/workbook.xml": b"<workbook><extLst>"})
    out = tmp_path / "output.xlsx"
    report = merge_workbook({"paths": [str(source)]}, out)
    assert not report["can_export"] and report["issues"]
    assert not out.exists()


def test_attached_chart_warns_and_exports_cells_without_chart(tmp_path):
    source = fixture(tmp_path / "input.xlsx")
    wb = load_workbook(source)
    chart = BarChart()
    chart.add_data(Reference(wb["数据页"], min_col=2, min_row=1, max_row=2), titles_from_data=True)
    wb["汇总"].add_chart(chart, "H2")
    wb.save(source)
    wb.close()
    out = tmp_path / "output.xlsx"
    report = merge_workbook({"paths": [str(source)]}, out)
    assert report["can_export"], report
    assert any("图表" in message and "未保留" in message for message in report["warnings"])
    assert_basic_cells(out)
    wb = load_workbook(out)
    assert not wb["汇总"]._charts
    wb.close()
    with ZipFile(out) as archive:
        assert not any(name.startswith(("xl/drawings/", "xl/charts/")) for name in archive.namelist())


def test_printer_settings_warn_and_leave_no_dangling_page_relationship(tmp_path):
    source = fixture(tmp_path / "input.xlsx")
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

    def with_page_id(raw):
        root = ET.fromstring(raw)
        root.find(f"{{{ns}}}pageSetup").set(f"{{{rel_ns}}}id", "rIdPrinter")
        return ET.tostring(root, encoding="utf-8")

    def with_printer_relationship(raw):
        root = ET.fromstring(raw)
        ET.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                      Id="rIdPrinter", Type=rel_ns + "/printerSettings",
                      Target="../printerSettings/printerSettings1.bin")
        return ET.tostring(root, encoding="utf-8")

    rewrite_package(source, {
        "xl/worksheets/sheet1.xml": with_page_id,
        "xl/worksheets/_rels/sheet1.xml.rels": with_printer_relationship,
        "xl/printerSettings/printerSettings1.bin": b"device-specific printer settings",
    })
    with ZipFile(source) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        assert root.find(f"{{{ns}}}pageSetup").get(f"{{{rel_ns}}}id") == "rIdPrinter"
    out = tmp_path / "output.xlsx"
    report = merge_workbook({"paths": [str(source)]}, out)
    assert report["can_export"], report
    assert any("打印机专用配置未保留" in message for message in report["warnings"])
    assert_basic_cells(out)
    wb = load_workbook(out)
    assert wb["汇总"].page_setup.id is None
    wb.close()
    with ZipFile(out) as archive:
        assert not any(name.startswith("xl/printerSettings/") for name in archive.namelist())


def test_custom_table_style_warns_and_retains_cells_and_table(tmp_path):
    source = fixture(tmp_path / "input.xlsx", with_table=True)
    wb = load_workbook(source)
    index = wb._differential_styles.add(DifferentialStyle(fill=PatternFill("solid", fgColor="FF00FF00")))
    wb._table_styles.tableStyle.append(TableStyle(
        name="TeamStyle", tableStyleElement=[TableStyleElement(type="wholeTable", dxfId=index)]))
    wb["数据页"].tables["Items"].tableStyleInfo.name = "TeamStyle"
    wb.save(source)
    wb.close()
    out = tmp_path / "output.xlsx"
    report = merge_workbook({"paths": [str(source)]}, out)
    assert report["can_export"], report
    assert any("自定义表样式未保留" in message for message in report["warnings"])
    assert_basic_cells(out)
    wb = load_workbook(out)
    assert wb["数据页"].tables["Items"].ref == "A1:B2"
    assert wb["数据页"].tables["Items"].tableStyleInfo is None
    assert wb["汇总"]["F4"].value == "=SUM(Items[数量])"
    wb.close()


def test_reader_warning_is_reported_and_real_reader_exception_still_blocks(tmp_path, monkeypatch):
    source = fixture(tmp_path / "input.xlsx")
    out = tmp_path / "output.xlsx"

    def warned_reader(*args, **kwargs):
        warnings.warn("附加外观信息未读取", UserWarning)
        return load_workbook(*args, **kwargs)

    monkeypatch.setattr(native_book, "load_workbook", warned_reader)
    report = merge_workbook({"paths": [str(source)]}, out)
    assert report["can_export"], report
    assert any("读取信息（UserWarning）" in message and "附加外观信息未读取" in message
               for message in report["warnings"])
    assert_basic_cells(out)

    def broken_reader(*args, **kwargs):
        raise ValueError("无法解析单元格数据")

    monkeypatch.setattr(native_book, "load_workbook", broken_reader)
    broken_out = tmp_path / "broken.xlsx"
    report = merge_workbook({"paths": [str(source)]}, broken_out)
    assert not report["can_export"] and not broken_out.exists()
    assert any("无法解析单元格数据" in issue["message"] for issue in report["issues"])
