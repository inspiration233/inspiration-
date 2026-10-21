"""Validate native concatenation at the OOXML package/reference level."""
from copy import copy
from pathlib import Path
import posixpath
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, Protection, NamedStyle
from openpyxl.workbook.views import BookView
from openpyxl.worksheet.views import SheetView

from engine.tools.excel_merge.merge import Options, merge


MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"m": MAIN}


def check_package(path):
    """Independent XML checks; an openpyxl read alone misses foreign style IDs."""
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        parts = {name: archive.read(name) for name in archive.namelist()}
    trees = {name: ET.fromstring(raw) for name, raw in parts.items()
             if name.endswith((".xml", ".rels"))}
    style = trees["xl/styles.xml"]
    tables = {name: style.find("m:" + name, NS) for name in
              ("fonts", "fills", "borders", "cellXfs", "cellStyleXfs", "numFmts", "dxfs")}
    for table in tables.values():
        if table is not None and "count" in table.attrib:
            assert int(table.attrib["count"]) == len(table)
    custom_formats = {int(n.attrib["numFmtId"]) for n in tables["numFmts"]} if tables["numFmts"] is not None else set()
    for table_name in ("cellXfs", "cellStyleXfs"):
        for xf in tables[table_name]:
            for attribute, table in (("fontId", "fonts"), ("fillId", "fills"),
                                     ("borderId", "borders"), ("xfId", "cellStyleXfs")):
                if attribute in xf.attrib:
                    assert 0 <= int(xf.attrib[attribute]) < len(tables[table]), (attribute, xf.attrib)
            fmt = int(xf.attrib.get("numFmtId", "0"))
            assert fmt < 164 or fmt in custom_formats
    book_views = len(trees["xl/workbook.xml"].find("m:bookViews", NS))
    for name, tree in trees.items():
        if name.endswith(".rels"):
            base = "" if name == "_rels/.rels" else name.rsplit("/_rels/", 1)[0]
            for relationship in tree:
                if relationship.attrib.get("TargetMode") == "External":
                    continue
                target = relationship.attrib["Target"]
                resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                assert resolved in parts, (name, resolved)
        if name.startswith("xl/worksheets/") and name.endswith(".xml"):
            for element in tree.iter():
                local = element.tag.rsplit("}", 1)[-1]
                attr = "style" if local == "col" else "s" if local in ("c", "row") else None
                if attr and attr in element.attrib:
                    assert 0 <= int(element.attrib[attr]) < len(tables["cellXfs"])
                if local == "sheetView":
                    assert 0 <= int(element.attrib["workbookViewId"]) < book_views
                if "{" + REL + "}id" in element.attrib:
                    rels_name = posixpath.join(posixpath.dirname(name), "_rels", posixpath.basename(name) + ".rels")
                    assert rels_name in trees
                    assert element.attrib["{" + REL + "}id"] in {r.attrib["Id"] for r in trees[rels_name]}
    return trees


def make_source(tmp_path, label, all_attributes):
    wb = Workbook(); ws = wb.active; ws.title = "计划"
    ws.append([label + "计划", None]); ws.append(["编号", "数量"]); ws.append([label, 2])
    for i in range(8):
        cell = ws.cell(50 + i, 20, "synthetic style filler")
        cell.font = Font(name=f"UnusedFont{i}")
        cell.fill = PatternFill("solid", fgColor=f"FF{i:06X}")
        cell.border = Border(bottom=Side(style="thin", color=f"FF{i:06X}"))
        cell.alignment = Alignment(textRotation=i + 1)
        cell.number_format = f'0.00 "unused{i}"'
    wb.add_named_style(NamedStyle(name="UnrelatedOne"))
    wb.add_named_style(NamedStyle(name="UnrelatedTwo"))
    for row_index in (1, 2):
        dim = ws.row_dimensions[row_index]
        dim.height = 24 + row_index
        dim.font = Font(name=f"{label}RowOnly{row_index}", size=17, bold=True)
        dim.fill = PatternFill("solid", fgColor="FF123456")
        if all_attributes:
            dim.border = Border(bottom=Side(style="thick", color="FF112233"))
            dim.alignment = Alignment(horizontal="center", textRotation=15)
            dim.protection = Protection(locked=False, hidden=True)
            dim.number_format = f'0.0000 "{label}-row"'
            dim.quotePrefix = True
            dim.pivotButton = True
            dim._style.xfId = 2
    path = tmp_path / f"{label}.xlsx"; wb.save(path); wb.close()
    return {"path": str(path), "sheet": "计划", "selection": {
        "axis": "rows", "header_start": 2, "header_end": 2,
        "data_start": 3, "data_end": 3, "field_start": 1, "field_end": 2}}


def run_merge(tmp_path, specs):
    out = tmp_path / "export.xlsx"
    before = {s["path"]: Path(s["path"]).read_bytes() for s in specs}
    merge([s["path"] for s in specs], out, Options(
        selections=specs, preserve_template_preamble=True, preserve_blank_rows=True,
        remove_summary=False, strip_serial_rows=False, renumber_serial=False, metadata_columns=False))
    assert all(Path(name).read_bytes() == content for name, content in before.items())
    return out


@pytest.mark.parametrize("all_attributes", [False, True])
@pytest.mark.parametrize("source_count", [1, 2])
def test_row_dimension_styles_are_registered_in_the_output_book(tmp_path, all_attributes, source_count):
    specs = [make_source(tmp_path, label, all_attributes) for label in ("甲", "乙")[:source_count]]
    out = run_merge(tmp_path, specs)
    check_package(out)
    actual = load_workbook(out); ws = actual.active
    for spec in specs:
        original = load_workbook(spec["path"]); source = original.active
        title = source["A1"].value
        row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == title)
        for offset in (0, 1):
            expected, got = source.row_dimensions[1 + offset], ws.row_dimensions[row + offset]
            assert got.height == expected.height
            for attribute in ("font", "fill", "border", "alignment", "protection"):
                assert copy(getattr(got, attribute)) == copy(getattr(expected, attribute))
            assert got.number_format == expected.number_format
            assert got.quotePrefix == expected.quotePrefix and got.pivotButton == expected.pivotButton
            assert got._style.xfId == 0
        original.close()
    actual.close()


def test_source_window_ids_are_rebased_to_the_single_output_window(tmp_path):
    spec = make_source(tmp_path, "甲", False)
    wb = load_workbook(spec["path"])
    wb.views = [BookView(), BookView()]
    ws = wb.active
    ws.views.sheetView = [SheetView(workbookViewId=1, zoomScale=135), SheetView(workbookViewId=0)]
    ws.freeze_panes = "B3"
    expected_pane, expected_selection = copy(ws.sheet_view.pane), [copy(s) for s in ws.sheet_view.selection]
    wb.save(spec["path"]); wb.close()
    out = run_merge(tmp_path, [spec])
    trees = check_package(out)
    views = trees["xl/worksheets/sheet1.xml"].findall("m:sheetViews/m:sheetView", NS)
    assert len(views) == 1 and views[0].attrib["workbookViewId"] == "0"
    actual = load_workbook(out)
    assert actual.active.sheet_view.zoomScale == 135
    assert actual.active.sheet_view.pane == expected_pane
    assert list(actual.active.sheet_view.selection) == expected_selection
    actual.close()


def test_printer_binary_relationship_is_omitted_without_dangling_id(tmp_path):
    spec = make_source(tmp_path, "甲", False)
    wb = load_workbook(spec["path"])
    wb.active.page_setup.paperSize = 9
    wb.active.page_setup.orientation = "landscape"
    wb.active.page_margins.left = 0.6
    wb.save(spec["path"]); wb.close()
    with ZipFile(spec["path"]) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    sheet = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    sheet.find("m:pageSetup", NS).set("{" + REL + "}id", "printer1")
    parts["xl/worksheets/sheet1.xml"] = ET.tostring(sheet)
    relationships = ET.Element("{" + PACKAGE_REL + "}Relationships")
    ET.SubElement(relationships, "{" + PACKAGE_REL + "}Relationship", Id="printer1",
                  Type=REL + "/printerSettings", Target="/xl/printerSettings/printerSettings1.bin")
    parts["xl/worksheets/_rels/sheet1.xml.rels"] = ET.tostring(relationships)
    parts["xl/printerSettings/printerSettings1.bin"] = b"synthetic device-specific settings"
    types = ET.fromstring(parts["[Content_Types].xml"])
    ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                  PartName="/xl/printerSettings/printerSettings1.bin",
                  ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.printerSettings")
    parts["[Content_Types].xml"] = ET.tostring(types)
    with ZipFile(spec["path"], "w", ZIP_DEFLATED) as archive:
        for name, raw in parts.items(): archive.writestr(name, raw)
    out = run_merge(tmp_path, [spec])
    trees = check_package(out)
    setup = trees["xl/worksheets/sheet1.xml"].find("m:pageSetup", NS)
    assert "{" + REL + "}id" not in setup.attrib
    actual = load_workbook(out)
    assert actual.active.page_setup.paperSize == 9
    assert actual.active.page_setup.orientation == "landscape"
    assert actual.active.page_margins.left == 0.6
    actual.close()
