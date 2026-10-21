"""Independent regressions for the actual preview -> automatic-selection UI path."""
import json

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from engine.tools.excel_merge.preview import preview
from engine.tools.excel_merge.workbench import merge_export, merge_preview


def workbook(path, rows, merges=(), header_rows=1):
    book = Workbook()
    sheet = book.active
    sheet.title = "明细"
    for row in rows:
        sheet.append(row)
    for area in merges:
        sheet.merge_cells(area)
    for row in sheet.iter_rows(max_row=header_rows):
        for cell in row:
            cell.font = Font(name="宋体", bold=True, color="FFFFFFFF", size=13)
            cell.fill = PatternFill("solid", fgColor="FF234567")
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=header_rows + 1):
        for cell in row:
            cell.font = Font(name="仿宋", size=12, color="FF123456")
    sheet.column_dimensions["B"].width = 23
    sheet.freeze_panes = f"B{header_rows + 1}"
    book.save(path)
    book.close()
    return path


def ui_payload(*paths, **options):
    views = [preview({"path": str(path), "row_limit": 1, "col_limit": 1}) for path in paths]
    return {
        "paths": [str(path) for path in paths],
        "template": {"path": str(paths[0]), "sheet": views[0]["sheet"]},
        "selections": [{"path": str(path), "sheet": view["sheet"], "selection": view["selection"],
                        "selection_origin": "auto"} for path, view in zip(paths, views)],
        "options": {"mode": "single", **options},
    }, views


def export(tmp_path, payload):
    out = tmp_path / "合并结果.xlsx"
    report = merge_export({**payload, "out": str(out)})
    assert report["can_export"] and out.exists(), report
    return load_workbook(out), report


def record(sheet, identity):
    matches = [list(row) for row in sheet.iter_rows(values_only=True) if row[0] == identity]
    assert len(matches) == 1, (identity, matches)
    return matches[0]


def test_ui_normalization_keeps_six_columns_original_headings_and_styles(tmp_path):
    headings = ["客户", "月份", "项目", "数量", "AMOUNT", "成本 名称"]
    first = workbook(tmp_path / "甲.xlsx", [headings, ["A1", "1月", "事项甲", 2, 20, "服务甲"],
                                          ["A2", "2月", "事项乙", 3, 30, "服务乙"]])
    second = workbook(tmp_path / "乙.xlsx", [["客户", "月份", "项目", "数量", "ａｍｏｕｎｔ", "成本　　名称"],
                                            ["B1", "1月", "事项丙", 4, 40, "服务丙"],
                                            ["B2", "2月", "事项丁", 5, 50, "服务丁"]])
    payload, views = ui_payload(first, second)
    assert views[0]["fields"][4]["signature"] != views[1]["fields"][4]["signature"]
    assert views[0]["fields"][4]["match_signature"] == views[1]["fields"][4]["match_signature"]
    assert views[0]["fields"][5]["match_signature"] == views[1]["fields"][5]["match_signature"]
    actual, _ = export(tmp_path, payload)
    try:
        assert len(actual.worksheets) == 1
        sheet = actual.active
        assert sheet.max_column == 6
        assert [cell.value for cell in sheet[1]] == headings
        assert record(sheet, "B1") == ["B1", "1月", "事项丙", 4, 40, "服务丙"]
        assert record(sheet, "B2") == ["B2", "2月", "事项丁", 5, 50, "服务丁"]
        assert sheet["E1"].font.name == "宋体" and sheet["E1"].font.bold
        assert sheet["E1"].fill.fgColor.rgb == "FF234567"
        assert sheet["E4"].font.name == "仿宋" and sheet["E4"].font.color.rgb == "FF123456"
        assert sheet.column_dimensions["B"].width == 23
    finally:
        actual.close()


def test_match_signatures_normalize_each_layer_without_flattening_display(tmp_path):
    first = workbook(tmp_path / "多层甲.xlsx", [["客户", "指标", None], [None, "ＮＵＭ", "成本  名称"],
                    ["A1", 2, "服务甲"], ["A2", 3, "服务乙"]], ("A1:A2", "B1:C1"), 2)
    second = workbook(tmp_path / "多层乙.xlsx", [["客户", "指标", None], [None, "num", "成本　名称"],
                     ["B1", 4, "服务丙"], ["B2", 5, "服务丁"]], ("A1:A2", "B1:C1"), 2)
    literal = workbook(tmp_path / "单格斜杠.xlsx", [["客户", "指标 / num", "成本 名称"],
                      ["L1", 6, "服务戊"], ["L2", 7, "服务己"]])
    payload, views = ui_payload(first, second)
    assert views[0]["selection"]["header_end"] == views[1]["selection"]["header_end"] == 2
    field = views[0]["fields"][1]
    assert json.loads(field["signature"]) == ["指标", "ＮＵＭ"]
    assert json.loads(field["match_signature"]) == ["指标", "num"]
    assert field["match_signature"] == views[1]["fields"][1]["match_signature"]
    literal_view = preview({"path": str(literal)})
    assert field["match_signature"] != literal_view["fields"][1]["match_signature"]
    actual, _ = export(tmp_path, payload)
    try:
        sheet = actual.active
        assert len(actual.worksheets) == 1 and sheet.max_column == 3
        assert sheet["B1"].value == "指标" and sheet["B2"].value == "ＮＵＭ"
        assert sheet["C2"].value == "成本  名称"
        assert {str(area) for area in sheet.merged_cells.ranges} == {"A1:A2", "B1:C1"}
        assert record(sheet, "B1") == ["B1", 4, "服务丙"]
    finally:
        actual.close()


def multilevel(path, parent="计划", prefix="A"):
    return workbook(path, [["客户", parent, None], [None, "数量", "金额"],
                    [prefix + "1", 2, 20], [prefix + "2", 3, 30]], ("A1:A2", "B1:C1"), 2)


@pytest.mark.parametrize("flat_first", [False, True])
def test_ui_uniquely_identified_flat_and_multilevel_fields_match_both_directions(tmp_path, flat_first):
    layered = multilevel(tmp_path / "多层.xlsx")
    flat = workbook(tmp_path / "平铺.xlsx", [["客户", "数量", "金额"], ["B1", 4, 40], ["B2", 5, 50]])
    paths = (flat, layered) if flat_first else (layered, flat)
    payload, _ = ui_payload(*paths)
    actual, _ = export(tmp_path, payload)
    try:
        assert len(actual.worksheets) == 1
        sheet = actual.active
        assert sheet.max_column == 3
        assert record(sheet, "A1") == ["A1", 2, 20]
        assert record(sheet, "B2") == ["B2", 5, 50]
        if not flat_first:
            assert sheet["B1"].value == "计划" and "B1:C1" in sheet.merged_cells
        else:
            assert [cell.value for cell in sheet[1]] == ["客户", "数量", "金额"]
    finally:
        actual.close()


def test_different_parent_paths_do_not_share_one_numeric_column(tmp_path):
    plan = multilevel(tmp_path / "计划.xlsx", "计划", "A")
    actual_source = multilevel(tmp_path / "实际.xlsx", "实际", "B")
    payload, _ = ui_payload(plan, actual_source, auto_group=False, merge_shape="stack")
    actual, _ = export(tmp_path, payload)
    try:
        sheet = actual.active
        assert sheet.max_column == 5
        assert record(sheet, "A1") == ["A1", 2, 20, None, None]
        assert record(sheet, "B1") == ["B1", None, None, 2, 20]
    finally:
        actual.close()


def test_flat_leaf_is_not_guessed_between_two_parent_groups(tmp_path):
    layered = workbook(tmp_path / "两组.xlsx", [["客户", "计划", None, "实际", None],
                       [None, "数量", "金额", "数量", "金额"], ["A1", 2, 20, 1, 10],
                       ["A2", 3, 30, 2, 20]], ("A1:A2", "B1:C1", "D1:E1"), 2)
    flat = workbook(tmp_path / "未指明父级.xlsx", [["客户", "数量", "金额"], ["B1", 4, 40], ["B2", 5, 50]])
    payload, _ = ui_payload(layered, flat, auto_group=False, merge_shape="stack")
    actual, _ = export(tmp_path, payload)
    try:
        sheet = actual.active
        assert sheet.max_column == 7
        assert record(sheet, "B1") == ["B1", None, None, None, None, 4, 40]
        assert record(sheet, "A1") == ["A1", 2, 20, 1, 10, None, None]
    finally:
        actual.close()


@pytest.mark.skip(reason="重名字段已放宽: 同数量重名按位置自动对应,不再硬阻断")
def test_ui_duplicate_headings_require_a_positional_choice(tmp_path):
    first = workbook(tmp_path / "甲.xlsx", [["客户", "数量", "数量"], ["A1", 2, 20], ["A2", 3, 30]])
    second = workbook(tmp_path / "乙.xlsx", [["客户", "数量", "数量"], ["B1", 4, 40], ["B2", 5, 50]])
    payload, _ = ui_payload(first, second)
    report = merge_preview(payload)
    assert not report["can_export"]
    assert any("重名" in issue["message"] for issue in report["issues"])
    out = tmp_path / "不能猜测.xlsx"
    with pytest.raises(ValueError):
        merge_export({**payload, "out": str(out)})
    assert not out.exists()


def uncertain_source(tmp_path):
    return workbook(tmp_path / "方向需确认.xlsx", [["项目", "甲", "乙", "丙"],
                    *[[f"项{i}", "数据", "数据", "数据"] for i in range(1, 5)]])


def direction_notices(report):
    return [message for message in report.get("warnings", []) if "方向" in message and "确认" in message]


def test_auto_selection_retains_recognition_warning_through_preview_and_merge(tmp_path):
    source = uncertain_source(tmp_path)
    payload, views = ui_payload(source)
    assert direction_notices(views[0]), views[0]["warnings"]
    repeated = preview({"path": str(source), "sheet": "明细", "selection": views[0]["selection"],
                        "selection_origin": "auto"})
    assert direction_notices(repeated), repeated["warnings"]
    report = merge_preview(payload)
    assert report["can_export"], report
    assert direction_notices(report), report["warnings"]
    actual, exported = export(tmp_path, payload)
    try:
        assert direction_notices(exported), exported["warnings"]
        assert record(actual.active, "项1") == ["项1", "数据", "数据", "数据"]
    finally:
        actual.close()


@pytest.mark.parametrize("origin", ["manual", None])
def test_manual_selection_clears_only_resolved_recognition_warning(tmp_path, origin):
    source = uncertain_source(tmp_path)
    payload, views = ui_payload(source)
    selected = views[0]["selection"]
    assert direction_notices(views[0])
    view_request = {"path": str(source), "sheet": "明细", "selection": selected}
    spec = payload["selections"][0]
    if origin is None:
        spec.pop("selection_origin")  # Existing explicit callers remain manual.
    else:
        spec["selection_origin"] = origin
        view_request["selection_origin"] = origin
    manual = preview(view_request)
    assert not direction_notices(manual), manual["warnings"]
    report = merge_preview(payload)
    assert report["can_export"], report
    assert not direction_notices(report), report["warnings"]
