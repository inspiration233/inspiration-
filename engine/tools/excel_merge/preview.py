"""Read-only, paged native worksheet preview and explicit selection metadata."""
import datetime as dt
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter

from . import reader
from .recognize import normalize_name


def json_value(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def normalize_selection(raw, max_row, max_column, suggested=None):
    raw = dict(raw or {})
    base = dict(suggested or {})
    axis = raw.get("axis", base.get("axis", "rows"))
    if axis not in ("rows", "columns"):
        raise ValueError("记录方向必须是 rows 或 columns")
    records = max_row if axis == "rows" else max_column
    fields = max_column if axis == "rows" else max_row
    if base.get("axis") != axis:
        base = {}
    result = {"axis": axis}
    defaults = {"header_start": 1, "header_end": 1, "data_start": 2,
                "data_end": records, "field_start": 1, "field_end": fields}
    for key, default in defaults.items():
        result[key] = int(raw.get(key, base.get(key, default)))
    hs, he = result["header_start"], result["header_end"]
    ds, de = result["data_start"], result["data_end"]
    fs, fe = result["field_start"], result["field_end"]
    if not (1 <= hs <= he < ds <= records + 1 and ds - 1 <= de <= records):
        raise ValueError("表头和数据范围无效：表头须位于数据之前，范围不能超过工作表")
    if not 1 <= fs <= fe <= fields:
        raise ValueError("字段范围超出了工作表")
    return result


def anchor_cell(ws, row, col):
    cell = ws.cell(row, col)
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
            return ws.cell(merged.min_row, merged.min_col)
    return cell


def field_metadata(ws, selection):
    horizontal = selection["axis"] == "columns"
    result = []
    for index in range(selection["field_start"], selection["field_end"] + 1):
        parts = []
        for header in range(selection["header_start"], selection["header_end"] + 1):
            row, col = (index, header) if horizontal else (header, index)
            value = anchor_cell(ws, row, col).value
            text = str(json_value(value)).strip() if value is not None else ""
            if text and (not parts or parts[-1] != text):
                parts.append(text)
        label = " / ".join(parts)
        # Empty headings deliberately include position: they cannot be matched by name.
        signature = json.dumps(parts, ensure_ascii=False, separators=(",", ":")) if parts else f"@empty:{index}"
        match_parts = []
        for part in parts:
            normalized = normalize_name(part)
            if normalized and (not match_parts or match_parts[-1] != normalized):
                match_parts.append(normalized)
        match_signature = json.dumps(match_parts, ensure_ascii=False, separators=(",", ":")) if match_parts else f"@empty:{index}"
        result.append({"index": index, "label": label or
                       (f"第 {index} 行" if horizontal else f"{get_column_letter(index)} 列（空表头）"),
                       "signature": signature, "match_signature": match_signature, "empty": not bool(parts)})
    return result


def _color(color, default=None):
    if color is None:
        return default
    if color.type == "rgb" and isinstance(color.rgb, str):
        return "#" + color.rgb[-6:]
    if color.type == "indexed" and isinstance(color.indexed, int) and color.indexed < len(COLOR_INDEX):
        return "#" + COLOR_INDEX[color.indexed][-6:]
    # Theme colors need the theme palette; use neutral defaults instead of guessing.
    return default


def _style(cell):
    font, fill, align, border = cell.font, cell.fill, cell.alignment, cell.border
    result = {"fontFamily": font.name or "Calibri", "fontSize": f"{font.sz or 11}pt",
              "fontWeight": "bold" if font.b else "normal",
              "fontStyle": "italic" if font.i else "normal",
              "color": _color(font.color), "background": _color(fill.fgColor) if fill.patternType == "solid" else None,
              "textAlign": align.horizontal if align.horizontal in ("left", "center", "right", "justify") else
                           ("right" if isinstance(cell.value, (int, float)) else "left"),
              "verticalAlign": {"center": "middle"}.get(align.vertical, align.vertical or "bottom"),
              "whiteSpace": "pre-wrap" if align.wrap_text else "pre",
              "textDecoration": "line-through" if font.strike else ("underline" if font.u else "none")}
    for name in ("top", "bottom", "left", "right"):
        side = getattr(border, name)
        if side and side.style:
            width = "2px" if side.style in ("medium", "thick", "double") else "1px"
            kind = "dashed" if "dash" in side.style.lower() else ("dotted" if side.style == "dotted" else "solid")
            result["border" + name.capitalize()] = f"{width} {kind} {_color(side.color, '#8b95a3')}"
    return {k: v for k, v in result.items() if v is not None}


def _display(value, number_format):
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S") if value.time() != dt.time() else value.strftime("%Y-%m-%d")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (float, int)):
        first = (number_format or "General").split(";")[0]
        if "%" in first:
            decimals = len(first.split(".")[-1].split("%")[0]) if "." in first else 0
            return f"{value:.{min(decimals, 10)}%}"
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
    return str(value)


def automatic_selection(structure, max_row, max_col):
    axis = "columns" if structure and structure.orientation == "horizontal" else "rows"
    # 原生交叉表(记录沿列): 选区语义 = 字段=行, 记录=列, 表头=标签列(A列)。
    # 记录列取结构识别的表头行非空列范围, 而非"到表尾"(空列会变成空白组)。
    cross = getattr(structure, "cross_records", None) if structure else None
    if axis == "columns" and cross:
        return {"axis": axis, "header_start": 1, "header_end": 1,
                "data_start": cross[0], "data_end": min(cross[1], max_col),
                "field_start": 1, "field_end": max_row}
    end = max_col if axis == "columns" else max_row
    hs = max(1, structure.header_row if structure else 1)
    he = min(end, hs + max(1, structure.height if structure else 1) - 1)
    return {"axis": axis, "header_start": hs, "header_end": he, "data_start": he + 1,
            "data_end": end, "field_start": 1, "field_end": max_row if axis == "columns" else max_col}


def _suggest(path, sheet, max_row, max_col, workbook=None):
    rd = None
    try:
        if workbook is not None:
            rd = reader._XlsxReader(path, workbook=workbook)
            st = rd.build_structure(sheet)
        else:
            tables, rd = reader.list_tables(path, True)
            source = next((src for src, _ in tables if src is not None and src.sheet_name == sheet), None)
            st = source.structure if source else None
        recognition = {"low_confidence": bool(st and st.low_confidence),
                       "orientation_confident": bool(st and st.orientation_confident),
                       "native_cross": bool(getattr(st, "native_cross", False)),
                       "warnings": list(st.warnings) if st else ["未能识别表头，请核对所示范围。"]}
        if recognition["low_confidence"]:
            recognition["warnings"].append("表头识别把握较低，请核对表头和数据范围；需要时可在原表上调整。")
        if recognition["native_cross"]:
            recognition["warnings"].append(
                "横向配置表：字段排在行上（点左侧行号选择），记录沿列方向；范围由程序识别。")
        return automatic_selection(st, max_row, max_col), recognition
    except Exception:
        return ({"axis": "rows", "header_start": 1, "header_end": 1,
                 "data_start": 2, "data_end": max_row, "field_start": 1, "field_end": max_col},
                {"low_confidence": True, "orientation_confident": False, "native_cross": False,
                 "warnings": ["自动识别未完成，暂按第一行表头显示，请核对或调整范围。"]})
    finally:
        if rd:
            rd.close()


def _plain_preview(path, requested_sheet):
    """Legacy formats remain available for merge, with honest non-native preview."""
    from openpyxl import Workbook
    tables, rd = reader.list_tables(path, True)
    try:
        sources = [src for src, _ in tables if src is not None]
        src = next((s for s in sources if s.sheet_name == requested_sheet), None) if requested_sheet else sources[0]
        if src is None:
            raise ValueError("找不到选择的工作表")
        wb = Workbook()
        ws = wb.active
        ws.title = src.sheet_name[:31]
        for row, cells in enumerate(src.header_cells(), src.structure.header_row):
            for col, cell in enumerate(cells, 1):
                ws.cell(row, col, getattr(cell, "value", cell))
        for row, cells in src.iter_data_rows():
            for col, cell in enumerate(cells, 1):
                ws.cell(row, col, getattr(cell, "value", cell))
        return wb, ws, [{"name": s.sheet_name, "state": "visible"} for s in sources]
    finally:
        rd.close()


def preview(payload):
    path = Path(payload["path"]).resolve()
    if not path.is_file():
        raise ValueError("文件不存在")
    wb = cached = None
    warnings = []
    converted = None
    try:
        if path.suffix.lower() in (".csv", ".xls"):
            # csv/xls 统一转成 xlsx 走原生链路: 识别/选列/范围与 xlsx 一致
            from .plain_import import to_xlsx
            converted, notes = to_xlsx(path)
            path = Path(converted)
            warnings.extend(notes)
        if path.suffix.lower() == ".xlsx":
            wb = load_workbook(path, data_only=False)
            cached = load_workbook(path, data_only=True, read_only=True)
            sheet = payload.get("sheet") or next((s.title for s in wb if s.sheet_state == "visible"), wb.sheetnames[0])
            if sheet not in wb:
                raise ValueError("工作表不存在，请重新选择")
            ws = wb[sheet]
            sheets = [{"name": s.title, "max_row": s.max_row, "max_column": s.max_column,
                       "state": s.sheet_state} for s in wb]
        elif path.suffix.lower() in (".csv", ".xls"):
            wb, ws, sheets = _plain_preview(path, payload.get("sheet"))
            warnings.append("此格式提供内容预览；原生拆分与回填请先在 Excel/WPS 中另存为 .xlsx。")
        else:
            raise ValueError("支持 .xlsx、.xls 和 .csv")
        origin = payload.get("selection_origin") or ("manual" if payload.get("selection") else "auto")
        if origin == "manual" and payload.get("selection"):
            suggested = None
            recognition = {"low_confidence": False, "orientation_confident": True, "warnings": []}
        else:
            suggested, recognition = _suggest(path, ws.title, ws.max_row, ws.max_column, cached)
            suggested = _apply_split_meta(suggested, path, ws)
        recognition["confirmed"] = origin == "manual"
        warnings.extend(recognition["warnings"])
        selection = normalize_selection(payload.get("selection") if origin == "manual" else None,
                                        ws.max_row, ws.max_column, suggested)
        fields = field_metadata(ws, selection)
        sr, sc = max(1, int(payload.get("start_row", 1))), max(1, int(payload.get("start_col", 1)))
        er = min(ws.max_row, sr + min(200, max(1, int(payload.get("row_limit", 80)))) - 1)
        ec = min(ws.max_column, sc + min(100, max(1, int(payload.get("col_limit", 40)))) - 1)
        cached_values = {}
        if cached:
            for ri, row in enumerate(cached[ws.title].iter_rows(min_row=sr, max_row=er, min_col=sc, max_col=ec), sr):
                for ci, cell in enumerate(row, sc):
                    cached_values[(ri, ci)] = cell.value
        cells = []
        for row in ws.iter_rows(min_row=sr, max_row=er, min_col=sc, max_col=ec):
            for cell in row:
                formula = cell.value if cell.data_type == "f" else None
                value = cached_values.get((cell.row, cell.column)) if formula else cell.value
                cells.append({"row": cell.row, "col": cell.column, "value": json_value(cell.value),
                              "text": _display(value, cell.number_format) if value is not None else (formula or ""),
                              "formula": formula, "style": _style(cell),
                              "number_format": cell.number_format})
        merges = [{"min_row": m.min_row, "max_row": m.max_row, "min_col": m.min_col, "max_col": m.max_col}
                  for m in ws.merged_cells.ranges if m.max_row >= sr and m.min_row <= er and m.max_col >= sc and m.min_col <= ec]
        signature = hashlib.sha256(json.dumps(
            {"axis": selection["axis"], "depth": selection["header_end"] - selection["header_start"] + 1,
             "fields": [f["signature"] for f in fields]}, ensure_ascii=False).encode()).hexdigest()
        return {"path": str(path), "sheets": sheets, "sheet": ws.title,
                "max_row": ws.max_row, "max_column": ws.max_column,
                "cells": cells, "merges": merges,
                "row_heights": {str(i): round((ws.row_dimensions[i].height or 15) * 4 / 3, 1) for i in range(sr, er + 1)},
                "column_widths": {str(i): round((ws.column_dimensions[get_column_letter(i)].width or 13) * 7 + 5, 1)
                                  for i in range(sc, ec + 1)},
                "selection": selection, "selection_origin": origin, "recognition": recognition,
                "fields": fields, "signature": signature,
                "warnings": warnings, "truncated": sr > 1 or sc > 1 or er < ws.max_row or ec < ws.max_column,
                "viewport": {"start_row": sr, "end_row": er, "start_col": sc, "end_col": ec}}
    finally:
        if cached:
            cached.close()
        if wb:
            wb.close()

def _split_meta(path):
    """读取拆分时写进自定义属性的溯源信息(列对应表); 非本工具拆分产物返回 None。
    用标准 XML 解析取值——Excel/WPS 保存时会重打包并按自己的风格重写 custom.xml
    (属性顺序/命名空间前缀/转义/空白都可能变), 正则匹配会静默失配。"""
    import json
    import re
    import xml.etree.ElementTree as ET
    from zipfile import ZipFile
    try:
        with ZipFile(str(path)) as archive:
            if "docProps/custom.xml" not in archive.namelist():
                return None
            raw = archive.read("docProps/custom.xml")
    except Exception:
        return None
    payload = None
    try:
        for prop in ET.fromstring(raw).iter():
            if prop.tag.rsplit("}", 1)[-1] == "property" and prop.get("name") == "TOOLBOX_SPLIT":
                for child in prop:
                    if child.text:
                        payload = child.text
                        break
    except ET.ParseError:
        payload = None
    if payload is None:                      # 极端格式退回正则兜底
        text = raw.decode("utf-8", "ignore")
        m = re.search(r'name="TOOLBOX_SPLIT"[^>]*>(.*?)</property>', text, re.S)
        if m:
            m2 = re.search(r">([^<>]*\{[^<>]*)", m.group(1), re.S)
            payload = m2.group(1) if m2 else None
    if not payload:
        return None
    try:
        meta = json.loads(payload)
    except ValueError:
        return None
    return meta if isinstance(meta, dict) and isinstance(meta.get("columns"), dict) else None


def _apply_split_meta(suggested, path, ws):
    """本工具拆分产物带拆分时的结构区间: 表头/字段范围按拆分时的用。
    每组只剩几行数据的产物, 自动识别常把多级表头判成单行(字段丢掉上一级),
    与底稿的对应表签名全部失配——按拆分区间识别即恢复一致。"""
    meta = _split_meta(path)
    if not meta or not suggested or meta.get("sheet") != ws.title:
        return suggested
    ms = meta.get("selection") or {}
    keys = ("header_start", "header_end", "field_start", "field_end")
    if not all(isinstance(ms.get(k), int) for k in keys):
        return suggested
    if suggested.get("axis") != ms.get("axis", suggested.get("axis")):
        return suggested
    limit = ws.max_column if suggested["axis"] == "rows" else ws.max_row
    fixed = dict(suggested, header_start=ms["header_start"], header_end=ms["header_end"],
                 field_start=ms["field_start"], field_end=min(ms["field_end"], limit),
                 data_start=ms["header_end"] + 1)
    if fixed["data_end"] < fixed["data_start"]:
        fixed["data_end"] = fixed["data_start"]
    return fixed
