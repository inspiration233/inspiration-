# -*- coding: utf-8 -*-
"""CSV / XLS → XLSX 统一转换层。

拆分要"原样复制样式再删行"的工作方式, xls(1997 格式)与 csv 装不下现代
工作簿的全部内容——统一在入口转成临时 xlsx, 之后识别/选列/范围/拆分/
属性写入与 xlsx 完全同路。

- CSV: 编码依次尝试 utf-8-sig / gbk / gb18030; 纯数字转数值(保留前导 0
  和超长号), 无格式无合并, 明确警告。
- XLS(第一步): 值 + 合并单元格 + 列宽行高尽力保留; 字体/填充等样式
  不迁移, 警告说明。产物一律 xlsx。
"""
from __future__ import annotations

import csv as _csv
import datetime as _dt
from pathlib import Path
import tempfile
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# 同一源文件(路径+mtime+大小)复用同一份转换, 避免反复重转
_CACHE: dict = {}


def to_xlsx(path):
    """把 csv/xls 转成临时 xlsx。返回 (xlsx路径, 警告列表); 其它后缀原样返回。"""
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix not in (".csv", ".xls"):
        return str(source), []
    key = (str(source.resolve()), source.stat().st_mtime_ns, source.stat().st_size)
    if key in _CACHE:
        return _CACHE[key][0], list(_CACHE[key][1])
    if suffix == ".csv":
        target, warnings = _csv_to_xlsx(source)
    else:
        target, warnings = _xls_to_xlsx(source)
    _CACHE[key] = (target, list(warnings))
    return target, warnings


def _coerce(text):
    """CSV 文本尽量转数值: 纯数字转数值, 前导 0 与超长号码保持文本。"""
    if not isinstance(text, str) or not text.strip():
        return text
    s = text.strip()
    if len(s) >= 15:
        return text
    if len(s) > 1 and s[0] == "0" and s[1:].isdigit():
        return text
    try:
        if s.isdigit():
            return int(s)
        if s.lower() in ("nan", "inf", "-inf"):
            return text
        return float(s)
    except ValueError:
        return text


def _csv_to_xlsx(source):
    warnings = []
    rows = None
    for encoding in ("utf-8-sig", "gbk", "gb18030"):
        try:
            with open(source, "r", encoding=encoding, newline="") as handle:
                rows = list(_csv.reader(handle))
            if encoding != "utf-8-sig":
                warnings.append(f"CSV 按 {encoding} 编码读取。")
            break
        except UnicodeDecodeError:
            continue
    if rows is None:
        raise ValueError("CSV 编码无法识别（尝试了 UTF-8 / GBK / GB18030），请先在 WPS 中另存为标准编码")
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    r = 0
    for row in rows:
        if not any(str(cell).strip() for cell in row):
            continue                              # 空行(常见于文件尾)不产生记录
        r += 1
        for c, value in enumerate(row, 1):
            if value != "":
                ws.cell(r, c, _coerce(value))
    target = _save_temp(wb, source)
    warnings.append("CSV 没有格式与合并单元格，已按数据表处理；拆分输出为 .xlsx。")
    return target, warnings


def _xls_to_xlsx(source):
    import xlrd
    warnings = []
    try:
        book = xlrd.open_workbook(str(source), formatting_info=True)
    except Exception:                       # 部分 xls 不支持格式信息, 降级读值
        book = xlrd.open_workbook(str(source))
        warnings.append("该 XLS 的格式信息无法读取，列宽行高未保留。")
    wb = Workbook()
    wb.remove(wb.active)
    for index in range(book.nsheets):
        sh = book.sheet_by_index(index)
        ws = wb.create_sheet(title=sh.name[:31])
        for r in range(sh.nrows):
            for c in range(sh.ncols):
                value = _xlrd_value(sh.cell(r, c), book)
                if value is not None:
                    ws.cell(r + 1, c + 1, value)
        for rlo, rhi, clo, chi in sh.merged_cells.ranges if hasattr(sh.merged_cells, "ranges") else sh.merged_cells:
            ws.merge_cells(start_row=rlo + 1, start_column=clo + 1,
                           end_row=rhi, end_column=chi)
        try:
            for idx, info in (sh.colinfo_map or {}).items():
                if info.width:
                    ws.column_dimensions[get_column_letter(idx + 1)].width = info.width / 256
            for idx, info in (sh.rowinfo_map or {}).items():
                if info.height:
                    ws.row_dimensions[idx + 1].height = info.height
        except Exception:
            pass
    if not wb.sheetnames:                   # 空簿兜底
        wb.create_sheet(title="数据")
    warnings.append("XLS 的字体、颜色等样式未完整迁移（值、合并单元格、列宽行高已保留）；拆分输出为 .xlsx。")
    return _save_temp(wb, source), warnings


def _xlrd_value(cell, book):
    import xlrd
    if cell.ctype == xlrd.XL_CELL_TEXT:
        return cell.value
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        return cell.value
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            tt = xlrd.xldate_as_tuple(cell.value, book.datemode)
            return _dt.datetime(*tt) if tt[3:] != (0, 0, 0) else _dt.date(*tt[:3])
        except Exception:
            return cell.value
    return None                              # EMPTY / ERROR / BLANK


def _save_temp(wb, source):
    fd, target = tempfile.mkstemp(prefix=f".{source.stem}-", suffix=".xlsx")
    import os
    os.close(fd)
    wb.save(target)
    wb.close()
    return target
