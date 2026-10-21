"""Combine complete XLSX worksheets without flattening formulas or layouts.

Each input workbook gets its own sheet/name/table namespace. Referenced sheets
are included transitively. Presentation losses are reported; unresolved data
dependencies are rejected before writing.
"""
from collections import OrderedDict
from copy import copy, deepcopy
import colorsys
from pathlib import Path
import re
import warnings
from xml.etree import ElementTree
from zipfile import ZipFile, BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.formula import Tokenizer
from openpyxl.styles import Color
from openpyxl.utils import quote_sheetname
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.views import SheetViewList
from openpyxl.writer.theme import theme_xml

from .merge import Cancelled
from .compatibility import inspect_extensions


_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
_TABLE = re.compile(r"^(@?)([A-Za-z_\\][A-Za-z0-9_.\\]*)(\[.*)$")
_DYNAMIC = re.compile(r"(?:^|[^A-Z0-9_])(?:_XLFN\.)?INDIRECT\s*\(", re.I)


def _cancel(cancel):
    if cancel and cancel():
        raise Cancelled("工作簿合并已取消")


def _guard(path):
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"{path.name}：完整保真合并只支持 .xlsx")
    prefixes = ("xl/embeddings/", "xl/activex/", "xl/ctrlprops/", "xl/pivottables/", "xl/pivotcache/",
                "xl/slicers/", "xl/slicercaches/", "xl/threadedcomments/", "xl/persons/",
                "xl/richdata/", "customxml/", "customui/", "_xmlsignatures/",
                "xl/externallinks/", "xl/macrosheets/",
                "xl/dialogsheets/", "xl/model/", "xl/chartsheets/", "xl/querytables/")
    omitted = {
        "xl/drawings/": "图形、图片及图表等附加对象未保留；已继续合并单元格内容和基础格式。",
        "xl/charts/": "图形、图片及图表等附加对象未保留；已继续合并单元格内容和基础格式。",
        "xl/media/": "图形、图片及图表等附加对象未保留；已继续合并单元格内容和基础格式。",
        "xl/printersettings/": "打印机专用配置未保留；已继续保留页面方向、页边距等常规打印设置。",
    }
    notices = []

    def notice(message):
        message = f"{path.name}：{message}"
        if message not in notices:
            notices.append(message)

    with ZipFile(path) as archive:
        for name in archive.namelist():
            lower = name.lower()
            if (lower.startswith(prefixes) or "vbaproject" in lower
                    or lower in ("xl/connections.xml", "xl/metadata.xml", "xl/xmlmaps.xml")):
                raise ValueError(f"{path.name} 含无法完整保留的对象 {name}，已阻止导出")
            omitted_notice = next((message for prefix, message in omitted.items()
                                   if lower.startswith(prefix)), None)
            if omitted_notice:
                notice(omitted_notice)
                # These parts are not copied. Their extensions do not affect
                # the retained worksheet cells, formulas or namespaces.
                continue
            if lower.startswith("xl/") and lower.endswith(".xml"):
                # Let the shared XML policy inspect every retained part;
                # a byte search would miss extensions in UTF-16 XML.
                for message in inspect_extensions(name, archive.read(name)):
                    notice(message)
    return notices


def _unique(name, used, limit=None):
    root = name[:limit] if limit else name
    candidate, count = root, 1
    while candidate.casefold() in used:
        count += 1
        suffix = f"_{count}"
        candidate = (root[:limit - len(suffix)] if limit else root) + suffix
    used.add(candidate.casefold())
    return candidate


def _tokens(expression):
    if not expression:
        return [], False
    text = str(expression)
    if _DYNAMIC.search(text):
        raise ValueError("含 INDIRECT 动态引用，不能静态证明依赖页和重命名后的引用完整")
    equal = text.startswith("=")
    try:
        return Tokenizer(text if equal else "=" + text).items, equal
    except Exception as exc:
        raise ValueError(f"无法解析公式或命名范围：{text}") from exc


def _sheet_reference(token, titles):
    """Return referenced source sheet titles; reject missing/external sheets."""
    prefix = token.rsplit("!", 1)[0]
    if prefix.startswith("'") and prefix.endswith("'"):
        prefix = prefix[1:-1].replace("''", "'")
    if "[" in prefix or "]" in prefix:
        raise ValueError("含外部工作簿引用，不能保证合并后的计算独立完整")
    lookup = {name.casefold(): name for name in titles}
    ends = prefix.split(":")
    if len(ends) not in (1, 2) or any(name.casefold() not in lookup for name in ends):
        raise ValueError(f"公式引用了不存在的工作表：{prefix}")
    names = [lookup[name.casefold()] for name in ends]
    if len(names) == 1:
        return names
    first, last = sorted((titles.index(names[0]), titles.index(names[1])))
    return titles[first:last + 1]


def _expressions(ws):
    for cell in ws._cells.values():
        if cell.data_type == "f":
            value = cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", None)
            if value:
                yield value
        if cell.hyperlink and cell.hyperlink.location:
            yield cell.hyperlink.location
        elif cell.hyperlink and cell.hyperlink.target and cell.hyperlink.target.startswith("#"):
            yield cell.hyperlink.target[1:]
    for dv in ws.data_validations.dataValidation:
        for value in (dv.formula1, dv.formula2):
            if value:
                yield value
    for cf in ws.conditional_formatting:
        for rule in ws.conditional_formatting[cf]:
            yield from rule.formula or []
    for table in ws.tables.values():
        for column in table.tableColumns:
            for formula in (column.calculatedColumnFormula, column.totalsRowFormula):
                if formula is not None and formula.attr_text:
                    yield formula.attr_text
    for name in ws.defined_names.values():
        if name.attr_text:
            yield name.attr_text


class InputBook:
    def __init__(self, path, selected):
        self.path = path
        self.warnings = _guard(path)
        with warnings.catch_warnings(record=True) as emitted:
            warnings.simplefilter("always")
            self.book = load_workbook(path, rich_text=True, data_only=False)
        self.warnings.extend(f"{path.name} 读取信息（{w.category.__name__}）：{w.message}"
                             for w in emitted)
        self.custom_table_styles = {style.name.casefold()
                                    for style in self.book._table_styles.tableStyle}
        if self.custom_table_styles:
            self.warnings.append(f"{path.name} 的自定义表样式未保留；"
                                 "已继续合并表格内容、公式和单元格基础格式。")
        self.selected = set(self.book.sheetnames if selected is None else selected)
        missing = self.selected.difference(self.book.sheetnames)
        if missing:
            self.book.close()
            raise ValueError(f"{path.name} 缺少所选工作表：{', '.join(sorted(missing))}")
        self.required = set(self.selected)
        self.table_sheets = {alias.casefold(): ws.title for ws in self.book
                             for table in ws.tables.values()
                             for alias in (table.name, table.displayName) if alias}
        self.sheet_map, self.name_map, self.table_map = {}, {}, {}
        self.dxf_map = {}
        self.theme_colors, self.theme_fonts = self._theme()
        if (any(ws.tables for ws in self.book)
                and (self.theme_colors, self.theme_fonts) != self._theme(theme_xml)):
            self.warnings.append(f"{path.name} 的自定义主题与 Excel 表样式组合可能改变表格配色；"
                                 "已继续合并内容、公式和单元格基础格式，请核对表格外观。")

    def _theme(self, blob=None):
        colors, fonts = [], {}
        blob = self.book.loaded_theme if blob is None else blob
        if blob:
            root = ElementTree.fromstring(blob)
            scheme = root.find(".//a:clrScheme", _NS)
            if scheme is not None:
                by_name = {}
                for node in scheme:
                    value = next(iter(node), None)
                    if value is not None:
                        by_name[node.tag.rsplit("}", 1)[-1]] = value.get("val") if value.tag.endswith("srgbClr") else value.get("lastClr")
                colors = [by_name.get(name) for name in
                          ("lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3",
                           "accent4", "accent5", "accent6", "hlink", "folHlink")]
            for kind in ("major", "minor"):
                node = root.find(f".//a:{kind}Font/a:latin", _NS)
                if node is not None:
                    fonts[kind] = node.get("typeface")
        return colors, fonts

    def resolve_dependencies(self):
        pending = list(self.required)
        checked = set()
        globals_by_name = {n.name.casefold(): n for n in self.book.defined_names.values()}

        def collect(expression, sheet=None, seen=None):
            seen = set() if seen is None else seen
            result = set()
            local = {} if sheet is None else {
                n.name.casefold(): n for n in self.book[sheet].defined_names.values()}
            for token in _tokens(expression)[0]:
                if token.type != "OPERAND" or token.subtype != "RANGE":
                    continue
                value = token.value
                if "!" in value:
                    result.update(_sheet_reference(value, self.book.sheetnames))
                    continue
                match = _TABLE.match(value)
                if match:
                    table = match.group(2).casefold()
                    if table not in self.table_sheets:
                        raise ValueError(f"{self.path.name} 引用了不存在的 Excel 表：{match.group(2)}")
                    result.add(self.table_sheets[table])
                    continue
                key = value.casefold()
                defined = local.get(key) or globals_by_name.get(key)
                if defined and key not in seen:
                    result.update(collect(defined.attr_text, sheet, seen | {key}))
            return result

        # Preserve global names as workbook features, bringing in their dependent
        # sheets even if a name is not currently used by a visible cell formula.
        for name in self.book.defined_names.values():
            self.required.update(collect(name.attr_text))
        pending = list(self.required)
        while pending:
            sheet = pending.pop()
            if sheet in checked:
                continue
            checked.add(sheet)
            for expression in _expressions(self.book[sheet]):
                for dependency in collect(expression, sheet):
                    if dependency not in self.required:
                        self.required.add(dependency)
                        pending.append(dependency)
        return self.required

    def rewrite(self, expression, sheet=None):
        tokens, leading_equal = _tokens(expression)
        if not tokens:
            return expression
        local = set() if sheet is None else {n.casefold() for n in self.book[sheet].defined_names}
        for token in tokens:
            if token.type != "OPERAND" or token.subtype != "RANGE":
                continue
            value = token.value
            if "!" in value:
                names = _sheet_reference(value, self.book.sheetnames)
                if any(name not in self.sheet_map for name in names):
                    raise ValueError("内部错误：公式依赖未被纳入输出")
                ref = self.sheet_map[names[0]]
                if len(names) > 1:
                    ref += ":" + self.sheet_map[names[-1]]
                tail = value.rsplit("!", 1)[1]
                sheet_local = {n.casefold() for n in self.book[names[0]].defined_names}
                if len(names) == 1 and tail.casefold() not in sheet_local:
                    tail = self.name_map.get(tail.casefold(), tail)
                token.value = quote_sheetname(ref) + "!" + tail
            else:
                match = _TABLE.match(value)
                if match:
                    token.value = match.group(1) + self.table_map.get(match.group(2).casefold(),
                                                                     match.group(2)) + match.group(3)
                elif value.casefold() not in local:
                    token.value = self.name_map.get(value.casefold(), value)
        return ("=" if leading_equal else "") + "".join(t.value for t in tokens)

    def color(self, color):
        if color is None:
            return None
        if color.type == "theme":
            index = color.theme
            if index >= len(self.theme_colors) or not self.theme_colors[index]:
                raise ValueError(f"{self.path.name} 无法解析主题颜色 {index}")
            rgb = self.theme_colors[index]
        elif color.type == "indexed" and color.indexed < len(self.book._colors):
            rgb = self.book._colors[color.indexed][-6:]
        else:
            return copy(color)
        rgb = rgb[-6:]
        if color.tint:
            r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
            h, lightness, saturation = colorsys.rgb_to_hls(r, g, b)
            tint = color.tint
            lightness = lightness * (1 + tint) if tint < 0 else lightness * (1 - tint) + tint
            rgb = "".join(f"{round(v * 255):02X}" for v in colorsys.hls_to_rgb(h, lightness, saturation))
        return Color(rgb="FF" + rgb.upper())

    def font(self, font):
        result = copy(font)
        result.color = self.color(font.color)
        if font.scheme in self.theme_fonts:
            result.name = self.theme_fonts[font.scheme] or font.name
            result.scheme = None
        return result

    def fill(self, fill):
        # StyleProxy implements __copy__, but cannot be deep-copied directly.
        result = copy(fill)
        if hasattr(result, "fgColor"):
            result.fgColor, result.bgColor = self.color(fill.fgColor), self.color(fill.bgColor)
        elif hasattr(result, "stop"):
            for stop in result.stop:
                stop.color = self.color(stop.color)
        return result

    def border(self, border):
        result = copy(border)
        for name in ("left", "right", "top", "bottom", "diagonal", "vertical", "horizontal", "start", "end"):
            side = getattr(border, name, None)
            if side:
                copied = copy(side)
                copied.color = self.color(side.color)
                setattr(result, name, copied)
        return result

    def styles(self, source, target):
        # Set properties individually: _style contains indices into the SOURCE
        # workbook and must never be assigned directly across workbooks.
        target.font = self.font(source.font)
        target.fill = self.fill(source.fill)
        target.border = self.border(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)
        if hasattr(source, "quotePrefix"):
            target.quotePrefix = source.quotePrefix

    def _remap_dxfs(self, obj, destination, seen=None):
        """Filter/table style IDs refer to source resources just like cell styles."""
        if obj is None:
            return
        seen = set() if seen is None else seen
        if id(obj) in seen:
            return
        seen.add(id(obj))
        for attr in getattr(obj, "__attrs__", ()):
            value = getattr(obj, attr, None)
            if attr.lower().endswith("dxfid") and value is not None:
                if value not in self.dxf_map:
                    styles = self.book._differential_styles.styles
                    if not 0 <= value < len(styles):
                        raise ValueError(f"{self.path.name} 含无效差异样式引用 {value}")
                    style = deepcopy(styles[value])
                    for name, convert in (("font", self.font), ("fill", self.fill), ("border", self.border)):
                        if getattr(style, name) is not None:
                            setattr(style, name, convert(getattr(style, name)))
                    self.dxf_map[value] = destination._differential_styles.add(style)
                setattr(obj, attr, self.dxf_map[value])
            if attr.endswith("CellStyle") and value:
                raise ValueError(f"{self.path.name} 的 Excel 表引用命名单元格样式 {value}，"
                                 "暂不能安全合并该样式关联")
        for attr in getattr(obj, "__elements__", ()):
            value = getattr(obj, attr, None)
            for child in value if isinstance(value, (tuple, list)) else (value,):
                self._remap_dxfs(child, destination, seen)

    def copy_sheet(self, source, target, cancel):
        for index, (position, cell) in enumerate(source._cells.items()):
            if index % 512 == 0:
                _cancel(cancel)
            if isinstance(cell, MergedCell):
                continue
            dst = target.cell(*position)
            if cell.data_type == "f":
                if isinstance(cell.value, str):
                    dst.value = self.rewrite(cell.value, source.title)
                else:
                    dst.value = deepcopy(cell.value)
                    if getattr(dst.value, "text", None):
                        dst.value.text = self.rewrite(dst.value.text, source.title)
            else:
                dst.value = deepcopy(cell.value)
                dst.data_type = cell.data_type
                if isinstance(dst.value, CellRichText):
                    for run in dst.value:
                        if isinstance(run, TextBlock):
                            run.font.color = self.color(run.font.color)
                            if run.font.scheme in self.theme_fonts:
                                run.font.rFont = self.theme_fonts[run.font.scheme] or run.font.rFont
                                run.font.scheme = None
            self.styles(cell, dst)
            if cell.comment:
                dst.comment = copy(cell.comment)
            if cell.hyperlink:
                dst._hyperlink = copy(cell.hyperlink)
                if dst.hyperlink.location:
                    dst.hyperlink.location = self.rewrite(dst.hyperlink.location, source.title)
                elif dst.hyperlink.target and dst.hyperlink.target.startswith("#"):
                    dst.hyperlink.target = "#" + self.rewrite(dst.hyperlink.target[1:], source.title)
        for dimension_name in ("row_dimensions", "column_dimensions"):
            for key, dimension in getattr(source, dimension_name).items():
                cloned = copy(dimension)
                cloned.parent = target
                if dimension.has_style:
                    cloned._style = None
                    self.styles(dimension, cloned)
                getattr(target, dimension_name)[key] = cloned
        for attr in ("sheet_format", "sheet_properties", "page_margins", "page_setup",
                     "print_options", "HeaderFooter", "protection", "auto_filter",
                     "row_breaks", "col_breaks", "scenarios"):
            setattr(target, attr, deepcopy(getattr(source, attr)))
        # The output has one workbook window. Source view IDs belong to the
        # source book's windows; retaining them can leave dangling references.
        view = deepcopy(source.sheet_view)
        view.workbookViewId = 0
        target.views = SheetViewList(sheetView=[view])
        if target.sheet_properties.tabColor is not None:
            target.sheet_properties.tabColor = self.color(target.sheet_properties.tabColor)
        target.page_setup._parent = target
        # Printer binary settings and their relationships are not copied.
        target.page_setup.id = None
        self._remap_dxfs(target.auto_filter, target.parent)
        target.sheet_state = source.sheet_state
        for merged in source.merged_cells.ranges:
            target.merge_cells(str(merged))
        for position, cell in source._cells.items():
            if isinstance(cell, MergedCell) and cell.has_style:
                self.styles(cell, target.cell(*position))
        # Pane and selection are already copied together in the retained view.
        # Setting freeze_panes would rebuild a split pane as a frozen pane.
        target.print_title_rows = source.print_title_rows
        target.print_title_cols = source.print_title_cols
        if source.print_area:
            target.print_area = self.rewrite(source.print_area, source.title)
        target.data_validations = deepcopy(source.data_validations)
        for dv in target.data_validations.dataValidation:
            for attr in ("formula1", "formula2"):
                value = getattr(dv, attr)
                if value:
                    setattr(dv, attr, self.rewrite(value, source.title))
        for cf in source.conditional_formatting:
            for rule in source.conditional_formatting[cf]:
                rule = deepcopy(rule)
                rule.formula = [self.rewrite(formula, source.title) for formula in rule.formula or []]
                rule.dxfId = None
                if rule.dxf:
                    for attr, convert in (("font", self.font), ("fill", self.fill), ("border", self.border)):
                        if getattr(rule.dxf, attr) is not None:
                            setattr(rule.dxf, attr, convert(getattr(rule.dxf, attr)))
                if rule.colorScale:
                    rule.colorScale.color = [self.color(color) for color in rule.colorScale.color]
                if rule.dataBar and rule.dataBar.color:
                    rule.dataBar.color = self.color(rule.dataBar.color)
                target.conditional_formatting.add(str(cf.sqref), rule)
        for table in source.tables.values():
            table = deepcopy(table)
            self._remap_dxfs(table, target.parent)
            if (table.tableStyleInfo and table.tableStyleInfo.name
                    and table.tableStyleInfo.name.casefold() in self.custom_table_styles):
                table.tableStyleInfo = None
            table.name = table.displayName = self.table_map[table.name.casefold()]
            for col in table.tableColumns:
                for formula in (col.calculatedColumnFormula, col.totalsRowFormula):
                    if formula is not None:
                        formula.attr_text = self.rewrite(formula.attr_text, source.title)
            target.add_table(table)
        for name in source.defined_names.values():
            new = deepcopy(name)
            new.attr_text = self.rewrite(new.attr_text, source.title)
            new.localSheetId = target.parent.index(target)
            target.defined_names.add(new)


def merge_workbook(payload, out, cancel=None, progress=None):
    """Copy selected complete sheets into ``out`` (a caller-owned staging path)."""
    report = {"files": [], "totals": {"rows_written": 0, "sheets_merged": 0,
                                     "dependency_sheets": 0},
              "warnings": [], "issues": [], "can_export": False, "groups": []}
    books = []
    destination = None
    try:
        paths = payload.get("paths") or [item["path"] for item in payload.get("selections", [])]
        if not isinstance(paths, list) or not paths:
            raise ValueError("请至少选择一个 .xlsx 文件")
        unique_paths = list(dict.fromkeys(Path(p).expanduser().resolve() for p in paths))
        out = Path(out).expanduser().resolve()
        if out in unique_paths:
            raise ValueError("输出路径不能覆盖输入工作簿")
        if out.suffix.lower() != ".xlsx":
            raise ValueError("输出路径必须为 .xlsx")
        selections = payload.get("selections")
        chosen = None
        if selections is not None:
            if not isinstance(selections, list) or not selections:
                raise ValueError("请选择至少一个工作表")
            chosen = OrderedDict()
            for item in selections:
                path = Path(item["path"]).expanduser().resolve()
                if path not in unique_paths:
                    raise ValueError("所选工作表的文件不在输入文件列表中")
                if not item.get("sheet"):
                    raise ValueError("每个选择项必须明确指定工作表")
                chosen.setdefault(path, set()).add(item["sheet"])
            report["warnings"].append("各表独立成页模式复制完整工作表；所选区域仅在拼接模式中使用")
        for path in unique_paths:
            _cancel(cancel)
            if chosen is not None and path not in chosen:
                report["files"].append({"path": str(path), "status": "skipped"})
                continue
            source = InputBook(path, None if chosen is None else chosen[path])
            books.append(source)
            report["warnings"].extend(source.warnings)
            source.resolve_dependencies()
            report["files"].append({"path": str(path), "status": "ok"})
        if not books:
            raise ValueError("没有可合并的工作表")
        destination = Workbook()
        destination.remove(destination.active)
        # Sheet names and defined/table names occupy case-insensitive namespaces.
        sheet_names, object_names = set(), set()
        for source in books:
            for name in source.book.sheetnames:
                if name in source.required:
                    source.sheet_map[name] = _unique(name, sheet_names, 31)
                    destination.create_sheet(source.sheet_map[name])
            for name in source.book.defined_names.values():
                source.name_map[name.name.casefold()] = _unique(name.name, object_names)
            for ws in source.book:
                if ws.title in source.required:
                    for table in ws.tables.values():
                        output_name = _unique(table.displayName or table.name, object_names)
                        for alias in (table.name, table.displayName):
                            if alias:
                                source.table_map[alias.casefold()] = output_name
        total = sum(len(source.required) for source in books)
        count = 0
        for source in books:
            for name in source.book.sheetnames:
                if name not in source.required:
                    continue
                _cancel(cancel)
                target = destination[source.sheet_map[name]]
                source.copy_sheet(source.book[name], target, cancel)
                dependency = name not in source.selected
                if dependency:
                    report["warnings"].append(f"{source.path.name} 已带入引用依赖页：{name}")
                    report["totals"]["dependency_sheets"] += 1
                report["groups"].append({"sheet": target.title,
                                         "sources": [f"{source.path.name} / {name}"],
                                         "count": 1, "dependency": dependency})
                report["totals"]["rows_written"] += source.book[name].max_row
                count += 1
                if progress:
                    progress(count, total, f"{source.path.name} / {name}", report["totals"]["rows_written"])
            for defined in source.book.defined_names.values():
                new = deepcopy(defined)
                new.name = source.name_map[defined.name.casefold()]
                new.attr_text = source.rewrite(defined.attr_text)
                new.localSheetId = None
                destination.defined_names.add(new)
        if not any(ws.sheet_state == "visible" for ws in destination):
            raise ValueError("所有选中页及依赖页均隐藏；请至少选择一张可见工作表")
        destination.active = next(i for i, ws in enumerate(destination) if ws.sheet_state == "visible")
        destination.calculation.fullCalcOnLoad = True
        destination.calculation.forceFullCalc = True
        _cancel(cancel)
        destination.save(out)
        _cancel(cancel)
        report["totals"]["sheets_merged"] = count
        report["can_export"] = True
        report["out"] = str(out)
        return report
    except Cancelled:
        raise
    except (ValueError, OSError, BadZipFile, InvalidFileException, KeyError, TypeError,
            ElementTree.ParseError) as exc:
        report["issues"].append({"kind": "native_workbook", "message": str(exc), "choices": []})
        return report
    finally:
        for source in books:
            source.book.close()
        if destination:
            destination.close()
