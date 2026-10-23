"""Apply explicit source regions without changing or normalizing source files."""
from pathlib import Path
import json

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .preview import normalize_selection, field_metadata, automatic_selection
from .recognize import MergeRange


def source_spec(options, path, sheet):
    specs = getattr(options, "selections", None) or []
    if not specs:
        return None
    resolved = Path(path).resolve()
    return next((s for s in specs if Path(s["path"]).resolve() == resolved and s["sheet"] == sheet), False)


def apply_selection(src, spec):
    src.field_map = (spec or {}).get("field_map") or {}
    src.field_map_context = (spec or {}).get("field_map_context")
    if not spec or not spec.get("selection"):
        return
    st = src.structure
    # 自动模式的选区本来就是引擎结构识别结果的回显(automatic_selection 由
    # structure 派生), 不得据此重写结构: 重写会用"原始标签键"(同名不去重)
    # 覆盖读取器的去重键, 破坏重名表头的按位对齐; 合并单元格坐标也会被按
    # 转置语义改写。结构保持读取器检测结果即可。
    if spec.get("selection_origin", "manual") == "auto":
        return
    if getattr(st, "native_cross", False):
        raise ValueError("「%s」是横向配置表(记录沿列), 坐标即原表网格, 暂不支持手动框选; 请使用自动识别的范围。" % st.sheet_name)
    if Path(src.path).suffix.lower() != ".xlsx":
        raise ValueError("手动区域处理需 .xlsx 文件；.xls/.csv 可使用原有合并流程")
    wb = load_workbook(src.path, data_only=False)
    try:
        ws = wb[src.sheet_name]
        selected = normalize_selection(spec["selection"], ws.max_row, ws.max_column)
        src.selection = selected
        st.orientation = "horizontal" if selected["axis"] == "columns" else "vertical"
        st.header_row = selected["header_start"]
        st.height = selected["header_end"] - selected["header_start"] + 1
        metadata = field_metadata(ws, selected)
        src.selected_fields = metadata
        st.keys = [f["label"] for f in metadata]
        # Display labels are not identities: a literal "A / B" header differs
        # from a two-level A,B path, even though the rendered label is identical.
        st.key_signatures = [f["match_signature"] for f in metadata]
        st.empty = selected["data_end"] < selected["data_start"]
        st.low_confidence = False
        st.warnings = []
        st.orientation_confident = True
        field_offset = selected["field_start"] - 1
        st.field_offset = field_offset
        st.data_start = selected["data_start"]
        st.data_end = selected["data_end"]
        st.field_end = selected["field_end"]
        merges = []
        for mg in ws.merged_cells.ranges:
            a, b, c, d = (mg.min_col, mg.min_row, mg.max_col, mg.max_row)
            if st.orientation == "horizontal":
                a, b, c, d = b, a, d, c
            if a >= selected["field_start"] and c <= selected["field_end"]:
                merges.append(MergeRange(a - field_offset, b, c - field_offset, d))
            elif a <= selected["field_end"] and c >= selected["field_start"]:
                raise ValueError(f"{src.display} 的选择范围切断合并单元格 {mg}，请扩大字段范围")
        st.all_merges = merges
        st.merges = [m for m in merges if m.min_row >= st.header_row and m.max_row <= selected["header_end"]]
        if st.orientation == "vertical":
            st.col_widths = {i - field_offset: ws.column_dimensions[get_column_letter(i)].width
                             for i in range(selected["field_start"], selected["field_end"] + 1)}
            st.pre_rows = [[ws.cell(r, c).value for c in range(selected["field_start"], selected["field_end"] + 1)]
                           for r in range(1, st.header_row)]
        else:
            st.col_widths = {}
            st.pre_rows = [[ws.cell(r, c).value for r in range(selected["field_start"], selected["field_end"] + 1)]
                           for c in range(1, st.header_row)]
    finally:
        wb.close()


def _current_fields(source):
    if hasattr(source, "selected_fields"):
        return source.selected_fields
    from .align import key_parts
    offset = getattr(source.structure, "field_offset", 0)
    return [{"index": offset + index + 1,
             "match_signature": json.dumps(key_parts(key), ensure_ascii=False, separators=(",", ":")),
             "empty": not bool(key)} for index, key in enumerate(source.structure.keys)]


def _mapped_position(index, snapshot, current):
    """Follow one explicit field identity across a changed source/template layout."""
    index = int(index)
    identity = lambda field: field.get("match_signature", field.get("signature"))
    old = next((f for f in snapshot if int(f["index"]) == index), None)
    if old is None:
        raise ValueError("字段对应缺少原表头信息，请重新打开原表并确认字段对应")
    # Exact positional snapshots also preserve deliberate duplicate/empty mappings.
    if [(int(f["index"]), identity(f)) for f in snapshot] == [(int(f["index"]), identity(f)) for f in current]:
        return index
    signature = identity(old)
    previous = [f for f in snapshot if identity(f) == signature]
    matches = [f for f in current if identity(f) == signature]
    if (signature and not old.get("empty") and not str(signature).startswith("@empty:")
            and len(previous) == len(matches) == 1):
        return int(matches[0]["index"])
    raise ValueError("字段对应所依据的表头已变化，无法唯一确认原字段，请重新打开原表并确认字段对应")


def apply_field_maps(sources, template_index):
    if not sources:
        return
    template = sources[min(template_index, len(sources) - 1)]
    tf = getattr(template.structure, "field_offset", 0)
    for src in sources:
        mapping = getattr(src, "field_map", {})
        context = getattr(src, "field_map_context", None)
        if mapping and context:
            source_fields, target_fields = _current_fields(src), _current_fields(template)
            mapping = {
                str(_mapped_position(target, context.get("target", []), target_fields)):
                    None if source in (None, "") else _mapped_position(source, context.get("source", []), source_fields)
                for target, source in mapping.items()
            }
        sf = getattr(src.structure, "field_offset", 0)
        used, used_targets = set(), set()
        src.field_pairs = {}
        src.unmapped_targets = set()
        src.field_map_template = template
        for target, source in mapping.items():
            ti = int(target) - tf - 1
            if not 0 <= ti < len(template.structure.keys):
                raise ValueError("字段对应超出已选范围")
            if ti in used_targets:
                raise ValueError("同一目标字段不能重复指定")
            used_targets.add(ti)
            if source in (None, ""):
                src.unmapped_targets.add(ti)
                continue
            si = int(source) - sf - 1
            if not 0 <= si < len(src.structure.keys):
                raise ValueError("字段对应超出已选范围")
            if si in used:
                raise ValueError("同一来源字段不能同时对应多个目标字段")
            used.add(si)
            # Keep positional correspondence, not just a renamed key. Renaming
            # duplicate labels would collapse two distinct physical columns.
            src.field_pairs[si] = ti
