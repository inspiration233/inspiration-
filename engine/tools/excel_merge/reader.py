# -*- coding: utf-8 -*-
"""统一读取层：把 .xlsx / .xls / .csv 抹平成同一种"表"抽象。

- 流式：数据行迭代器逐行产出（value, number_format, style_source)，不整表载入。
- 结构扫描只读文件头部（表头探测窗口 + 表头区高度）。
- 合并单元格范围：xlsx 走 zip 内 XML 流式扫描（read_only 模式拿不到）；
  xls 走 xlrd merged_cells；csv 无合并概念。
"""
import csv
import io
import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from . import recognize
from .recognize import MergeRange

SUPPORTED_EXTS = {".xlsx", ".xls", ".csv"}
XLSX_WINDOW = 16   # 结构扫描读取的行数（探测窗口10 + 表头区高4 + 余量）


# --------------------------------------------------------------- 文件级错误

class FileError(Exception):
    """文件不可读（损坏/被占用/不支持的格式）——进报告，不中断。"""

    def __init__(self, path, reason):
        super().__init__(reason)
        self.path = str(path)
        self.reason = reason


# --------------------------------------------------------------- 结构与表源

@dataclass
class SheetStructure:
    sheet_name: str
    header_row: int                      # 1-based
    height: int                          # 表头区行数
    keys: list                           # 展平内部对齐键
    merges: list                         # 表头区内原始 MergeRange（输出模板用）
    all_merges: list = field(default_factory=list)   # 全表合并范围（数据区还原用）
    col_widths: dict = field(default_factory=dict)   # 列宽（保留样式复制用）
    orientation: str = "vertical"         # vertical 常规 | horizontal 横向表(转置读取)
    orientation_confident: bool = True    # 方向检测置信度(低→体检提示人工确认)
    pre_rows: list = field(default_factory=list)       # 表头前置区行(值, 共性检测/拼接用)
    title: str = None                    # 识别到的标题文本（不进总表，进报告）
    low_confidence: bool = False
    warnings: list = field(default_factory=list)
    empty: bool = False                  # 整表无数据（含仅表头）


@dataclass
class TableSource:
    """一张可合并的表 = 某文件的一个 Sheet（csv 视为单 Sheet 文件）。"""
    path: Path
    sheet_name: str
    display: str                         # 文件名 或 文件名/Sheet名
    structure: SheetStructure
    _reader: object = None

    # ---- 数据行迭代（含表头之后的所有行，未做任何清理）----
    def iter_data_rows(self):
        return self._reader.iter_data_rows(self)

    # ---- 表头区快照（含单元格对象，供输出模板与样式复制）----
    def header_cells(self):
        return self._reader.header_cells(self)


# --------------------------------------------------------------- CSV 编码探测

def detect_csv_encoding(path):
    raw = Path(path).open("rb").read(4096)
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    for enc in ("utf-8", "gbk"):
        try:
            raw.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
    return "gb18030"


# --------------------------------------------------------------- xlsx 合并范围（XML 流式）

_MERGE_RE = re.compile(rb'<mergeCell ref="([A-Z0-9:]+)"\s*/>')
_COL_RE = re.compile(rb'<col\s([^>]*?)/>')


def _sheet_xml_paths(zf):
    """sheet名 -> worksheets xml 路径（走 workbook.xml + rels，命名空间无关）。"""
    def local(tag):
        return tag.rsplit("}", 1)[-1]

    wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
    rid_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    name_to_rid = {}
    for el in wb_root.iter():
        if local(el.tag) == "sheet":
            name_to_rid[el.get("name")] = el.get(rid_attr)

    rels_root = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rid_to_target = {}
    for el in rels_root.iter():
        if local(el.tag) == "Relationship":
            rid_to_target[el.get("Id")] = el.get("Target")

    out = {}
    for name, rid in name_to_rid.items():
        target = rid_to_target.get(rid, "")
        if target.startswith("/"):
            target = target.lstrip("/")
        elif not target.startswith("xl/"):
            target = "xl/" + target
        out[name] = target
    return out


def xlsx_column_widths(path, sheet_name):
    """流式解析 <cols> 块（位于 sheetData 之前，扫描很早结束）：自定义列宽。"""
    with zipfile.ZipFile(path) as zf:
        target = _sheet_xml_paths(zf).get(sheet_name)
        if not target:
            return {}
        widths = {}
        carry = b""
        with zf.open(target) as f:
            while True:
                chunk = f.read(1 << 18)
                if not chunk:
                    break
                data = carry + chunk
                for m in _COL_RE.finditer(data):
                    attrs = dict(re.findall(rb'(\w+)="([^"]*)"', m.group(1)))
                    try:
                        mn, mx = int(attrs[b"min"]), int(attrs[b"max"])
                        w = float(attrs[b"width"])
                    except (KeyError, ValueError):
                        continue
                    for c in range(mn, min(mx, mn + 200) + 1):
                        widths[c] = w
                if b"<sheetData" in data:
                    break
                carry = data[-128:]
        return widths


def xlsx_merge_ranges(path, sheet_name):
    """流式扫描 sheet XML 的 mergeCell（大文件不整块载入）。返回 [MergeRange]。"""
    with zipfile.ZipFile(path) as zf:
        target = _sheet_xml_paths(zf).get(sheet_name)
        if not target:
            return []
        refs = []
        carry = b""
        with zf.open(target) as f:
            while True:
                chunk = f.read(1 << 20)
                if not chunk:
                    break
                data = carry + chunk
                for m in _MERGE_RE.finditer(data):
                    refs.append(m.group(1).decode())
                # 保留尾部防标签被切断
                carry = data[-64:]
        return [MergeRange.parse(r) for r in refs]


# --------------------------------------------------------------- 各格式读取器

class _XlsxReader:
    def __init__(self, path):
        self.path = Path(path)
        try:
            self.wb = load_workbook(self.path, read_only=True, data_only=True)
        except Exception as e:
            raise FileError(self.path, f"文件无法读取（{type(e).__name__}）") from e

    def close(self):
        self.wb.close()

    def sheets(self):
        return [(ws.title, ws.sheet_state == "visible") for ws in self.wb.worksheets]

    def build_structure(self, sheet_name, force=0, override=None):
        ws = self.wb[sheet_name]
        window = []
        for row in ws.iter_rows(min_row=1, max_row=XLSX_WINDOW, values_only=True):
            window.append(list(row))
            if len(window) >= XLSX_WINDOW:
                break
        if not window or not any(recognize._row_nonempty(r) for r in window):
            return SheetStructure(sheet_name, 0, 0, [], [], empty=True)

        ranges = xlsx_merge_ranges(self.path, sheet_name)
        orientation, confident = "vertical", True
        if override in ("vertical", "horizontal"):
            orientation, confident = override, True      # 人工指定优先
        elif force == 0 and window and (ws.max_row or 0) <= 20000 and (ws.max_column or 0) <= 500:
            orientation, confident = recognize.detect_orientation(window)
        if orientation == "horizontal":
            return self._structure_transposed(ws, sheet_name, ranges, confident)
        if force > 0:
            det = recognize.detect_header_row(window)
            header_row = force
        else:
            det = recognize.detect_header_row(window)
            header_row = det.header_row or 1
            header_row = recognize.extend_header_up(window, header_row, ranges)
        height = recognize.detect_region_height(window, header_row, ranges)
        keys = recognize.flatten_keys(window, header_row, height, ranges)
        st = SheetStructure(
            sheet_name=sheet_name,
            header_row=header_row,
            height=height,
            keys=keys,
            merges=[mg for mg in ranges if mg.min_row >= header_row
                    and mg.max_row <= header_row + height - 1],
            all_merges=ranges,
            col_widths=xlsx_column_widths(self.path, sheet_name),
            orientation_confident=confident,
            pre_rows=[list(r) for r in window[:header_row - 1]],
            title=recognize.detect_title(window, header_row, ranges),
            low_confidence=det.low_confidence,
        )
        warn = recognize.out_of_bounds_merge_warning(ranges, header_row, height)
        if warn:
            st.warnings.append(warn)
        if det.header_row == 0:
            st.warnings.append("未探测到表头，已按第 1 行处理（可手动指定）")
        if not confident:
            st.warnings.append("方向待确认（疑似交叉表），可在体检中手动切换纵向/横向")
        return st

    def _structure_transposed(self, ws, sheet_name, ranges, confident=True):
        """横向表：在转置视图（行=原列）上跑同一套识别逻辑。"""
        buf = [list(row) for row in ws.iter_rows(values_only=True)]
        ncols = max((len(r) for r in buf), default=0)
        tv = [[(buf[r][c] if c < len(buf[r]) else None) for r in range(len(buf))]
              for c in range(min(ncols, XLSX_WINDOW))]
        ranges = [mg.transposed() for mg in ranges]
        det = recognize.detect_header_row(tv)
        header_row = det.header_row or 1
        header_row = recognize.extend_header_up(tv, header_row, ranges)
        height = recognize.detect_region_height(tv, header_row, ranges)
        keys = recognize.flatten_keys(tv, header_row, height, ranges)
        st = SheetStructure(
            sheet_name=sheet_name, header_row=header_row, height=height,
            keys=keys, orientation="horizontal", orientation_confident=confident,
            pre_rows=[list(r) for r in tv[:header_row - 1]],
            merges=[mg for mg in ranges if mg.min_row >= header_row
                    and mg.max_row <= header_row + height - 1],
            all_merges=ranges)
        warn = recognize.out_of_bounds_merge_warning(ranges, header_row, height)
        if warn:
            st.warnings.append(warn)
        if det.header_row == 0:
            st.warnings.append("未探测到表头，已按第 1 行处理（可手动指定）")
        return st

    def iter_data_rows(self, source):
        ws = self.wb[source.sheet_name]
        st = source.structure
        start = st.header_row + st.height
        if st.orientation == "horizontal":
            # 横向表：原列=记录。整表缓冲（方向探测已限 20000行×500列）
            buf = [list(row) for row in ws.iter_rows()]
            ncols = max((len(r) for r in buf), default=0)
            for c in range(start - 1, ncols):
                yield c + 1, [(buf[r][c] if c < len(buf[r]) else None)
                              for r in range(len(buf))]
            return
        for row_no, row in enumerate(ws.iter_rows(min_row=start), start=start):
            yield row_no, [c for c in row]   # ReadOnlyCell 列表（带 number_format/font/fill）

    def header_cells(self, source):
        ws = self.wb[source.sheet_name]
        st = source.structure
        if st.orientation == "horizontal":
            buf = [list(row) for row in ws.iter_rows()]
            return [[(buf[r][c] if c < len(buf[r]) else None) for r in range(len(buf))]
                    for c in range(st.header_row - 1, st.header_row - 1 + st.height)]
        rows = []
        for row in ws.iter_rows(min_row=st.header_row, max_row=st.header_row + st.height - 1):
            rows.append([c for c in row])
        return rows


class _XlsReader:
    def __init__(self, path):
        import xlrd
        self.path = Path(path)
        try:
            try:
                self.book = xlrd.open_workbook(str(self.path), formatting_info=True)
            except NotImplementedError:
                self.book = xlrd.open_workbook(str(self.path))
        except Exception as e:
            raise FileError(self.path, f"文件无法读取（{type(e).__name__}）") from e
        self._has_merge = hasattr(self.book.sheet_by_index(0), "merged_cells") \
            and getattr(self.book.sheet_by_index(0), "merged_cells", None) is not None

    def close(self):
        pass

    def sheets(self):
        return [(s.name, True) for s in self.book.sheets()]

    def _conv(self, sheet, r, c):
        import xlrd
        try:
            cell = sheet.cell(r, c)
        except Exception:
            return None
        if cell.ctype == xlrd.XL_CELL_TEXT:
            return cell.value
        if cell.ctype == xlrd.XL_CELL_NUMBER:
            v = cell.value
            return int(v) if float(v).is_integer() else v
        if cell.ctype == xlrd.XL_CELL_DATE:
            try:
                tt = xlrd.xldate_as_tuple(cell.value, self.book.datemode)
                return dt_datetime(tt)
            except Exception:
                return cell.value
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return bool(cell.value)
        return None

    def build_structure(self, sheet_name, force=0, override=None):
        sheet = self.book.sheet_by_name(sheet_name)
        if sheet.nrows == 0:
            return SheetStructure(sheet_name, 0, 0, [], [], empty=True)
        window = [[self._conv(sheet, r, c) for c in range(sheet.ncols)]
                  for r in range(min(sheet.nrows, XLSX_WINDOW))]
        ranges = [MergeRange(c1 + 1, r1 + 1, c2 + 1, r2 + 1)
                  for r1, r2, c1, c2 in (sheet.merged_cells if self._has_merge else [])]
        det = recognize.detect_header_row(window)
        header_row = force if force > 0 else recognize.extend_header_up(
            window, det.header_row or 1, ranges)
        height = recognize.detect_region_height(window, header_row, ranges)
        keys = recognize.flatten_keys(window, header_row, height, ranges)
        widths = {}
        try:
            ci = getattr(sheet, "colinfo_map", None) or {}
            widths = {c + 1: round(v.width / 256.0, 2) for c, v in ci.items()}
        except Exception:
            pass
        return SheetStructure(
            sheet_name, header_row, height, keys,
            [mg for mg in ranges if mg.min_row >= header_row and mg.max_row <= header_row + height - 1],
            all_merges=ranges, col_widths=widths,
            pre_rows=[list(r) for r in window[:header_row - 1]],
            title=recognize.detect_title(window, header_row, ranges),
            low_confidence=det.low_confidence,
            warnings=(["未探测到表头，已按第 1 行处理（可手动指定）"] if det.header_row == 0 else []),
        )

    def _row_cells(self, sheet, r, ncols):
        return [_PseudoCell(self._conv(sheet, r, c)) for c in range(ncols)]

    def iter_data_rows(self, source):
        sheet = self.book.sheet_by_name(source.sheet_name)
        st = source.structure
        start = st.header_row + st.height - 1          # 0-based 起始行
        for r in range(start, sheet.nrows):
            yield r + 1, [self._conv(sheet, r, c) for c in range(sheet.ncols)]  # 原生值直传

    def header_cells(self, source):
        sheet = self.book.sheet_by_name(source.sheet_name)
        st = source.structure
        return [self._row_cells(sheet, r, sheet.ncols)
                for r in range(st.header_row - 1, st.header_row - 1 + st.height)]


def dt_datetime(tt):
    import datetime as _dt
    return _dt.datetime(tt[0], tt[1], tt[2], tt[3], tt[4], int(tt[5]))


class _PseudoCell:
    """xls/csv 行单元格的最小替身：值 + General 格式。"""
    __slots__ = ("value", "number_format", "font", "fill", "border", "alignment")

    def __init__(self, value):
        self.value = value
        self.number_format = "General"
        self.font = None
        self.fill = None
        self.border = None
        self.alignment = None


class _CsvReader:
    def __init__(self, path):
        self.path = Path(path)
        self.encoding = detect_csv_encoding(self.path)
        try:
            with self.path.open("r", encoding=self.encoding, newline="") as f:
                head = [row for _, row in zip(range(XLSX_WINDOW), csv.reader(f))]
        except Exception as e:
            raise FileError(self.path, f"文件无法读取（{type(e).__name__}）") from e
        self.window = head
        self.sheet_name = self.path.stem

    def close(self):
        pass

    def sheets(self):
        return [(self.sheet_name, True)]

    def _reopen(self):
        return csv.reader(self.path.open("r", encoding=self.encoding, newline=""))

    def build_structure(self, sheet_name, force=0, override=None):
        window = [r for r in self.window]  # 前导空行参与窗口但计分为 0
        if not any(recognize._row_nonempty(r) for r in window):
            return SheetStructure(sheet_name, 0, 0, [], [], empty=True)
        orientation, confident = "vertical", True
        if override in ("vertical", "horizontal"):
            orientation, confident = override, True
        elif force == 0:
            orientation, confident = recognize.detect_orientation(window)
        if orientation == "horizontal":
            rows = [r for r in self._reopen()]
            ncols = max((len(r) for r in rows), default=0)
            tv = [[(rows[r][c] if c < len(rows[r]) else None)
                   for r in range(len(rows))] for c in range(min(ncols, XLSX_WINDOW))]
            det = recognize.detect_header_row(tv)
            header_row = det.header_row or 1
            keys = recognize.flatten_keys(tv, header_row, 1, [])
            return SheetStructure(sheet_name, header_row, 1, keys, [],
                                  orientation="horizontal",
                                  orientation_confident=confident,
                                  pre_rows=[list(r) for r in tv[:header_row - 1]],
                                  low_confidence=det.low_confidence)
        det = recognize.detect_header_row(window)
        header_row = force if force > 0 else (det.header_row or 1)
        keys = recognize.flatten_keys(window, header_row, 1, [])
        return SheetStructure(
            sheet_name, header_row, 1, keys, [],
            pre_rows=[list(r) for r in window[:header_row - 1]],
            title=recognize.detect_title(window, header_row, []),
            low_confidence=det.low_confidence,
            warnings=(["未探测到表头，已按第 1 行处理（可手动指定）"] if det.header_row == 0 else []),
        )

    def iter_data_rows(self, source):
        start = source.structure.header_row          # csv 高度恒 1
        if source.structure.orientation == "horizontal":
            rows = [r for r in self._reopen()]
            ncols = max((len(r) for r in rows), default=0)
            for c in range(start, ncols + 1):
                yield c, [((rows[r][c - 1] if c - 1 < len(rows[r]) else "") or None)
                          for r in range(len(rows))]
            return
        for i, row in enumerate(self._reopen(), start=1):
            if i > start:
                yield i, [v if v != "" else None for v in row]   # 原生值直传

    def header_cells(self, source):
        r = source.structure.header_row
        if source.structure.orientation == "horizontal":
            rows = [row for row in self._reopen()]
            col = [((rows[i][r - 1] if r - 1 < len(rows[i]) else "") or None)
                   for i in range(len(rows))]
            return [[_PseudoCell(v) for v in col]]
        row = self.window[r - 1] if r - 1 < len(self.window) else []
        return [[_PseudoCell(v if v != "" else None) for v in row]]


# --------------------------------------------------------------- 入口

def open_reader(path):
    ext = Path(path).suffix.lower()
    if ext not in SUPPORTED_EXTS:
        raise FileError(path, f"不支持的格式 {ext}（支持 xlsx/xls/csv）")
    if not Path(path).exists():
        raise FileError(path, "文件不存在")
    if ext == ".xlsx":
        return _XlsxReader(path)
    if ext == ".xls":
        return _XlsReader(path)
    return _CsvReader(path)


def list_tables(path, include_hidden=False, force_header_row=0, orientation_override=None):
    """打开文件 → [(TableSource, 已跳过原因?)]。损坏文件抛 FileError 由上层记报告。"""
    reader = open_reader(path)
    out = []
    for name, visible in reader.sheets():
        if not visible and not include_hidden:
            out.append((None, f"隐藏 Sheet「{name}」已跳过"))
            continue
        st = reader.build_structure(name, force_header_row, orientation_override)
        src = TableSource(Path(path), name, display_table(Path(path), name),
                          st, reader)
        out.append((src, None))
    return out, reader


def display_table(path, sheet_name):
    stem = Path(path).stem
    if Path(path).suffix.lower() == ".csv":
        return stem
    return f"{stem}/{sheet_name}"
